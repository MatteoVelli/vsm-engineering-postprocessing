"""Live Statistics agro values; Python reporting remains the numerical reference."""
import math

import numpy as np
from openpyxl.utils import get_column_letter, quote_sheetname

from .agro_reporting import reporting_metric_rows
from .excel_formulas import write_formula


def write_agro_statistics_formulas(workbook, result, metadata):
    agro = result.reporting_agro
    if agro is None:
        return
    metrics = tuple(m for m in (agro.capacity, agro.rate, agro.mass_rate, agro.net_mass_rate) if m is not None)
    formulas = {metric.metric_id: '="N/A"' for metric in metrics}
    support = _Support(workbook['Formula Support'], result.dataset.quality.sample_count)

    def source_values(name):
        item = result.resolution.resolved.get(name)
        channel = item.channel if item else next(c for c in result.dataset.channels if c.source_name == name)
        return result.dataset.values[:, result.dataset.channel_index(channel.channel_id)]

    mass_guard = None
    binding = getattr(workbook, '_report_agro_source', {})
    if agro.mass_values is not None:
        mass = support.extent(binding['mass'])
        raw = support.extent(binding['raw'])
        minimum_max = '>0' if agro.trace['mass_source']['policy'] == 'force_inference' else '>=0'
        mass_guard = f'{support.complete(raw)},{support.complete(mass)},MAX({mass}){minimum_max}'
        formulas[agro.capacity.metric_id] = support.checked(mass_guard, f'MAX({mass})')

    d_item = result.resolution.resolved.get('track_distance')
    if d_item is not None and d_item.effective_unit in {'m', 'km'} and support.end >= 6:
        distance_values = source_values('track_distance')
        scale = 1000 if d_item.effective_unit == 'km' else 1
        if agro.mass_values is not None:
            source_distance = support.extent(binding['distance'])
            distance = support.vector('Agro travelled-distance input [m]', distance_values * scale,
                lambda row: f'IF(ISNUMBER({support.at(source_distance, row)}),{support.at(source_distance, row)}*{scale},"N/A")')
        else:
            # Direct volume/flow can remain valid without a mass source.
            distance = support.vector('Agro cumulative distance [m]: ' + d_item.channel.source_name,
                                      distance_values * scale)
        travelled = f'({support.at(distance, support.end)}-{support.at(distance, 5)})'
        distance_guard = (f'{support.complete(distance)},MIN({distance})>=0,{travelled}>0,'
                          f'{support.monotonic(distance)},{support.no_jump(distance, distance_values * scale)}')
        if mass_guard is not None:
            # The net kg/km KPI deliberately uses the final counter rather than
            # travelled distance, and permits intermediate refills with warnings.
            if agro.net_mass_rate is not None:
                final_distance = support.at(distance, support.end)
                net_guard = (f'{mass_guard},MIN({mass})>=0,{support.complete(distance)},MIN({distance})>=0,'
                             f'{support.monotonic(distance)},{final_distance}>0,'
                             f'{support.at(mass,5)}>={support.at(mass,support.end)}')
                formulas[agro.net_mass_rate.metric_id] = support.checked(net_guard,
                    f'({support.at(mass,5)}-{support.at(mass,support.end)})/({final_distance}/1000)')
            mass_steps = support.no_jump(mass, agro.mass_values)
            diagnostic_guard = (f'{mass_guard},MIN({mass})>=0,{distance_guard},'
                                f'{support.monotonic(mass, decreasing=True)},{mass_steps}')
            lost_mass = f'({support.at(mass, 5)}-{support.at(mass, support.end)})'
            formulas[agro.mass_rate.metric_id] = support.checked(diagnostic_guard, f'{lost_mass}/{travelled}')
            if not metadata.reporting_inputs.agro_volume_source and not metadata.reporting_inputs.agro_flow_source:
                density = support.parameter(metadata.reporting_inputs.density_kg_per_L,
                    'Agro density [kg/L]; ' + (metadata.reporting_inputs.density_provenance or 'not supplied'))
                formulas[agro.rate.metric_id] = support.checked(
                    f'{diagnostic_guard},ISNUMBER({density}),{density}>0',
                    f'{lost_mass}/({density}*{travelled})')

        # Preserve existing explicitly selected volumetric modes. They must never
        # silently switch to mass/density when their configured source is invalid.
        if agro.rate.value is not None and agro.trace['method'] in {
                'direct cumulative volume', 'direct volumetric flow integration'}:
            inputs = metadata.reporting_inputs
            name = inputs.agro_volume_source or inputs.agro_flow_source
            values = source_values(name)
            volume_source = support.vector('Agro configured volume/flow: ' + name, values)
            guard = f'{distance_guard},{support.complete(volume_source)},MIN({volume_source})>=0'
            guard += ',' + support.no_jump(volume_source, values)
            if inputs.agro_volume_source:
                guard += ',' + support.monotonic(volume_source)
                volume = f'({support.at(volume_source, support.end)}-{support.at(volume_source, 5)})'
            else:
                times = support.vector('Agro flow time [s]', source_values('track_time'))
                guard += f',{support.complete(times)},{support.monotonic(times, strict=True)}'
                next_f, previous_f = support.adjacent(volume_source)
                next_t, previous_t = support.adjacent(times)
                unit = agro.trace['direct_flow_resolution']['raw_unit'].lower()
                divisor = 60 if unit.endswith('/min') else 3600 if unit.endswith('/h') else 1
                volume = f'SUMPRODUCT(({next_f}+{previous_f})/2,({next_t}-{previous_t}))/{divisor}'
            formulas[agro.rate.metric_id] = support.checked(guard, f'({volume})/{travelled}')

    first_row = len(result.statistics) + len(result.kpis) + 8
    by_id = {metric.metric_id: metric for metric in metrics}
    for row, item in enumerate(reporting_metric_rows(result, metadata), first_row):
        if item['metric_id'] in formulas:
            metric = by_id[item['metric_id']]
            write_formula(workbook['Statistics'].cell(row, 3), formulas[metric.metric_id],
                          metric.value if metric.value is not None else 'N/A')


class _Support:
    """Append auditable helpers without moving existing data or parameter cells."""
    def __init__(self, sheet, sample_count):
        self.sheet = sheet
        self.end = sample_count + 4
        self.column = sheet.max_column + 1
        self.parameter_row = max(sheet.max_row, self.end) + 2

    def vector(self, label, values, formula=None):
        column = self.column
        self.column += 1
        self.sheet.cell(3, column, label)
        for row, value in enumerate(values, 5):
            cached = float(value) if math.isfinite(float(value)) else 'N/A'
            cell = self.sheet.cell(row, column)
            if formula:
                write_formula(cell, '=' + formula(row), cached)
            else:
                cell.value = cached
        letter = get_column_letter(column)
        return f'{quote_sheetname(self.sheet.title)}!${letter}$5:${letter}${self.end}'

    def extent(self, location):
        sheet, column = location
        letter = get_column_letter(column)
        return f'{quote_sheetname(sheet.title)}!{letter}5:{letter}{self.end}'

    def parameter(self, value, description):
        row = self.parameter_row
        self.parameter_row += 1
        self.sheet.cell(row, 1, description)
        self.sheet.cell(row, 2, value)
        return f'{quote_sheetname(self.sheet.title)}!$B${row}'

    @staticmethod
    def at(extent, row):
        return f'INDEX({extent},{row - 4})'

    def adjacent(self, extent):
        return f'{self.at(extent, 6)}:{self.at(extent, self.end)}', f'{self.at(extent, 5)}:{self.at(extent, self.end-1)}'

    @staticmethod
    def complete(extent):
        return f'COUNT({extent})=ROWS({extent})'

    def monotonic(self, extent, decreasing=False, strict=False):
        following, previous = self.adjacent(extent)
        violation = '>' if decreasing else '<=' if strict else '<'
        return f'SUMPRODUCT(--(({following}){violation}({previous})))=0'

    def no_jump(self, extent, values):
        # MEDIAN ignores text blanks. Only nonzero absolute steps participate,
        # exactly like calculate_agro_kpis.jump(), without array-entered formulas.
        column = self.column
        self.column += 1
        self.sheet.cell(3, column, 'Agro nonzero absolute steps: ' + extent)
        for row, delta in enumerate(np.abs(np.diff(values)), 6):
            step = f'ABS({self.at(extent, row)}-{self.at(extent, row-1)})'
            formula = f'=IFERROR(IF({step}>0,{step},""),"N/A")'
            cached = float(delta) if math.isfinite(float(delta)) and delta > 0 else '' if delta == 0 else 'N/A'
            write_formula(self.sheet.cell(row, column), formula, cached)
        letter = get_column_letter(column)
        steps = f'{quote_sheetname(self.sheet.title)}!${letter}$6:${letter}${self.end}'
        return f'IF(COUNT({steps})>=5,MAX({steps})<=20*MEDIAN({steps}),TRUE)'

    @staticmethod
    def checked(guard, body):
        return f'=IFERROR(IF(AND({guard}),{body},"N/A"),"N/A")'
