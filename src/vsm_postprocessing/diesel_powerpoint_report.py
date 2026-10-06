"""Diesel presentation content using the approved Astauto template and D1 results.

This adapter changes presentation only. It consumes the Excel result's canonical
metrics and existing plot assets; it never calculates engineering quantities.
"""
from __future__ import annotations

from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path
from typing import Any

import numpy as np
from pptx import Presentation
from pptx.util import Inches

from . import profile_powerpoint_report_engine as ppt


# Template page, content type, title, subtitle. Metric/plot selections live in YAML.
_SLIDES = {
    "cover": (1, "cover", None, "Deterministic Diesel simulation report"),
    "system_overview": (2, "overview", "System & Simulation Overview", "Diesel profile and source-data context"),
    "executive_results": (3, "kpi_grid", "Executive Results", "Vehicle, engine and fuel results from the validated profile"),
    "vehicle_operation": (4, "plot_full", "Vehicle Operation", "Measured speed and road gradient over the drive cycle"),
    "diesel_engine": (5, "plot_pair", "Diesel Engine Performance", "Engine speed, torque and mechanical power; negative overrun retained"),
    "diesel_fuel_energy": (6, "plot_pair", "Fuel Consumption & Mechanical Energy", "Cumulative fuel mass, volumetric flow and net engine energy"),
    "diesel_driveline": (7, "plot_pair", "Driveline & Wheel Loads", "Corner driveshaft torques and total tyre vertical loads"),
    "diesel_controls": (8, "plot_pair", "Engine Load & Throttle", "Source engine load torque [Nm] and throttle command [%]"),
    "diesel_thermal": (9, "plot_full", "Engine Thermal Behaviour", "Source oil temperature; no derived thermal assessment"),
    "simulation_summary": (10, "conclusion", "Simulation Summary", "Deterministic Diesel result snapshot"),
}
_OPTIONAL_SLIDES = {"diesel_driveline", "diesel_controls", "diesel_thermal"}
_PANELS = {4: (20,), 5: (28, 29), 6: (24, 25), 7: (28, 29), 8: (20, 21), 9: (28, 29)}
_LABELS = {
    "chassis_speed_max": "MAX SPEED", "fuel_consumed_kg": "FUEL CONSUMED",
    "engine_speed_max": "MAX ENGINE SPEED", "engine_speed_average": "AVG ENGINE SPEED",
    "engine_torque_max": "MAX ENGINE TORQUE", "engine_power_max": "MAX ENGINE POWER",
    "engine_power_min": "MIN ENGINE POWER", "engine_power_rms": "ENGINE POWER RMS",
    "engine_mechanical_energy_last": "NET MECHANICAL ENERGY",
    "fuel_volume_last": "INTEGRATED FUEL VOLUME", "fuel_flow_average": "AVG FUEL FLOW",
    "fuel_flow_max": "MAX FUEL FLOW", "engine_load_max": "MAX ENGINE LOAD TORQUE",
    "engine_throttle_max": "MAX THROTTLE",
}


def _statistics(result):
    """Presentation labels only; canonical values and source units are retained."""
    statistics = ppt._profile_statistics_as_powerpoint_statistics(result)
    statistics.statistics = [replace(s, display_name=_LABELS.get(s.statistic_id, s.display_name))
                             for s in statistics.statistics]
    return statistics


def _display(stat):
    text = ppt._format_statistic_value(stat).replace(" l/h", " L/h")
    return text[:-2] + " L" if text.endswith(" l") else text


def _available_plots(result):
    # A present, nonzero constant trace is still useful (e.g. a held temperature).
    return {p.plot_id for p in result.plotting_result.rendered_plots
            if any(not s.is_all_zero for s in result.plotting_result.series_summaries[p.plot_id])}


def diesel_powerpoint_config(result: ppt.ProfileExcelReportResult, *, output_filename: str) -> dict[str, Any]:
    profile = result.profile
    statistics = {s.statistic_id: s for s in _statistics(result).statistics}
    plots = _available_plots(result)
    definitions = profile.presentation.slides_by_id()
    missing = set(_SLIDES) - definitions.keys()
    if missing:
        raise ppt.PowerPointReportError("Diesel presentation is missing slides: " + ", ".join(sorted(missing)))
    value = lambda name: _display(statistics[name])
    bodies = {
        "executive_results": ("Canonical simulation results; no acceptance criteria have been applied.",),
        "diesel_driveline": ("Available corners are shown individually; missing optional channels are omitted.",),
        "simulation_summary": (
            f"Run covered {value('distance_km')} in {value('drive_cycle_minutes')}; maximum speed {value('chassis_speed_max')}.",
            f"Fuel consumed: {value('fuel_consumed_kg')}; net mechanical energy: {value('engine_mechanical_energy_last')}.",
            f"Engine power ranged from {value('engine_power_min')} to {value('engine_power_max')}; maximum engine speed {value('engine_speed_max')}.",
            "Results describe this simulation; acceptance conclusions require explicit engineering criteria.",
        ),
    }
    slides = []
    for slide_id, (_, kind, title, subtitle) in _SLIDES.items():
        definition = definitions[slide_id]
        plot_ids = [p for p in definition.plots if p in plots]
        if slide_id in _OPTIONAL_SLIDES and not plot_ids:
            continue
        slides.append({
            "slide_id": slide_id, "type": kind, "title": title or result.report_metadata.report_title,
            "subtitle": subtitle, "statistics": [s for s in definition.statistics if s in statistics],
            "plots": plot_ids, "body": list(bodies.get(slide_id, ())),
        })
    return {
        "version": 1,
        "presentation": {
            "title": result.report_metadata.report_title, "output_filename": output_filename,
            "author": "VSM Engineering", "subject": "Deterministic Diesel engineering report",
            "keywords": "VSM, Diesel, deterministic engineering report",
            "footer": profile.presentation.footer, "comments": f"Generated from canonical profile results by v{ppt.__version__}.",
        },
        "slides": slides,
    }


def _remove(shape):
    shape.element.getparent().remove(shape.element)


def _metadata_pills(shapes, page, result):
    indices = (6, 8, 10, 12, 14) if page == 1 else (5, 7, 9, 11, 13)
    texts = (f"Machine {result.report_metadata.machine_name}", "Powertrain Diesel", "Source VSM simulation",
             f"Samples {result.sample_count:,}", f"Tool v{ppt.__version__}")
    left = Inches(.6)
    for index, width, text in zip(indices, (3.2, 1.9, 2.8, 1.8, 1.79), texts):
        for shape in (shapes[index - 1], shapes[index]):
            shape.left, shape.width = left, Inches(width)
        ppt._set_shape_text(shapes[index], text)
        left += Inches(width + .16)


def _overview(shapes, result):
    ppt._set_shape_text(shapes[17], "Source & Drive Cycle")
    ppt._set_overview_card_body(shapes[18], "\n".join((
        # Zero-width spaces allow wrapping at separators without changing the filename.
        "Source: " + result.report_metadata.source_filename.replace("_", "_\u200b"),
        f"Profile: {result.profile.metadata.name}",
        f"Imported samples: {result.sample_count:,}; source channels: {result.source_raw_channel_count:,}.",
    )))
    ppt._set_shape_text(shapes[22], "Diesel Data Context")
    height = "Road height available." if "track_height" in result.resolution.resolved else "Road height unavailable; gradient is measured."
    flow = "Fuel flow available." if "fuel_flow" in result.resolution.resolved else "Fuel flow unavailable; cumulative mass retained."
    ppt._set_overview_card_body(shapes[23], "\n".join((
        "Engine speed, torque and cumulative fuel mass resolved.", flow, height,
        f"Optional raw channels unavailable: {len(result.resolution.missing_optional)}.",
    )))


def _replace_plots(slide, shapes, page, plot_ids, plots_by_id, assets, slide_width):
    slots = ppt._ELECTRIC_LAYOUT.plot_slots.get(page, ())
    panels = _PANELS.get(page, ())
    if len(plot_ids) > len(slots):
        raise ppt.PowerPointReportError(f"Diesel slide {page} exceeds the template plot capacity")
    for index, (slot, panel_index) in enumerate(zip(slots, panels)):
        picture, panel = shapes[slot.picture], shapes[panel_index]
        if index >= len(plot_ids):
            _remove(picture)
            _remove(panel)
            continue
        if len(plot_ids) == 1:
            # Keep the approved dimensions and aspect ratio; centre the existing pair slot.
            picture.left = int((slide_width - picture.width) / 2)
            panel.left = picture.left - Inches(.12)
        image = ppt._profile_plot_path(plots_by_id[plot_ids[index]], assets)
        ppt._replace_picture(slide, picture, image)


def build_diesel_powerpoint(
    excel_result: ppt.ProfileExcelReportResult, config_path: Path, output_dir: Path, *, plot_assets_dir: Path,
) -> ppt.PowerPointReportResult:
    config = ppt.load_powerpoint_report_config(config_path)
    template = ppt._resolve_visual_template()
    prs = Presentation(template)
    if len(prs.slides) != 10:
        raise ppt.PowerPointReportError("Diesel requires the approved ten-slide Astauto template")
    statistics_result = _statistics(excel_result)
    statistics = {s.statistic_id: s for s in statistics_result.statistics}
    plots = {p.plot_id: p for p in excel_result.plotting_result.rendered_plots}
    assets = Path(plot_assets_dir).resolve()
    original_slides = list(prs.slides)
    # Reuse the template's neutral overview icon instead of battery/charge pictograms.
    neutral_icon = original_slides[1].shapes[1].image.blob
    retained_pages = {_SLIDES[s.slide_id][0] for s in config.slides}
    for page in range(10, 0, -1):
        if page not in retained_pages:
            slide_id = prs.slides._sldIdLst[page - 1]
            prs.part.drop_rel(slide_id.rId)
            prs.slides._sldIdLst.remove(slide_id)
    for number, definition in enumerate(config.slides, 1):
        page = _SLIDES[definition.slide_id][0]
        slide = original_slides[page - 1]
        shapes = list(slide.shapes)
        slots = ppt._ELECTRIC_LAYOUT.text_slots[page]
        ppt._set_shape_text(shapes[slots.title], definition.title)
        ppt._set_shape_text(shapes[slots.subtitle], definition.subtitle or "")
        ppt._update_footer_preserve_template(shapes[slots.footer], excel_result, config.footer)
        ppt._update_page_number_preserve_template(shapes[slots.page_number], number, len(config.slides))
        if page in (1, 2):
            _metadata_pills(shapes, page, excel_result)
        if page == 2:
            _overview(shapes, excel_result)
            picture = shapes[21]
            rect = picture.left, picture.top, picture.width, picture.height
            _remove(picture)
            slide.shapes.add_picture(BytesIO(shapes[16].image.blob), *rect)
        elif page == 3:
            ppt._set_shape_text(shapes[4], "")  # Remove the template's electrical category legend.
            ppt._set_shape_text(shapes[37], definition.body[0])
        elif page == 7:
            ppt._set_shape_text(shapes[32], definition.body[0])
        elif page == 10:
            banner = f"{_display(statistics['distance_km'])} covered   |   Fuel {_display(statistics['fuel_consumed_kg'])}   |   Net energy {_display(statistics['engine_mechanical_energy_last'])}"
            ppt._set_summary_banner_preserve_runs(shapes[4], banner)
            ppt._set_overview_card_body(shapes[6], "\n".join(definition.body))
        selected = [statistics[s] for s in definition.statistic_ids]
        ppt._replace_kpi_slots(shapes, page, selected, ppt._ELECTRIC_LAYOUT)
        for slot, stat in zip(ppt._ELECTRIC_LAYOUT.kpi_slots[page], selected):
            ppt._set_value_text_preserve_runs(shapes[slot.value], _display(stat))
        _replace_plots(slide, shapes, page, definition.plot_ids, plots, assets, prs.slide_width)
        if page in (5, 6, 7, 8, 9):
            picture = shapes[1]
            rect = picture.left, picture.top, picture.width, picture.height
            _remove(picture)
            slide.shapes.add_picture(BytesIO(neutral_icon), *rect)
        if page == 9:
            values = excel_result.plotting_result.values_by_semantic_name
            oil, coolant = values.get("engine_oil_temperature"), values.get("engine_water_temperature")
            note = "Oil temperature shown; coolant remains available in Excel." if coolant is not None else "Oil temperature shown; coolant source unavailable."
            if coolant is not None and np.array_equal(oil, coolant):
                note = "Oil and coolant source traces are identical in this run; one trace is shown."
            ppt._add_text_box(slide, note, Inches(.6), Inches(6.65), Inches(12.13), Inches(.24),
                              font_size=10.5, color=ppt._ASTAUTO_MUTED, italic=True)
        ppt._add_astauto_logo(slide, prs)
        # Keep the exact dataset identity and omission detail accessible without crowding slides.
        slide.notes_slide.notes_text_frame.text = (
            f"Source: {excel_result.report_metadata.source_filename}\nProfile: {excel_result.profile.profile_id}\n"
            + "\n".join(f"{s.statistic_id}: {s.value!r} {s.channel_unit}" for s in selected)
        )
        # Removed template charts must not remain as unused electrical image assets.
        used_images = {node.get(ppt.qn("r:embed")) for node in slide.element.xpath(".//a:blip")}
        for relationship in list(slide.part.rels.values()):
            if relationship.reltype == ppt.RT.IMAGE and relationship.rId not in used_images:
                slide.part.drop_rel(relationship.rId)
    prs.core_properties.title = config.title
    prs.core_properties.subject = config.subject
    prs.core_properties.author = config.author
    prs.core_properties.keywords = config.keywords
    prs.core_properties.comments = config.comments
    path = Path(output_dir) / config.output_filename
    try:
        prs.save(path)
    except Exception as exc:
        raise ppt.PowerPointReportError(f"Could not save Diesel PowerPoint '{path}': {exc}") from exc
    result = ppt.PowerPointReportResult(
        presentation_path=path, manifest_path=Path(output_dir) / "powerpoint_report_manifest.json",
        summary_path=Path(output_dir) / "powerpoint_report_summary.txt", plot_assets_dir=assets,
        config_path=config_path, config=config, statistics_result=statistics_result,
        plotting_result=excel_result.plotting_result,
    )
    ppt._write_powerpoint_metadata(result)
    path.with_name("diesel_template_reuse.json").write_text(json.dumps({
        "template": template.name, "template_pages": sorted(retained_pages),
        "slide_count": result.slide_count, "plot_count": result.plot_count,
        "omitted_plots": sorted(set(plots) - _available_plots(excel_result)),
        "missing_optional_channels": [c.definition.semantic_name for c in excel_result.resolution.missing_optional],
    }, indent=2), encoding="utf-8")
    return result
