"""Live report formulas with deterministic initial values for non-calculating readers."""

import ast
from io import BytesIO
import math
import re
from zipfile import ZipFile

import numpy as np
from openpyxl.utils import get_column_letter


FORMULA_OPERATIONS = {"max", "min", "last", "first"}


def profile_arithmetic_formula(expression, channel_columns, constants):
    """Translate only pointwise arithmetic, preserving the engine's AST grouping.

    Return a row template; functions/integration or unexported dependencies retain
    Python values. COUNT guards prevent Excel treating missing inputs as zero.
    """
    references = []
    operators = {ast.Add: "+", ast.Sub: "-", ast.Mult: "*", ast.Div: "/", ast.Pow: "^"}

    def translate(node):
        if isinstance(node, ast.Constant) and type(node.value) in {int, float}:
            return repr(float(node.value))
        if isinstance(node, ast.Name):
            if node.id in constants:
                return repr(float(constants[node.id]))
            if node.id not in channel_columns:
                raise ValueError("dependency not exported")
            reference = "$" + get_column_letter(channel_columns[node.id]) + "{row}"
            if reference not in references:
                references.append(reference)
            return reference
        if isinstance(node, ast.BinOp) and type(node.op) in operators:
            return f"({translate(node.left)}{operators[type(node.op)]}{translate(node.right)})"
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            return f"({'+' if isinstance(node.op, ast.UAdd) else '-'}{translate(node.operand)})"
        raise ValueError("not pointwise arithmetic")

    try:
        body = translate(ast.parse(expression, mode="eval").body)
    except (ValueError, SyntaxError):
        return None
    if references:
        return f"=IF(COUNT({','.join(references)})={len(references)},{body},NA())"
    return "=" + body


def write_statistic(cell, operation, start_row, end_row, values, value, nan_policy="error"):
    """Keep unsupported engineering statistics numeric and expose simple statistics."""
    cell.value = value
    if operation not in FORMULA_OPERATIONS:
        return
    col = get_column_letter(cell.column)
    extent = f"{col}{start_row}:{col}{end_row}"
    if operation in {"max", "min"}:
        formula = f"{operation.upper()}({extent})"
    else:
        formula = f"{col}{start_row if operation == 'first' else end_row}"
    if not np.isfinite(values).all():
        if nan_policy == "propagate":
            formula = f"IF(COUNT({extent})=ROWS({extent}),{formula},NA())"
        elif nan_policy == "omit" and operation == "first":
            formula = f"INDEX({extent},MATCH(TRUE,INDEX(ISNUMBER({extent}),0),0))"
        elif nan_policy == "omit" and operation == "last":
            formula = f"LOOKUP(2,1/ISNUMBER({extent}),{extent})"
    write_formula(cell, "=" + formula, value)


def write_formula(cell, formula, value):
    cell.value = formula
    workbook = cell.parent.parent
    if not hasattr(workbook, "_report_formula_values"):
        workbook._report_formula_values = {}
    workbook._report_formula_values.setdefault(cell.parent.title, {})[cell.coordinate] = value


def statistic_formula(operation, extent, nan_policy="omit"):
    """Formula equivalent of sample statistics over the specified numeric range."""
    bodies = {
        "rms": f"SQRT(SUMSQ({extent})/COUNT({extent}))",
        "max": f"MAX({extent})", "min": f"MIN({extent})",
        "last": f"LOOKUP(2,1/ISNUMBER({extent}),{extent})",
        "first": f"INDEX({extent},MATCH(TRUE,INDEX(ISNUMBER({extent}),0),0))",
        "sum": f"SUM({extent})",
    }
    if operation not in bodies:
        return None
    condition = f"COUNT({extent})>0" if nan_policy == "omit" else f"COUNT({extent})=ROWS({extent})"
    return f"=IF({condition},{bodies[operation]},NA())"


def profile_temporal_formula(expression, columns, row, start_row=5, output_column=None):
    """Live observed-interval integration/cumulative formulas for revised profiles.

    Restrict translation to the shared engine's documented functions; unsupported
    expressions keep their deterministic Python results and equation notes.
    """
    node = ast.parse(expression, mode="eval").body
    divisor = 1.0
    if isinstance(node, ast.BinOp) and isinstance(node.op, ast.Div) and isinstance(node.right, ast.Constant):
        divisor = float(node.right.value)
        node = node.left
    if not isinstance(node, ast.Call) or not isinstance(node.func, ast.Name):
        return None
    if not all(isinstance(a, ast.Name) and a.id in columns for a in node.args):
        return None
    args = [get_column_letter(columns[a.id]) for a in node.args]
    name = node.func.id
    if name == "cumulative_sum" and len(args) == 1:
        current = f"${args[0]}{row}"
        if row == start_row:
            return f"=IF(ISNUMBER({current}),{current}/{divisor},NA())"
        if output_column is None:
            return None
        previous = f"${get_column_letter(output_column)}{row-1}"
        return f"=IF(COUNT({previous},{current})=2,{previous}+{current}/{divisor},NA())"
    if name not in {"interval_energy_kwh", "cumulative_trapezoid"} or len(args) != 2:
        return None
    p, t = args
    if row == start_row:
        return f"=IF(COUNT(${p}{row},${t}{row})=2,0,NA())"
    previous = row - 1
    guard = f"COUNT(${p}{previous}:${p}{row},${t}{previous}:${t}{row})=4,${t}{row}>${t}{previous}"
    if name == "interval_energy_kwh":
        body = f"(${p}{previous}/2+${p}{row}/2)*(${t}{row}-${t}{previous})/3600/{divisor}"
    else:
        if output_column is None:
            return None
        last = f"${get_column_letter(output_column)}{previous}"
        guard += f",ISNUMBER({last})"
        body = f"{last}+(${p}{previous}/2+${p}{row}/2)*(${t}{row}-${t}{previous})/{divisor}"
    return f"=IF(AND({guard}),{body},NA())"


def save_report_workbook(workbook, output_path):
    """openpyxl writes formulas but cannot write their cached numeric values.

    Fill only the registered formula caches in the serialized worksheet XML;
    preserve all other package parts. Excel recalculates the live dependencies
    on opening and after edits, rather than treating caches as authoritative.
    """
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    buffer = BytesIO()
    workbook.save(buffer)
    caches = getattr(workbook, "_report_formula_values", {})
    by_part = {f"xl/worksheets/sheet{i}.xml": caches[sheet.title]
               for i, sheet in enumerate(workbook.worksheets, 1) if sheet.title in caches}
    with ZipFile(buffer) as source, ZipFile(output_path, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename in by_part:
                values = by_part[entry.filename]

                def fill_cache(match):
                    coordinate = match[1].decode("ascii")
                    if coordinate not in values or not math.isfinite(values[coordinate]):
                        return match[0]
                    cached = repr(float(values[coordinate])).encode("ascii")
                    return match[0].replace(b"<v></v>", b"<v>" + cached + b"</v>")

                data = re.sub(rb'<c\b[^>]*\br="([A-Z]+[0-9]+)"[^>]*>.*?</c>', fill_cache, data)
            target.writestr(entry, data)
