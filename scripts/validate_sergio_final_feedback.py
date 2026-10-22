"""Regenerate local client reports and independently audit the final refinements.

Requires the pre-change baseline captured under outputs/sergio_final_2026_10_17.
Does not edit source data, commit code or create a release package.
"""
from dataclasses import replace
from io import BytesIO
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple
from pptx import Presentation
from PIL import Image

from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report


ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "outputs/sergio_final_2026_10_17"


def main():
    audits = []
    for baseline in json.loads((DEST / "baseline.json").read_text()):
        assert hashlib.sha256(Path(baseline["source"]).read_bytes()).hexdigest() == baseline["source_sha256"]
        kind = baseline["kind"]
        name = "full_size_sprayer_diesel" if kind == "diesel" else "robosprayer_" + kind
        profile = ROOT / "config/report_profiles/sergio_2026_10" / (name + ".yaml")
        folder = DEST / "reports" / kind
        print("Generating " + kind, flush=True)
        excel = generate_profile_excel_report(baseline["source"], profile, folder / "excel")
        ppt = build_profile_powerpoint_report(excel, folder / "powerpoint")
        actual_hashes = {k: hashlib.sha256(v.tobytes()).hexdigest() for k, v in excel.math_result.values_by_semantic_name.items()}
        assert actual_hashes == baseline["baseline_math"]
        stats = {x.definition.statistic_id: x.value for x in excel.statistics_result.statistics}
        expected = {k: v for k, v in baseline["baseline_statistics"].items() if k != "engine_load_max"}
        assert stats == expected
        assert {x.definition.kpi_id: x.value for x in excel.statistics_result.kpis} == baseline["baseline_kpis"]
        prs = Presentation(ppt.presentation_path)
        assert any(s.has_text_frame and s.text == "Source: " + Path(baseline["source"]).name for s in prs.slides[0].shapes)
        values = excel.plotting_result.values_by_semantic_name
        numerical = {}
        slide_ids = [s.slide_id for s in ppt.powerpoint_result.config.slides]
        previous = Presentation(baseline["presentation"])
        retained = [s for i, s in enumerate(prs.slides) if slide_ids[i] != "tyre_total"]
        assert len(retained) == len(previous.slides)
        preserved_images = 0
        for index, (old_slide, new_slide) in enumerate(zip(previous.slides, retained)):
            old_images = [s.image.blob for s in old_slide.shapes if s.shape_type == 13 and s.top > 3*914400]
            new_images = [s.image.blob for s in new_slide.shapes if s.shape_type == 13 and s.top > 3*914400]
            if kind == "diesel" and index == 7:
                old_images = old_images[1:]  # Only the rejected load-torque plot is removed.
            assert len(old_images) == len(new_images)
            for old_image, new_image in zip(old_images, new_images):
                np.testing.assert_array_equal(np.asarray(Image.open(BytesIO(old_image))), np.asarray(Image.open(BytesIO(new_image))))
            preserved_images += len(new_images)
        if kind == "electric":
            i = slide_ids.index("tyre_total")
            assert slide_ids[i:i+3] == ["tyre_total", "soil_bulldozing", "soil_compaction"]
            raw = []
            for mechanism in ("rolling", "bulldozing", "compaction"):
                raw.append(sum(values[f"tyre_{mechanism}resistancepower_{w}"] for w in ("fl", "fr", "rl", "rr")))
            power = sum(raw)
            t = values["track_time"]
            energy = np.r_[0, np.cumsum((power[:-1]/2+power[1:]/2)*np.diff(t)/3600)]
            np.testing.assert_allclose(values["tyre_total_resistance_power"], power, rtol=3e-14, atol=1e-12)
            np.testing.assert_allclose(values["tyre_total_resistance_energy_accumulated"], energy, rtol=3e-14, atol=1e-12)
            pstats = {s.statistic_id:s.value for s in ppt.powerpoint_result.statistics_result.statistics}
            expected_cards = (max(power), np.trapezoid(power,t)/(t[-1]-t[0]), np.sqrt(np.mean(power**2)), energy[-1], energy[-1]/pstats["distance_km_last"],pstats["distance_km_last"])
            ids = ppt.powerpoint_result.config.slides[i].statistic_ids
            np.testing.assert_allclose([pstats[k] for k in ids], expected_cards,rtol=3e-14,atol=1e-12)
            numerical = dict(zip(ids, expected_cards))
            numerical["soft_soil_energy_kwh"] = float(np.trapezoid(raw[1]+raw[2], t)/3600)
            numerical["rolling_energy_kwh"] = float(np.trapezoid(raw[0], t)/3600)
        if kind == "diesel":
            assert "engine_load" not in {c.channel_id for c in excel.report_channels}
            assert "engine_load" in values
            assert stats["engine_torque_max"] == baseline["baseline_statistics"]["engine_torque_max"]
            assert abs(stats["engine_torque_max"]-936.59) < .01
            text = " ".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
            assert "Engine Load" not in text and "936.59 Nm" in text
            numerical = {"engine_torque_max_nm": stats["engine_torque_max"], "excluded_raw_engine_load_max_nm":float(max(values["engine_load"]))}
        book = load_workbook(excel.report_path)
        cached = load_workbook(excel.report_path, data_only=True)
        references = 0
        formulas = 0
        try:
            columns = {c.channel_id:i for i,c in enumerate(excel.report_channels,1)}
            sheet = book.worksheets[0]
            for channel in excel.report_channels:
                actual = np.array([cached.worksheets[0].cell(r,columns[channel.channel_id]).value for r in range(5,excel.sample_count+5)])
                np.testing.assert_allclose(actual, values[channel.channel_id], rtol=3e-14,atol=1e-12)
            for ws in book:
                for row in ws:
                    for c in row:
                        assert c.data_type != "e" and "#REF!" not in str(c.value)
                        if kind == "diesel":
                            assert "Engine Load Torque" not in str(c.value)
                        formulas += c.data_type == "f"
                for chart in ws._charts:
                    for sub in chart._charts:
                        for series in sub.ser:
                            for axis in (series.xVal, series.yVal):
                                if axis and axis.numRef:
                                    title, bounds = range_to_tuple(axis.numRef.f)
                                    assert title in book.sheetnames
                                    c1,r1,c2,r2 = bounds
                                    assert 1 <= c1 <= c2 <= len(excel.report_channels)
                                    assert r1 == 5 and r2 == excel.sample_count+4
                                    references += 1
            assert len(sheet._charts) == excel.plot_count
        finally:
            book.close(); cached.close()
        # Reuse the generated plot assets to check unusually long metadata without
        # repeating calculations. The basename remains exactly as supplied.
        variants = []
        for label, filename in (("long", "Simulation source (final)_" + "configuration_"*15 + ".xlsx"), ("unbroken", "W"*249+".csv")):
            altered = replace(excel, report_metadata=replace(excel.report_metadata,source_filename=filename))
            variant = build_profile_powerpoint_report(altered, folder / label, output_filename="cover_check.pptx")
            variants.append({"label":label,"presentation":str(variant.presentation_path),"source_filename":filename})
        audit = {"profile":kind,"kind":kind,"source":baseline["source"],"workbook":str(excel.report_path),
                 "presentation":str(ppt.presentation_path),"slides":len(prs.slides),"slide_ids":slide_ids,
                 "charts":excel.plot_count,"samples":excel.sample_count,"formula_cells":formulas,"chart_references":references,
                 "unchanged_math_arrays":len(actual_hashes),"unchanged_statistics":len(stats),"numerical_checks":numerical,
                 "pixel_identical_retained_plot_images":preserved_images,
                 "variants":variants,"status":"PASS"}
        audits.append(audit)
        (DEST/"audit.json").write_text(json.dumps(audits,indent=2))
        print(kind+": numerical, source metadata, workbook and deck structure PASS",flush=True)


if __name__ == "__main__":
    main()
