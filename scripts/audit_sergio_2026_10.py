"""Reproduce the private-reference audit without editing any source workbook.

Run from the repository root: python scripts/audit_sergio_2026_10.py
The JSON evidence is written under outputs/sergio_2026_10_validation.
"""
import json
from pathlib import Path
from zipfile import ZipFile

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from vsm_postprocessing.importer import load_data_file
from vsm_postprocessing.math_engine import _interval_energy_kwh
from vsm_postprocessing.profile_math import calculate_profile_math_channels
from vsm_postprocessing.report_profile import load_reporting_profile
from vsm_postprocessing.utils import normalized_name, normalize_display_unit, sha256_file


ROOT = Path(__file__).resolve().parents[1]
REFERENCES = ROOT / "reference_files"
PROFILES = ROOT / "config/report_profiles/sergio_2026_10"
TEMPLATES = {
    "diesel": "Diesel Template_03.xlsx",
    "electric": "Robo_Sprayer_Electrification_Tamplate_Electric_07.xlsx",
    "hybrid": "Robo_Sprayer_Electrification_Tamplate_Hybrid_07.xlsx",
}
PREVIOUS = {"diesel": "Diesel Template.xlsx",
            "electric": "Robo_Sprayer_Electrification_Tamplate_Electric_05.xlsx",
            "hybrid": "Robo_Sprayer_Electrification_Tamplate_Hybrid_06.xlsx"}
RAW = {
    "diesel": "Full_Size_Sprayer_8500Kg_Diesel_Test_03_Crop_Field_MatLab_20kph_Gradient.csv",
    "electric": "RoboSprayer_3500Kg_Electric_12kph_Batt_50kW_Mot_63RPM_Susp_Cool_Rough_Grad_Discharge.csv",
    "hybrid": "Sprayer_Caiman_SP_9300Kg_Hybrid_12x1Km_4000Kg_Chem.csv",
}


def label_key(value):
    return normalized_name(str(value).replace("Accomulated", "Accumulated").split(" [")[0])


def selection(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book.worksheets[1]
        entries = {str(row[1]).strip(): {"report_name": row[2], "type": row[3], "row": i}
                   for i, row in enumerate(sheet.values, 1)
                   if len(row) > 3 and isinstance(row[1], str) and i > 2}
        return entries
    finally:
        book.close()


def audit_workbook(kind):
    path = REFERENCES / TEMPLATES[kind]
    if not path.exists():
        return {"missing_file": str(path)}
    entries = selection(path)
    book = load_workbook(path, read_only=True, data_only=False)
    top_formulas = {c.coordinate: c.value for row in book.worksheets[0].iter_rows(min_row=1, max_row=2)
                    for c in row if c.data_type == "f"}
    first_formulas = {c.coordinate: c.value for row in book.worksheets[0].iter_rows(min_row=5, max_row=7)
                      for c in row if c.data_type == "f"}
    book.close()
    book = load_workbook(path, read_only=True, data_only=True)
    rows = list(book.worksheets[0].values)
    book.close()
    headers, units = rows[2], rows[3]
    data = np.asarray([[float(v) if isinstance(v, (int, float)) else np.nan for v in row]
                       for row in rows[4:] if isinstance(row[0], (int, float))])
    header_columns = {}
    for i, name in enumerate(headers):
        header_columns.setdefault(label_key(name), []).append(i)
    profile_name = "full_size_sprayer_diesel" if kind == "diesel" else "robosprayer_" + kind
    profile = load_reporting_profile(PROFILES / (profile_name + ".yaml"))
    mapped = {}
    for channel in profile.raw_channels:
        entry = entries.get(channel.source_name)
        label = entry["report_name"] if entry else channel.report_name
        indices = header_columns.get(label_key(label), [])
        if channel.semantic_name == "track_height" and not indices:
            indices = header_columns.get(label_key("Road Heigh"), [])
        if len(indices) != 1:
            continue
        i = indices[0]
        # Cached MATH mass must not be mistaken for the raw force in Newtons.
        if normalize_display_unit(str(units[i])) != normalize_display_unit(channel.unit):
            continue
        mapped[channel.semantic_name] = (i, data[:, i])
    evidence = {"file": path.name, "sha256": sha256_file(path), "samples": len(data),
                "top_formulas": top_formulas, "raw_selection_entries": len(entries),
                "source_mappings": {}, "calculated_checks": {}}

    def compare_cached(label, calculated):
        indices = header_columns.get(label_key(label), [])
        if len(indices) != 1:
            return {"cached_comparison": "label not uniquely mapped"}
        cached = data[:, indices[0]]
        mask = np.isfinite(cached) & np.isfinite(calculated)
        return {"cached_column": get_column_letter(indices[0] + 1),
                "compared_samples": int(mask.sum()),
                "max_absolute_difference": float(np.max(np.abs(cached[mask] - calculated[mask]))),
                "first_formula": first_formulas.get(get_column_letter(indices[0] + 1) + "5"),
                "first_calculated": float(calculated[0]), "first_cached": float(cached[0]),
                "last_calculated": float(calculated[-1]), "last_cached": float(cached[-1])}
    with ZipFile(path) as archive:
        evidence["chart_count"] = sum(name.startswith("xl/charts/chart") and name.endswith(".xml") for name in archive.namelist())
    previous = REFERENCES / PREVIOUS[kind]
    if previous.exists():
        old = selection(previous)
        evidence["previous_template"] = previous.name
        evidence["added_source_identifiers"] = sorted(set(entries) - set(old))
        evidence["removed_source_identifiers"] = sorted(set(old) - set(entries))
    else:
        evidence["missing_previous_template"] = str(previous)
    for name, (i, values) in mapped.items():
        if name.startswith(("driveshaft_torque_", "engine_auxiliary_torque_")) or name in {
            "auxiliary_high_voltage_power_consumption", "track_height", "engine_water_temperature"}:
            evidence["source_mappings"][name] = {"source": profile.raw_by_semantic_name()[name].source_name,
                "selection_row": entries.get(profile.raw_by_semantic_name()[name].source_name, {}).get("row"),
                "data_column": get_column_letter(i + 1), "unit": units[i],
                "finite_samples": int(np.isfinite(values).sum()), "min": float(np.nanmin(values)),
                "max": float(np.nanmax(values)), "all_zero": bool(np.all(values == 0))}
    time = data[:, 0]
    for component in ("bulldozing", "compaction"):
        names = [f"tyre_{component}resistancepower_{wheel}" for wheel in ("fl", "fr", "rl", "rr")]
        if all(name in mapped for name in names):
            power = sum(mapped[name][1] for name in names)
            energy = np.cumsum(_interval_energy_kwh(power, time))
            legacy_rectangles = np.cumsum(power * np.r_[time[1] - time[0], np.diff(time)] / 3600)
            evidence["calculated_checks"][component] = {"signed_total_power_min_kw": float(power.min()),
                "signed_total_power_max_kw": float(power.max()), "trapezoidal_final_energy_kwh": float(energy[-1]),
                "cached_power_comparison": compare_cached(profile.math_by_semantic_name()[f"total_{component}_resistance_power"].report_name, power),
                "cached_energy_comparison": compare_cached(profile.math_by_semantic_name()[f"tyre_{component}_resistance_energy_accumulated"].report_name, energy),
                "cached_energy_legacy_rectangle_reconstruction": compare_cached(profile.math_by_semantic_name()[f"tyre_{component}_resistance_energy_accumulated"].report_name, legacy_rectangles)}
    high, low = "auxiliary_high_voltage_power_consumption", "electricsystem_auxiliary_lowvoltagepowerconsumption"
    if high in mapped and low in mapped:
        evidence["calculated_checks"]["auxiliary_final_energy_kwh"] = float(_interval_energy_kwh(mapped[high][1] + mapped[low][1], time).sum())
    if kind != "diesel":
        evidence["wheel_real_world_validation"] = "BLOCKED: measured torque channels are all zero; nonzero synthetic cases are tested separately"
        evidence["stale_rms_divisor_ratio"] = float(np.sqrt(len(data) / 17417))
        evidence["rms_investigation"] = {}
        for semantic, cached_column in (("electricsystem_battery_power", 30), ("electricsystem_battery_heatflow", 32)):
            if semantic in mapped:
                values = mapped[semantic][1]
                expected = float(np.sqrt(np.mean(np.square(values))))
                evidence["rms_investigation"][semantic] = {
                    "valid_samples": int(np.isfinite(values).sum()), "python_sample_rms_kw": expected,
                    "cached_cell": get_column_letter(cached_column + 1) + "2",
                    "cached_rms_kw": rows[1][cached_column],
                    "stale_divisor_reconstruction_kw": float(np.sqrt(np.sum(np.square(values)) / 17417))}
    elif "engine_speed" in mapped:
        rpm = mapped["engine_speed"][1]
        for i in (1, 2):
            name = f"engine_auxiliary_torque_{i}"
            if name in mapped:
                torque = mapped[name][1]
                power = rpm * torque * 2 * np.pi / 60000
                evidence["calculated_checks"][f"auxiliary_{i}"] = {"min_kw": float(power.min()), "max_kw": float(power.max()),
                    "exact_si_comparison": compare_cached(profile.math_by_semantic_name()[f"auxiliary_power_{i}"].report_name, power),
                    "template_divisor_comparison": compare_cached(profile.math_by_semantic_name()[f"auxiliary_power_{i}"].report_name, rpm * torque / 9548.8)}
        total = np.zeros(len(data))
        for wheel in ("fl", "fr", "rl", "rr"):
            if f"wheel_speed_{wheel}" in mapped and f"driveshaft_torque_{wheel}" in mapped:
                power = mapped[f"wheel_speed_{wheel}"][1] * mapped[f"driveshaft_torque_{wheel}"][1] * 2 * np.pi / 60000
                total += power
                evidence["calculated_checks"]["wheel_power_" + wheel] = compare_cached("Wheel Power " + wheel.upper(), power)
        evidence["calculated_checks"]["wheel_power_total"] = compare_cached("Wheel Power Total", total)
    return evidence


def audit_raw(kind):
    path = REFERENCES / RAW[kind]
    if not path.exists():
        return {"missing_file": str(path)}
    name = "full_size_sprayer_diesel" if kind == "diesel" else "robosprayer_" + kind
    profile = load_reporting_profile(PROFILES / (name + ".yaml"))
    data = load_data_file(path)
    result = calculate_profile_math_channels(data, profile)
    sources = {c.source_name for c in data.channels}
    checks = {}
    if kind == "diesel":
        raw = {name: data.values[:, data.channel_index(item.channel.channel_id)] for name, item in result.resolution.resolved.items()}
        for i in (1, 2):
            expected = raw[f"engine_auxiliary_torque_{i}"] * raw["engine_speed"] * 2 * np.pi / 60000
            checks[f"auxiliary_{i}_max_absolute_error_kw"] = float(np.max(np.abs(expected - result.values_by_semantic_name[f"auxiliary_power_{i}"])))
    return {"file": path.name, "sha256": sha256_file(path), "samples": data.quality.sample_count,
            "channels": len(data.channels), "has_new_gearbox_torque": [f"KERS_MotGenUnit_GearboxOutputTorque_{i}" in sources for i in range(1, 5)],
            "has_high_voltage_source": "User_defined_Data_High_Voltage_Auxiliary_signal1_0" in sources,
            "has_track_height": "Track_Height" in sources, "has_legacy_coolant_source": "Engine_Water_Temperature" in sources,
            "unavailable_math": [c.definition.semantic_name for c in result.unavailable],
            "diagnostics": result.resolution.diagnostics, "checks": checks}


def main():
    evidence = {"templates": {kind: audit_workbook(kind) for kind in TEMPLATES},
                "raw_exports": {kind: audit_raw(kind) for kind in RAW}}
    output = ROOT / "outputs/sergio_2026_10_validation/reference_audit.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(evidence, indent=2, allow_nan=False), encoding="utf-8")
    print(output)
    for kind, result in evidence["templates"].items():
        print(kind, result.get("samples"), result.get("chart_count"), result.get("calculated_checks"))


if __name__ == "__main__":
    main()
