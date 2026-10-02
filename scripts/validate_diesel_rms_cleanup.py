"""Generate before/after reports and audit the Diesel RMS-only cleanup.

Run ``--phase baseline`` before editing, then ``--phase final`` afterwards.
Artifacts are kept under outputs/diesel_rms_cleanup by default.
"""
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
import re
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report


ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    "diesel": "Full_Size_Sprayer_8500Kg_Diesel_Test_03_Crop_Field_MatLab_20kph_Gradient.csv",
    "electric": "RoboSprayer_3500Kg_Electric_12kph_Batt_50kW_Mot_63RPM_Susp_Cool_Rough_Grad_Discharge.csv",
    "hybrid": "Sprayer_Caiman_SP_9300Kg_Hybrid_12x1Km_4000Kg_Chem.csv",
}
PROFILES = {"diesel": "full_size_sprayer_diesel", "electric": "robosprayer_electric", "hybrid": "robosprayer_hybrid"}


def package_differences(before, after):
    def stable_content(name, data):
        if name.startswith("xl/worksheets/"):
            # Metadata B3 records generation time. Normalize just that timestamp;
            # every value, formula, style, layout and chart remains comparable.
            return re.sub(rb'(<c r="B3"[^>]* t="inlineStr"><is><t>)\d{4}-\d{2}-\d{2}T[^<]+(</t>)',
                          rb'\1GENERATED_TIMESTAMP\2', data)
        return data

    with ZipFile(before) as old, ZipFile(after) as new:
        assert set(old.namelist()) == set(new.namelist())
        # Creation/modification timestamps are unrelated to report behavior.
        return [name for name in old.namelist()
                if name != "docProps/core.xml" and stable_content(name, old.read(name)) != stable_content(name, new.read(name))]


def audit_diesel(result, baseline):
    book = load_workbook(result.report_path, data_only=False)
    cached = load_workbook(result.report_path, data_only=True)
    old = load_workbook(baseline, data_only=False)
    old_cached = load_workbook(baseline, data_only=True)
    sheet, numeric, previous = book.worksheets[0], cached.worksheets[0], old.worksheets[0]
    try:
        assert sheet.freeze_panes == "B6"
        columns = {c.channel_id: i for i, c in enumerate(result.excel_result.report_channels, 1)}
        count = len(columns)
        end = result.sample_count + 4
        assert [sheet.cell(end + i, 1).value for i in range(1, 5)] == ["MAX", "MIN", "LAST", "FIRST"]
        assert not any(isinstance(cell.value, str) and (
            "Diesel Executive Results" in cell.value or "Complete results: Statistics." in cell.value
            or cell.value == "RMS") for row in sheet for cell in row)
        assert all(sheet.cell(end + 5, col).value is None for col in range(1, count + 2))
        for col in range(1, count + 1):
            for row in range(3, end + 5):
                assert sheet.cell(row, col).value == previous.cell(row, col).value
        for col in range(2, count + 1):
            letter = get_column_letter(col)
            expected = [f"=MAX({letter}5:{letter}{end})", f"=MIN({letter}5:{letter}{end})", f"={letter}{end}", f"={letter}5"]
            assert [sheet.cell(end + i, col).value for i in range(1, 5)] == expected
        metrics = {s.definition.statistic_id: s.value for s in result.excel_result.statistics_result.statistics}
        rms = {}
        for target, label, unit in [("engine_torque", "Engine Torque RMS", "Nm"), ("engine_power", "Engine Power RMS", "kW")]:
            header = next(c for row in sheet.iter_rows(min_row=1, max_row=1) for c in row if c.value == label)
            col = columns[target]
            value = sheet.cell(2, col)
            assert header.column <= col
            letter = get_column_letter(col)
            assert value.value == f"=SQRT(SUMSQ({letter}5:{letter}{end})/COUNT({letter}5:{letter}{end}))"
            series = [numeric.cell(row, col).value for row in range(5, end + 1)]
            expected = math.sqrt(math.fsum(x * x for x in series) / len(series))
            assert math.isclose(expected, metrics[target + "_rms"], rel_tol=2e-15)
            assert math.isclose(numeric[value.coordinate].value, expected, rel_tol=2e-15)
            assert numeric.cell(4, col).value == unit
            rms[target] = {"value": metrics[target + "_rms"], "unit": unit, "cell": value.coordinate, "formula": value.value}
        support = {}
        stats = {row[0]: row[4] for row in cached["Statistics"].iter_rows(values_only=True) if row[0] in metrics}
        before_stats = {row[0]: row[4] for row in old["Statistics"].iter_rows(values_only=True) if row[0] in metrics}
        for corner in ("fl", "fr", "rl", "rr"):
            mid = f"driveshaft_torque_{corner}_rms"
            assert stats[mid] == before_stats[mid]
            assert math.isclose(stats[mid], metrics[mid], rel_tol=2e-15)
            support[mid] = stats[mid]
        removed = {"engine_speed_rms", *(f"chassis_acceleration_{axis}_rms" for axis in ("longitudinal", "lateral", "vertical"))}
        assert not removed & {s.statistic_id for s in result.excel_result.profile.statistics}
        assert not removed & stats.keys()
        assert len(sheet._charts) == len(previous._charts) == 10
        for chart, before in zip(sheet._charts, previous._charts):
            assert chart.to_tree() is not None
            from lxml import etree
            assert etree.tostring(chart._write()) == etree.tostring(before._write())
            assert chart.anchor._from.col == before.anchor._from.col
            assert chart.anchor._from.row == before.anchor._from.row - 3
            assert chart.anchor.ext == before.anchor.ext
            assert chart.x_axis.majorGridlines == before.x_axis.majorGridlines
            assert chart.x_axis.majorGridlines is not None
        # The former banner rows now contain the existing summary and values.
        summary_start = count + 2
        for col in range(summary_start, previous.max_column + 1):
            label = previous.cell(3, col).value
            if label and not label.startswith("Engine Power RMS"):
                found = next(c for c in sheet[1] if c.value == label)
                assert numeric.cell(2, found.column).value == old_cached.worksheets[0].cell(4, col).value
        return {"headline_rms": rms, "supporting_rms": support, "removed": sorted(removed),
                "samples": result.sample_count, "channels": count, "charts": len(sheet._charts),
                "bottom_formulas": 4 * (count - 1), "freeze_panes": sheet.freeze_panes, "status": "PASS"}
    finally:
        book.close()
        cached.close()
        old.close()
        old_cached.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--phase", choices=("baseline", "final"), required=True)
    parser.add_argument("--output-root", type=Path, default=ROOT / "outputs/diesel_rms_cleanup")
    args = parser.parse_args()
    payload = {}
    for name, filename in SOURCES.items():
        print(f"Generating {args.phase} {name}", flush=True)
        result = generate_reporting_profile_engineering_report(
            ROOT / "reference_files" / filename, ROOT / "config/report_profiles" / (PROFILES[name] + ".yaml"),
            args.output_root / args.phase / name)
        entry = {"xlsx": str(result.report_path), "pptx": str(result.presentation_path)}
        if args.phase == "final":
            baseline = args.output_root / "baseline" / name
            if name == "diesel":
                entry.update(audit_diesel(result, baseline / "profile_excel_report" / result.report_path.name))
            else:
                entry["xlsx_changed_parts"] = package_differences(baseline / "profile_excel_report" / result.report_path.name, result.report_path)
                assert not entry["xlsx_changed_parts"], entry
            entry["pptx_changed_parts"] = package_differences(baseline / "profile_powerpoint_report" / result.presentation_path.name, result.presentation_path)
            assert not entry["pptx_changed_parts"], entry
        payload[name] = entry
    (args.output_root / (args.phase + ".json")).write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(json.dumps(payload, indent=2))


if __name__ == "__main__":
    main()
