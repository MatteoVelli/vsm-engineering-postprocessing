"""Regenerate template-driven reports and prepare independent native Excel checks."""
import csv
import json
from pathlib import Path
import re
from zipfile import ZipFile

import numpy as np
from openpyxl import load_workbook

import audit_machine_specifications as reference_audit
from vsm_postprocessing.importer import ImportOptions, load_data_file
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.report_profile import load_reporting_profile
from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'outputs/template_agro_workflow_2026_10_24'
IDS = ('agro_capacity_observed_kg', 'agro_discharge_rate_L_per_m',
       'agro_mass_discharge_rate_kg_per_m', 'agro_discharge_rate_kg_per_km')


def clear_cache(match):
    cell = match.group(0)
    if b'<f' not in cell:
        return cell
    head, body = cell.split(b'>', 1)
    return re.sub(rb'\s+t="[^"]*"', b'', head) + b'>' + re.sub(rb'<v>.*?</v>', b'', body)


def audit_book(case, path):
    book, cached = load_workbook(path), load_workbook(path, data_only=True)
    try:
        rows = {r[0].value:r for r in book['Statistics'] if r[0].value in IDS}
        metrics = {}
        for name in IDS:
            cell = rows[name][2]
            assert cell.data_type == 'f'
            metrics[name] = {'value_cell':cell.coordinate, 'type':'f', 'formula':cell.value,
                             'cached_value':cached['Statistics'][cell.coordinate].value}
        assert 'Agro Mass Source' not in book
        cache_free = path.with_name(path.stem+'_cache_cleared.xlsx')
        with ZipFile(path) as source, ZipFile(cache_free,'w') as target:
            for part in source.infolist():
                content = source.read(part.filename)
                if part.filename.startswith('xl/worksheets/') and part.filename.endswith('.xml'):
                    content = re.sub(rb'<c\b[^>]*>.*?</c>',clear_cache,content)
                target.writestr(part,content)
        return {'case':case, 'workbook':str(path), 'cache_cleared_workbook':str(cache_free), 'metrics':metrics,
                'visible_sheets':[s.title for s in book if s.sheet_state=='visible'],
                'hidden_sheets':[s.title for s in book if s.sheet_state!='visible']}
    finally:
        book.close(); cached.close()


def main():
    reference_audit.OUTPUT = OUTPUT
    reference_audit.main()
    references = json.loads((OUTPUT/'audit.json').read_text())
    previous = {r['case']:r for r in json.loads((ROOT/'outputs/agro_source_policy_2026_10_23/audit.json').read_text())}
    native = []
    legacy_checks = []
    for entry in references:
        assert entry['agro'] == previous[entry['case']]['agro']
        native.append(audit_book(entry['case'],Path(entry['excel'])))
        stats = calculate_profile_statistics(load_data_file(entry['source'],ImportOptions(strict=False)),
                                             load_reporting_profile(entry['profile_file']))
        arrays = dict(stats.math_result.values_by_semantic_name)
        arrays.update({'stat_'+s.definition.statistic_id:np.asarray(s.value) for s in stats.statistics})
        arrays.update({'kpi_'+k.definition.kpi_id:np.asarray(k.value) for k in stats.kpis})
        with np.load(ROOT/'outputs/machine_audit_2026_10_23'/(entry['case']+'_HEAD.npz')) as before:
            assert set(arrays) == set(before.files)
            assert all(np.array_equal(value,before[key],equal_nan=True) for key,value in arrays.items())
        legacy_checks.append({'case':entry['case'], 'legacy_math_arrays':len(stats.math_result.values_by_semantic_name),
            'statistics':len(stats.statistics),'kpis':len(stats.kpis),'comparison':'exactly equal to HEAD baseline'})
    # A derived regression fixture, explicitly not a newly simulated physical run.
    electric = next(r for r in references if r['case']=='electric8500')
    with Path(electric['source']).open(newline='',encoding='utf-8-sig') as stream:
        rows = list(csv.reader(stream))
    force = rows[0].index('HitchRear_Force_Z_VehicleCoordinates')
    for row in rows[2:]:
        row[force] = str(float(row[force])/2)
    fixture = OUTPUT/'derived_fixture/Future_Carrier_Unknown_Cycle.csv'
    fixture.parent.mkdir(parents=True,exist_ok=True)
    with fixture.open('w',newline='',encoding='utf-8') as stream:
        csv.writer(stream).writerows(rows)
    report = generate_reporting_profile_engineering_report(fixture,electric['profile_file'],OUTPUT/'future_machine')
    assert report.report_metadata.configuration_detection == 'unmatched'
    assert report.excel_result.statistics_result.reporting_agro.capacity.value == electric['agro']['capacity']['value']/2
    native.append(audit_book('future_machine',report.report_path))
    (OUTPUT/'native_audit.json').write_text(json.dumps(native,indent=2),encoding='utf-8')
    (OUTPUT/'legacy_engine_comparison.json').write_text(json.dumps(legacy_checks,indent=2),encoding='utf-8')
    (OUTPUT/'future_machine_proof.json').write_text(json.dumps({
        'fixture':'Derived from the Electric 8500 reference, force samples halved; not a new physical simulation',
        'configuration_detection':report.report_metadata.configuration_detection,
        'reporting_inputs':report.report_metadata.reporting_inputs.__dict__,
        'agro':report.excel_result.statistics_result.reporting_agro.to_dict(),
        'excel':str(report.report_path),'presentation':str(report.presentation_path)},indent=2),encoding='utf-8')
    print('All four formulas, source visibility, unchanged numerical engine and unknown-machine load change PASS',flush=True)


if __name__=='__main__':
    main()
