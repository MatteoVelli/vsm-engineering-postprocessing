"""Generate fresh Diesel reports and independently audit the original CSV values.

Run with the project Python, passing a source CSV and a new output directory.
This does not use the profile expression evaluator for the expected values.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
import hashlib
import json
import math
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple
from pptx import Presentation

from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report


def close(actual, expected):
    assert math.isclose(actual, expected, rel_tol=3e-14, abs_tol=1e-12), (actual, expected)


def integral(values, time):
    result = [0.0]
    for i in range(1, len(time)):
        result.append(result[-1] + (values[i - 1] + values[i]) * (time[i] - time[i - 1]) / 2)
    return result


def audit(source: Path, destination: Path):
    with source.open(encoding="cp1252", newline="") as handle:
        names, units, *rows = list(csv.reader(handle))
    duplicates = [name for name, count in Counter(names).items() if count > 1]
    malformed = [i + 3 for i, row in enumerate(rows) if len(row) != len(names)]
    assert not duplicates and not malformed
    numeric = [[float(value) for value in row] for row in rows]
    assert all(math.isfinite(value) for row in numeric for value in row)
    source_columns = {name: list(column) for name, column in zip(names, zip(*numeric))}
    profile_path = Path(__file__).resolve().parents[1] / "config/report_profiles/full_size_sprayer_diesel.yaml"
    result = generate_reporting_profile_engineering_report(source, profile_path, destination)
    excel = result.excel_result
    profile = excel.profile
    expected = {c.semantic_name: source_columns[c.source_name] for c in profile.raw_channels if c.source_name in source_columns}
    time = expected["track_time"]
    duration = time[-1] - time[0]
    power = [torque * rpm * math.tau / 60000 for torque, rpm in zip(expected["engine_torque"], expected["engine_speed"])]
    expected["engine_power"] = power
    expected["engine_mechanical_energy"] = [x / 3600 for x in integral(power, time)]
    if "fuel_flow" in expected:
        expected["fuel_volume"] = [x / 3600 for x in integral(expected["fuel_flow"], time)]
    operations = {"max": max, "min": min, "first": lambda x: x[0], "last": lambda x: x[-1],
                  "rms": lambda x: math.sqrt(math.fsum(v * v for v in x) / len(x))}
    metrics = {s.statistic_id: operations[s.operation](expected[s.target]) for s in profile.statistics if s.target in expected}
    kpis = {
        "drive_cycle_minutes": duration / 60,
        "distance_km": (expected["track_distance"][-1] - expected["track_distance"][0]) / 1000,
        "fuel_consumed_kg": expected["engine_fuel_consumption"][-1] - expected["engine_fuel_consumption"][0],
        "engine_speed_average": integral(expected["engine_speed"], time)[-1] / duration,
    }
    if "fuel_flow" in expected:
        kpis["fuel_flow_average"] = integral(expected["fuel_flow"], time)[-1] / duration
    metrics.update(kpis)
    for metric in excel.statistics_result.canonical_metrics:
        close(metric.value, metrics[metric.metric_id])
    book = load_workbook(result.report_path, data_only=False)
    sheet = book.worksheets[0]
    columns = [c.channel_id for c in excel.report_channels]
    assert len(set(columns)) == len(columns) == len(expected)
    maximum_error = 0.0
    for column, name in enumerate(columns, 1):
        for row, value in enumerate(expected[name], 5):
            stored = sheet.cell(row, column).value
            close(stored, value)
            maximum_error = max(maximum_error, abs(stored - value))
    for row in book["Statistics"].iter_rows(values_only=True):
        if row[0] in metrics:
            stored = row[2] if row[0] in kpis else row[4]
            close(stored, metrics[row[0]])
    for worksheet in book:
        assert not any(cell.data_type in {"e", "f"} for row in worksheet for cell in row)
    assert "Plot Templates" not in book.sheetnames
    assert not sheet._images
    chart_audit = []
    rendered_ids = {p.plot_id for p in excel.plotting_result.rendered_plots}
    plotted_definitions = [p for p in profile.plots if p.plot_id in rendered_ids]
    assert len(sheet._charts) == len(plotted_definitions)
    for chart, definition in zip(sheet._charts, plotted_definitions):
        series = [s for component in chart._charts for s in component.series]
        available = [s for s in definition.series if s.semantic_name in expected]
        assert len(series) == len(available)
        for series_item, expected_series in zip(series, available):
            ref_sheet, (col, first, last_col, last) = range_to_tuple(series_item.yVal.numRef.f)
            assert ref_sheet == sheet.title and columns[col - 1] == expected_series.semantic_name
            assert col == last_col and (first, last) == (5, len(rows) + 4)
            _, (xcol, xfirst, _, xlast) = range_to_tuple(series_item.xVal.numRef.f)
            assert columns[xcol - 1] == definition.x and (xfirst, xlast) == (first, last)
        chart_audit.append({"plot": definition.plot_id, "series": [s.semantic_name for s in available]})
    book.close()
    prs = Presentation(result.presentation_path)
    ppt = result.powerpoint_result.powerpoint_result
    card_count = 0
    for slide, definition in zip(prs.slides, ppt.config.slides):
        texts = [s.text for s in slide.shapes if s.has_text_frame]
        assert texts and definition.title in texts
        assert not any(word in "\n".join(texts).lower() for word in ("battery", "hybrid", "electric", "placeholder"))
        notes = slide.notes_slide.notes_text_frame.text
        note_values = dict(line.split(": ", 1) for line in notes.splitlines() if ": " in line)
        for metric_id in definition.statistic_ids:
            close(float(note_values[metric_id].split()[0]), metrics[metric_id])
            card_count += 1
        for shape in slide.shapes:
            assert 0 <= shape.left <= shape.left + shape.width <= prs.slide_width + 9144
            assert 0 <= shape.top <= shape.top + shape.height <= prs.slide_height + 9144
    for path in (result.report_path, result.presentation_path):
        with ZipFile(path) as package:
            assert package.testzip() is None
    inventory = {
        "source": str(source.resolve()), "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "samples": len(rows), "channels": len(names), "duplicate_headers": duplicates, "malformed_rows": malformed,
        "nonfinite_cells": 0, "time": {"source": "Track_Time", "unit": "s", "first": time[0], "last": time[-1], "elapsed_intervals": len(time) - 1},
        "raw_channels": [{"source": c.source_name, "unit": c.unit, "required": c.required, "present": c.source_name in names,
                          "source_unit": units[names.index(c.source_name)] if c.source_name in names else None} for c in profile.raw_channels],
        "math": [{"name": c.semantic_name, "unit": c.unit, "expression": c.expression, "first": expected[c.semantic_name][0],
                  "last": expected[c.semantic_name][-1]} for c in profile.math_channels if c.semantic_name in expected],
        "statistics_configured": len(profile.statistics), "statistics_available": excel.statistic_count,
        "independent_metrics": metrics, "maximum_excel_absolute_error": maximum_error,
        "checked_channel_samples": len(expected) * len(rows), "plots": chart_audit,
        "slide_count": len(prs.slides), "checked_card_notes": card_count,
        "workbook": str(result.report_path), "presentation": str(result.presentation_path),
        "status": "PASS",
    }
    (destination / "independent_audit.json").write_text(json.dumps(inventory, indent=2), encoding="utf-8")
    print(json.dumps({key: inventory[key] for key in ("status", "samples", "channels", "checked_channel_samples", "maximum_excel_absolute_error", "slide_count", "checked_card_notes")}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    if args.destination.exists():
        parser.error("Use a new audit directory; existing artifacts are preserved.")
    audit(args.source, args.destination)
