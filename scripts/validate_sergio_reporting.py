"""Regenerate latest reports and compare every legacy numeric result with the captured baseline."""
from dataclasses import asdict
import hashlib
import json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from pptx import Presentation
from vsm_postprocessing.importer import ImportOptions
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from vsm_postprocessing.agro_reporting import reporting_metric_rows
from validate_battery_agro_reporting import inspect

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / 'outputs/sergio_reporting_2026_10_22'


def main():
    audits = []
    for baseline in json.loads((DEST / 'baseline.json').read_text()):
        kind = baseline['kind']
        source = Path(baseline['source'])
        assert hashlib.sha256(source.read_bytes()).hexdigest() == baseline['source_sha256']
        print('Generating ' + kind, flush=True)
        excel = generate_profile_excel_report(source, baseline['profile'], DEST / kind / 'excel', ImportOptions(strict=False))
        ppt = build_profile_powerpoint_report(excel, DEST / kind / 'powerpoint')
        assert {k: hashlib.sha256(v.tobytes()).hexdigest() for k, v in excel.math_result.values_by_semantic_name.items()} == baseline['baseline_math']
        assert {x.definition.statistic_id: x.value for x in excel.statistics_result.statistics} == baseline['baseline_statistics']
        assert {x.definition.kpi_id: x.value for x in excel.statistics_result.kpis} == baseline['baseline_kpis']
        details = inspect(excel, ppt, {'electric': 15, 'hybrid': 14, 'diesel': 10}[kind])
        agro = excel.statistics_result.reporting_agro
        assert agro.capacity.value == baseline['max_mass_kg']
        assert agro.capacity.status == 'inferred'
        assert agro.rate.value is None and agro.rate.unit == 'L/m'
        if kind != 'hybrid':
            assert agro.mass_rate.value == baseline['mass_rate_kg_per_m']
            assert 'density required' in agro.rate.reason
        else:
            assert agro.mass_rate.value is None and 'negative' in agro.rate.reason
        inputs = excel.report_metadata.reporting_inputs
        assert inputs.battery_mass_kg == (470 if kind == 'electric' else None)
        metrics = reporting_metric_rows(excel.statistics_result, excel.report_metadata)
        book = load_workbook(excel.report_path, data_only=True)
        rows = {r[0]: r for r in book['Statistics'].iter_rows(values_only=True) if r[0]}
        prs = Presentation(ppt.presentation_path)
        all_text = '\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
        for metric in metrics:
            row = rows[metric['metric_id']]
            expected = metric['display'] if metric['metric_id'] == 'battery_mass_kg' else metric['value'] if metric['value'] is not None else 'N/A'
            if isinstance(expected, (float, int)):
                np.testing.assert_allclose(row[2], expected, rtol=2e-15)
            else:
                assert row[2] == expected
            if metric['metric_id'] != 'battery_mass_kg':
                assert row[3] == metric['unit']
                assert book['Statistics'].cell(next(i for i, r in enumerate(book['Statistics'].iter_rows(values_only=True), 1) if r[0] == metric['metric_id']), 3).number_format == metric['number_format']
            if metric['value'] is not None:
                assert metric['display'] in all_text, metric
        if kind == 'electric':
            assert 'BATTERY PACK MASS: 470 kg (confirmed)' in all_text
        if kind == 'diesel':
            assert '1,700 kg' not in all_text and '470 kg' not in all_text
        book.close()
        audit = {'scenario': kind, 'source': str(source), 'source_sha256': baseline['source_sha256'],
                 'excel': str(excel.report_path), 'powerpoint': str(ppt.presentation_path),
                 'legacy_math_arrays_statistics_kpis': 'PASS', 'expected': {
                     'capacity_kg': baseline['max_mass_kg'],
                     'mass_rate_kg_per_m': baseline['mass_rate_kg_per_m'] if kind != 'hybrid' else None,
                     'volume_rate_L_per_m': None, 'battery_mass_kg': 470 if kind == 'electric' else None},
                 'actual': agro.to_dict(), 'reporting_metrics': metrics,
                 'configuration_detection': excel.report_metadata.configuration_detection,
                 'reporting_inputs': asdict(inputs), 'validation': 'PASS', **details}
        audits.append(audit)
        (DEST / 'audit.json').write_text(json.dumps(audits, indent=2, allow_nan=False))
        print(kind + ' PASS', flush=True)

if __name__ == '__main__':
    main()
