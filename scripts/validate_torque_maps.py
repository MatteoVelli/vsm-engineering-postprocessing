"""Generate and independently inspect torque maps against optional baseline evidence."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re
import sys
from time import perf_counter
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import numpy as np
from openpyxl import load_workbook
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.importer import load_data_file
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from vsm_postprocessing.report_profile import load_reporting_profile


def numeric_record(data, statistics):
    return {"shape": list(data.values.shape), "raw_sha256": hashlib.sha256(data.values.tobytes()).hexdigest(),
            "math": {key: hashlib.sha256(value.tobytes()).hexdigest()
                     for key, value in statistics.math_result.values_by_semantic_name.items()},
            "statistics": {s.definition.statistic_id: s.value for s in statistics.statistics},
            "kpis": {k.definition.kpi_id: k.value for k in statistics.kpis}}


def inspect_excel(report):
    book = load_workbook(report.report_path)
    try:
        sheet = book["Torque Maps"]
        assert len(sheet._charts) == 2 and not sheet._images
        assert sheet.freeze_panes is None and sheet.sheet_view.pane is None
        counts = []
        for chart, key in zip(sheet._charts, ("torque_map_powertrain", "torque_map_wheels")):
            records = [r for r in report.plotting_result.series_diagnostics[key] if r["included"]]
            assert chart.scatterStyle == "marker"
            assert len(chart.series) == (len(records) if records else 1)
            if not records:
                assert chart.series[0].marker.symbol is None
                assert sheet["A32"].value is None and sheet["B32"].value is None
            for series, record in zip(chart.series, records):
                assert series.graphicalProperties.line.noFill and series.marker.symbol == "circle"
                for reference, semantic in ((series.xVal.numRef.f, record["x"]), (series.yVal.numRef.f, record["y"])):
                    name, extent = reference.split("!")
                    assert name == "'Torque Maps'"
                    actual = [row[0].value for row in sheet[extent.replace("$", "")]]
                    # Office XML serialises up to 16 significant decimal digits.
                    np.testing.assert_allclose(actual, report.plotting_result.values_by_semantic_name[semantic], rtol=1e-14, atol=0)
            counts.append(len(records))
        return {"sheets": book.sheetnames, "native_map_charts": 2, "series_counts": counts}
    finally:
        book.close()


def inspect_powerpoint(ppt):
    prs = Presentation(ppt.presentation_path)
    assert [s.slide_id for s in ppt.powerpoint_result.config.slides[-2:]] == ["torque_maps", "simulation_summary"]
    for number, slide in enumerate(prs.slides, 1):
        text = " ".join(" ".join(s.text.split()) for s in slide.shapes if s.has_text_frame)
        assert f"{number} / {len(prs.slides)}" in text
    slide = prs.slides[-2]
    subtitle = "VSM simulation operating points showing the relationship between rotational speed and torque."
    assert any(s.has_text_frame and s.text == subtitle for s in slide.shapes)
    assert all(s.left >= 0 and s.top >= 0 and s.left+s.width <= prs.slide_width+10
               and s.top+s.height <= prs.slide_height+10 for s in slide.shapes)
    plots = [s for s in slide.shapes if s.shape_type == 13 and s.top > 2*914400]
    if len(plots) == 2:
        assert plots[0].left+plots[0].width < plots[1].left
    return {"slides": len(prs.slides), "map_images": len(plots), "map_slide": len(prs.slides)-1}


def compare_existing_reports(root, kind, report, ppt):
    baseline = root / "baseline_reports" / kind
    old_xlsx = next((baseline / "excel").glob("*.xlsx"), None)
    old_pptx = next((baseline / "powerpoint").glob("*.pptx"), None)
    result = {}
    if old_xlsx:
        with ZipFile(old_xlsx) as before, ZipFile(report.report_path) as after:
            assert before.read("xl/worksheets/sheet1.xml") == after.read("xl/worksheets/sheet1.xml")
            charts = [name for name in before.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", name)]
            for name in charts:
                assert before.read(name) == after.read(name), name
            result.update(main_worksheet_xml_unchanged=True, existing_chart_xml_unchanged=len(charts))
            formulas = 0
            ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            for name in before.namelist():
                if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
                    continue
                def expressions(data):
                    return {cell.get("r"): cell.find("m:f", ns).text
                            for cell in ET.fromstring(data).findall(".//m:c", ns) if cell.find("m:f", ns) is not None}
                old, new = expressions(before.read(name)), expressions(after.read(name))
                assert old == new, "Existing formulas changed: "+name
                formulas += len(old)
            result["existing_formulas_unchanged"] = formulas
    if old_pptx:
        before, after = Presentation(old_pptx), Presentation(ppt.presentation_path)
        assert len(after.slides) == len(before.slides)+1
        for index, slide in enumerate(before.slides):
            current = after.slides[index if index < len(before.slides)-1 else -1]
            old_xml, new_xml = ET.fromstring(slide.element.xml), ET.fromstring(current.element.xml)
            for element in (old_xml, new_xml):
                for text in element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"):
                    if text.text and re.fullmatch(r"\s*\d+\s*/\s*\d+\s*", text.text):
                        text.text = "PAGE / TOTAL"
                    elif text.text and re.fullmatch(r"Rendered profile plots: [\d,]+\.", text.text):
                        text.text = "Rendered profile plots: COUNT."
            assert ET.tostring(old_xml) == ET.tostring(new_xml), f"existing {kind} slide {index+1} changed"
        result["existing_slides_unchanged_except_numbering_and_plot_count"] = len(before.slides)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("outputs/torque_maps_2026_10_28"))
    parser.add_argument("--pattern", default="*.csv")
    args = parser.parse_args()
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=True)
    baseline_path = root / "baseline_numerics.json"
    baseline = json.loads(baseline_path.read_text()) if baseline_path.exists() else {}
    result_path = root / "validation.json"
    results = json.loads(result_path.read_text()) if result_path.exists() else {}
    checked = 0
    for source in Path("reference_files").glob(args.pattern):
        start = perf_counter()
        kind = "diesel" if "Diesel" in source.name else "hybrid" if "Hybrid" in source.name else "electric"
        name = "full_size_sprayer_diesel" if kind == "diesel" else "robosprayer_"+kind
        profile_path = Path("config/report_profiles/sergio_2026_10") / (name+".yaml")
        profile = load_reporting_profile(profile_path)
        data = load_data_file(source)
        statistics = calculate_profile_statistics(data, profile)
        record = numeric_record(data, statistics)
        assert statistics.is_complete
        if source.name in baseline:
            assert record == baseline[source.name], "Numerical baseline changed: "+source.name
        # Keep repeated long VSM basenames below the Windows path-length limit.
        label = source.stem[:22]+"_"+hashlib.sha256(source.name.encode()).hexdigest()[:6]
        dest = root / "reports" / label
        report = generate_profile_excel_report(source, profile_path, dest/"excel")
        ppt = build_profile_powerpoint_report(report, dest/"powerpoint")
        result = {"kind": kind, "source": str(source.resolve()), "excel": str(report.report_path),
                  "powerpoint": str(ppt.presentation_path), "numerical_baseline_equal": source.name in baseline,
                  "excel_checks": inspect_excel(report), "powerpoint_checks": inspect_powerpoint(ppt),
                  "series": report.plotting_result.series_diagnostics}
        baseline_match = (kind == "electric" and "Mot_63RPM" in source.name or kind == "hybrid"
                          or kind == "diesel" and "LSD_07" in source.name)
        if baseline_match:
            result["preservation"] = compare_existing_reports(root, kind, report, ppt)
        result["elapsed_seconds"] = round(perf_counter()-start, 3)
        results[source.name] = result
        checked += 1
        (root/"validation.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
        print("PASS", source.name, result["excel_checks"]["series_counts"], result["powerpoint_checks"], flush=True)
    hashes = root/"reference_hashes_before.json"
    if hashes.exists():
        for path, digest in json.loads(hashes.read_text()).items():
            assert hashlib.sha256(Path(path).read_bytes()).hexdigest() == digest, "Reference modified: "+path
    print("Validated", checked, "source datasets; reference integrity passed", flush=True)


if __name__ == "__main__":
    main()
