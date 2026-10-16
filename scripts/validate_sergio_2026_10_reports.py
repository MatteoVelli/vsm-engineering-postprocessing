"""Generate clearly labelled synthetic acceptance reports for Office validation."""
import csv
import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from vsm_postprocessing.report_profile import load_reporting_profile


ROOT = Path(__file__).resolve().parents[1]
DESTINATION = ROOT / "outputs/sergio_2026_10_validation/synthetic_reports"


def main(destination=DESTINATION):
    evidence = []
    for kind in ("diesel", "electric", "hybrid"):
        name = "full_size_sprayer_diesel" if kind == "diesel" else "robosprayer_" + kind
        profile_path = ROOT / "config/report_profiles/sergio_2026_10" / (name + ".yaml")
        profile = load_reporting_profile(profile_path)
        folder = destination / kind
        folder.mkdir(parents=True, exist_ok=True)
        source = folder / (kind + "_SYNTHETIC.csv")
        values = {c.semantic_name: [1., 1., 1., 1.] for c in profile.raw_channels}
        values.update(track_time=[0., 900., 1800., 3600.], track_distance=[0., 100., 200., 300.])
        if kind == "diesel":
            values.update(engine_fuel_consumption=[0., .5, 1., 2.], engine_speed=[1800.] * 4,
                          engine_auxiliary_torque_1=[10.] * 4, engine_auxiliary_torque_2=[20.] * 4,
                          engine_water_temperature=[70., 75., 80., 85.],
                          agrochemical_discharge_force=[-9.80665 * mass for mass in [100, 90, 80, 70]])
        else:
            values.update(electricsystem_battery_energy=[50., 49., 48., 47.],
                          electricsystem_battery_soc=[100., 98., 96., 94.],
                          auxiliary_high_voltage_power_consumption=[8.] * 4,
                          electricsystem_auxiliary_lowvoltagepowerconsumption=[2.] * 4)
        for i, wheel in enumerate(("fl", "fr", "rl", "rr"), 1):
            values["driveshaft_torque_" + wheel] = [v * i for v in [0., 100., -150., 50.]]
            speed_name = ("wheel_speed_" if kind == "diesel" else "wheel_rotationalspeed_") + wheel
            values[speed_name] = [0., 60., 120., -30.]
            for component in ("bulldozing", "compaction"):
                name = f"tyre_{component}resistancepower_{wheel}"
                if name in values:
                    values[name] = [2.5] * 4
        with source.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow([c.source_name for c in profile.raw_channels])
            writer.writerow([c.unit for c in profile.raw_channels])
            writer.writerows(zip(*(values[c.semantic_name] for c in profile.raw_channels)))
        report = generate_profile_excel_report(source, profile_path, folder / "excel", machine_name_override="SYNTHETIC validation")
        ppt = build_profile_powerpoint_report(report, folder / "powerpoint")
        workbook = load_workbook(report.report_path)
        try:
            live = [c.channel_id for col, c in enumerate(report.report_channels, 1)
                    if c.kind == "math" and workbook.worksheets[0].cell(5, col).data_type == "f"]
        finally:
            workbook.close()
        evidence.append({"profile": kind, "workbook": str(report.report_path), "presentation": str(ppt.presentation_path),
                         "samples": report.sample_count, "live_math": live,
                         "equation_inventory": list(report.math_result.calculation_order),
                         "plots": [p.plot_id for p in report.plotting_result.rendered_plots]})
        print(kind, report.report_path, flush=True)
    path = destination / "office_audit.json"
    path.write_text(json.dumps(evidence, indent=2), encoding="utf-8")
    print(path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DESTINATION)
    main(parser.parse_args().output_dir.resolve())
