"""Generate local October 08/04 acceptance reports; never package a release."""
from dataclasses import asdict
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from pptx import Presentation

from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / 'outputs/october_08_04_validation'
CASES = [
    ('electric', 'RoboSprayer_3500Kg_Electric_12kph_Batt_50kW_Mot_63RPM_Susp_Cool_Rough_Grad_Discharge.csv', 470),
    ('hybrid', 'Sprayer_Caiman_SP_9300Kg_Hybrid_12x1Km_4000Kg_Chem.csv', 725),
    ('diesel', 'Full_Size_Sprayer_8500Kg_Diesel_Test_03_Crop_Field_MatLab_20kph_Gradient.csv', None),
]


def main():
    DEST.mkdir(parents=True, exist_ok=True)
    audit = []
    for kind, filename, mass in CASES:
        print('Generating ' + kind, flush=True)
        profile = ROOT / 'config/report_profiles/sergio_2026_10' / (
            'full_size_sprayer_diesel.yaml' if kind == 'diesel' else 'robosprayer_' + kind + '.yaml')
        result = generate_reporting_profile_engineering_report(
            ROOT / 'reference_files' / filename, profile, DEST / kind,
            reporting_inputs={'battery_mass_kg': mass} if mass is not None else {})
        excel = result.excel_result
        net = excel.statistics_result.reporting_agro.net_mass_rate
        resolution = excel.resolution.resolved
        force = excel.dataset.values[:, excel.dataset.channel_index(resolution['agrochemical_discharge_force'].channel.channel_id)]
        distance = excel.dataset.values[:, excel.dataset.channel_index(resolution['track_distance'].channel.channel_id)]
        contents = -force / (9.80665 if kind == 'diesel' else 9.81)
        independent = (contents[0] - contents[-1]) / (distance[-1] / 1000)
        if np.any(contents < 0):
            assert net.value is None and 'negative' in net.reason
        else:
            assert np.isclose(net.value, independent, rtol=1e-13)
        book = load_workbook(result.report_path, data_only=True)
        try:
            rows = {r[0]: r for r in book['Statistics'].values if r[0]}
            assert rows['agro_discharge_rate_kg_per_km'][2] == (net.value if net.value is not None else 'N/A')
            charts = len(book.worksheets[0]._charts)
        finally:
            book.close()
        prs = Presentation(result.presentation_path)
        texts = [[s.text for s in slide.shapes if s.has_text_frame] for slide in prs.slides]
        expected = 'N/A' if net.value is None else f'{net.value:.2f} kg/km'
        assert any('AGRO DISCHARGE RATE' in page and expected in page for page in texts)
        if mass is not None:
            assert 'BATTERY PACK MASS' in texts[4] and f'{mass} kg' in texts[4]
            cards = [s for s in prs.slides[4].shapes if s.shape_type == 1 and s.has_text_frame and not s.text
                     and 1.5*914400 < s.top < 3*914400 and s.width > 914400 and s.height > .5*914400]
            assert len(cards) == 7
            assert all(c.left >= 0 and c.left + c.width <= prs.slide_width for c in cards)
        audit.append({'profile': kind, 'template': excel.profile.metadata.excel_template,
                      'source': filename, 'samples': excel.sample_count, 'excel': str(result.report_path),
                      'presentation': str(result.presentation_path), 'slide_count': len(prs.slides),
                      'chart_count': charts, 'battery_mass_kg': mass, 'rate': asdict(net),
                      'net_mass_trace': excel.statistics_result.reporting_agro.trace['net_mass_discharge']})
        (DEST / 'audit.json').write_text(json.dumps(audit, indent=2), encoding='utf-8')
        print(kind + ': PASS; rate=' + expected, flush=True)


if __name__ == '__main__':
    main()
