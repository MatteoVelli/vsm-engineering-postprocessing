"""Audit real-data tyre correction against a baseline captured before editing.

Run --baseline first, then without it. Sources come from the existing local
client audit (or --sources, a JSON list of {kind, source} entries). Artifacts
remain in ignored outputs; no source file or reference template is modified.
"""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple
from pptx import Presentation

from vsm_postprocessing.excel_report_engine import generate_profile_excel_report, _ordered_profile_plots
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from vsm_postprocessing.profile_statistics import build_canonical_report_metrics
from vsm_postprocessing.profile_soil_kpis import build_soil_powerpoint_metrics

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "outputs/tyre_rolling_correction_2026_10_18"
RETIRED = {"tyre_total_resistance_" + s for s in
           ("power", "energy_kwh", "energy_wh", "energy_accumulated")}
POWER = "total_rolling_resistance_power"
ENERGY = "tyre_rolling_resistance_energy_accumulated"


def digest(data):
    return hashlib.sha256(data).hexdigest()


def check_recalculated_errors(path):
    """Check error cells on all recalculated sheets, including hidden helpers."""
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                continue
            with archive.open(name) as stream:
                for _, element in ET.iterparse(stream, events=("end",)):
                    if element.tag == "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c":
                        assert element.get("t") != "e", (str(path), name, element.get("r"))
                        element.clear()


def capture(kind, source, phase):
    name = "full_size_sprayer_diesel" if kind == "diesel" else "robosprayer_" + kind
    folder = DEST / phase / kind
    report = generate_profile_excel_report(source, ROOT / "config/report_profiles/sergio_2026_10" / (name + ".yaml"), folder / "excel")
    ppt = build_profile_powerpoint_report(report, folder / "powerpoint")
    values = report.plotting_result.values_by_semantic_name
    statistics = {s.definition.statistic_id: s.value for s in report.statistics_result.statistics}
    kpis = {k.definition.kpi_id: k.value for k in report.statistics_result.kpis}
    prs = Presentation(ppt.presentation_path)
    canonical = build_canonical_report_metrics(report.statistics_result)
    data = dict(kind=kind, profile=name, source=str(source), source_sha256=digest(Path(source).read_bytes()),
                workbook=str(report.report_path), presentation=str(ppt.presentation_path), samples=report.sample_count,
                arrays={k: digest(v.tobytes()) for k, v in values.items()}, statistics=statistics, kpis=kpis,
                tyre_energy_kwh=float(values.get("tyre_total_resistance_energy_accumulated", values[ENERGY])[-1]),
                peak_tyre_power_kw=float(max(values.get("tyre_total_resistance_power", values[POWER]))),
                channels=[c.channel_id for c in report.report_channels],
                plots={p.plot_id: digest(Path(p.output_file).read_bytes()) for p in report.plotting_result.rendered_plots},
                slides=[s.slide_id for s in ppt.powerpoint_result.config.slides],
                slide_geometry=[[[s.shape_type, s.left, s.top, s.width, s.height] for s in slide.shapes] for slide in prs.slides],
                presentation_metrics={m.metric_id: dict(value=m.value, source=m.source_channel, unit=m.unit)
                                      for m in (*canonical, *build_soil_powerpoint_metrics(report, canonical))})
    return report, ppt, data


def validate(report, ppt, old, new):
    assert old["source_sha256"] == new["source_sha256"]
    assert set(old["arrays"]) - set(new["arrays"]) == (RETIRED if new["kind"] != "diesel" else set())
    assert all(old["arrays"][k] == v for k, v in new["arrays"].items())
    affected = {"tyre_total_energy_accumulated_last"}
    assert {k: v for k, v in old["statistics"].items() if k not in affected} == {
        k: v for k, v in new["statistics"].items() if k not in affected}
    assert {k: v for k, v in old["kpis"].items() if k != "total_energy_consumption_tyres_aux"} == {
        k: v for k, v in new["kpis"].items() if k != "total_energy_consumption_tyres_aux"}
    assert not RETIRED.intersection(new["channels"])
    assert len(new["channels"]) == len(set(new["channels"]))
    assert old["slides"] == new["slides"] and old["slide_geometry"] == new["slide_geometry"]
    changed_plots = {"tyres_energy_consumption", "tyre_total_resistance_power", "tyre_total_resistance_energy_accumulated"}
    assert old["plots"].keys() == new["plots"].keys()
    assert all(old["plots"][k] == v for k, v in new["plots"].items() if k not in changed_plots)
    affected_metrics = affected | {"total_energy_consumption_tyres_aux"} | {
        "tyre_total_" + suffix for suffix in ("max_power", "mean_power", "rms_power", "total_energy", "energy_per_km")}
    assert all(old["presentation_metrics"][k] == v for k, v in new["presentation_metrics"].items()
               if k not in affected_metrics)
    deck_metrics = {s.statistic_id: s for s in ppt.powerpoint_result.statistics_result.statistics}
    for name, metric in new["presentation_metrics"].items():
        assert np.isclose(deck_metrics[name].value, metric["value"])
    values = report.plotting_result.values_by_semantic_name
    if new["kind"] != "diesel":
        assert new["statistics"]["tyre_total_energy_accumulated_last"] == values[ENERGY][-1]
        assert ("total_energy_consumption_tyres_aux" in new["kpis"]) == ("total_energy_consumption_tyres_aux" in old["kpis"])
        if "total_energy_consumption_tyres_aux" in new["kpis"]:
            assert np.isclose(new["kpis"]["total_energy_consumption_tyres_aux"],
                              values[ENERGY][-1] + new["statistics"]["auxiliary_energy_accumulated_max"])
        assert [s.semantic_name for s in report.profile.plots_by_id()["tyres_energy_consumption"].series] == [POWER, ENERGY]
        time = values["track_time"]
        independent = np.r_[0, np.cumsum((values[POWER][:-1] / 2 + values[POWER][1:] / 2) * np.diff(time) / 3600)]
        np.testing.assert_allclose(values[ENERGY], independent, rtol=3e-14, atol=1e-12)
        if new["kind"] == "electric":
            expected = dict(max_power=max(values[POWER]),
                            mean_power=np.trapezoid(values[POWER], time) / (time[-1] - time[0]),
                            rms_power=np.sqrt(np.mean(values[POWER] ** 2)),
                            total_energy=values[ENERGY][-1] - values[ENERGY][0],
                            energy_per_km=(values[ENERGY][-1] - values[ENERGY][0]) / new["statistics"]["distance_km_last"])
            for suffix, value in expected.items():
                assert np.isclose(deck_metrics["tyre_total_" + suffix].value, value, rtol=3e-14, atol=1e-12)
    book = load_workbook(report.report_path)
    cached = load_workbook(report.report_path, data_only=True)
    previous = load_workbook(old["workbook"])
    formula_count = references = 0
    try:
        sheet = book.worksheets[0]
        columns = {c.channel_id: i for i, c in enumerate(report.report_channels, 1)}
        for channel, col in columns.items():
            np.testing.assert_allclose([cached.worksheets[0].cell(r, col).value for r in range(5, report.sample_count + 5)],
                                       values[channel], rtol=3e-14, atol=1e-12)
        assert len(sheet._charts) == len(previous.worksheets[0]._charts) == report.plot_count
        for chart, before, plot in zip(sheet._charts, previous.worksheets[0]._charts, _ordered_profile_plots(report.plotting_result.rendered_plots)):
            # The panel follows the exported data. Removing four columns shifts
            # it left by four columns while preserving its alignment and size.
            anchor = deepcopy(before.anchor)
            anchor._from.col -= len(old["channels"]) - len(new["channels"])
            assert chart.anchor == anchor, plot.plot_id
            assert chart.style == before.style
            for sub, old_sub in zip(chart._charts, before._charts):
                assert len(sub.ser) == len(old_sub.ser)
                for series, old_series in zip(sub.ser, old_sub.ser):
                    assert series.graphicalProperties == old_series.graphicalProperties
                    for axis in (series.xVal, series.yVal):
                        if axis and axis.numRef:
                            title, (c1, r1, c2, r2) = range_to_tuple(axis.numRef.f)
                            assert title == sheet.title and 1 <= c1 == c2 <= len(columns)
                            assert r1 == 5 and r2 == report.sample_count + 4
                            references += 1
            if plot.plot_id == "tyres_energy_consumption":
                actual = [s.yVal.numRef.f for sub in chart._charts for s in sub.ser]
                assert [range_to_tuple(ref)[1][0] for ref in actual] == [columns[POWER], columns[ENERGY]]
        for ws in book:
            for row in ws:
                for cell in row:
                    assert cell.data_type != "e" and "#REF!" not in str(cell.value)
                    formula_count += cell.data_type == "f"
        stats_sheet = book["Statistics"]
        for row in stats_sheet.iter_rows():
            if row[0].value == "tyre_total_energy_accumulated_last":
                assert row[4].data_type == "f"
                from openpyxl.utils import get_column_letter
                assert f"{get_column_letter(columns[ENERGY])}5:" in row[4].value.replace("$", "")
        for item in report.statistics_result.kpis:
            row = next(r for r in stats_sheet.iter_rows() if r[0].value == item.definition.kpi_id)
            assert row[2].data_type == "f"
            assert cached["Statistics"][row[2].coordinate].value == item.value
    finally:
        book.close(); cached.close(); previous.close()
    new.update(status="PASS", formula_cells=formula_count, chart_references=references,
               unchanged_arrays=len(new["arrays"]), unchanged_plot_images=len(new["plots"].keys() - changed_plots))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", action="store_true")
    parser.add_argument("--kind", choices=("electric", "hybrid", "diesel"))
    parser.add_argument("--sources", type=Path, default=ROOT / "outputs/sergio_final_2026_10_17/baseline.json")
    parser.add_argument("--prepare-recalculation", action="store_true")
    parser.add_argument("--audit-recalculation", action="store_true")
    args = parser.parse_args()
    DEST.mkdir(parents=True, exist_ok=True)
    audit_path = DEST / "audit.json"
    if args.prepare_recalculation:
        from validate_october_reporting import prepare_recalculation
        entries = json.loads(audit_path.read_text())
        for entry in entries:
            entry.setdefault("profile", entry["kind"])
        audit_path.write_text(json.dumps(entries, indent=2))
        prepare_recalculation(audit_path)
        return
    if args.audit_recalculation:
        entries = json.loads(audit_path.read_text())
        for entry in entries:
            path = DEST / "recalculated" / Path(entry["workbook"]).name
            expected = load_workbook(entry["workbook"], data_only=True)
            actual = load_workbook(path, data_only=True)
            checked = 0
            maximum = 0.
            try:
                for name in (expected.worksheets[0].title, "Statistics"):
                    for row in expected[name]:
                        for cell in row:
                            value = actual[name][cell.coordinate]
                            assert value.data_type != "e", (entry["kind"], name, cell.coordinate, value.value)
                            if not isinstance(cell.value, (int, float)):
                                continue
                            assert isinstance(value.value, (int, float)), (name, cell.coordinate, value.value)
                            assert np.isclose(value.value, cell.value, rtol=3e-14, atol=1e-12), (entry["kind"], name, cell.coordinate, value.value, cell.value)
                            maximum = max(maximum, abs(value.value - cell.value))
                            checked += 1
            finally:
                expected.close(); actual.close()
            check_recalculated_errors(path)
            entry.update(recalculation="PASS: desktop Excel, empty input formula caches",
                         recalculated_workbook=str(path), recalculated_numeric_cells=checked,
                         maximum_recalculation_difference=maximum, recalculated_error_cells=0)
            print(f"{entry['kind']}: {checked} independently recalculated numeric cells PASS; maximum difference {maximum:.12g}", flush=True)
        audit_path.write_text(json.dumps(entries, indent=2))
        return
    baseline = DEST / "baseline.json"
    if args.baseline and baseline.exists():
        raise FileExistsError("Keep the pre-change baseline: " + str(baseline))
    if args.baseline and args.kind:
        parser.error("Capture all baseline cases together; --kind is for targeted revalidation.")
    cases = json.loads(args.sources.read_text()) if args.baseline else json.loads(baseline.read_text())
    if args.kind:
        cases = [case for case in cases if case["kind"] == args.kind]
    results = ([entry for entry in json.loads(audit_path.read_text()) if entry["kind"] != args.kind]
               if args.kind and audit_path.exists() else [])
    for case in cases:
        print(("Baseline " if args.baseline else "Corrected ") + case["kind"], flush=True)
        report, ppt, data = capture(case["kind"], case["source"], "before" if args.baseline else "after")
        if not args.baseline:
            validate(report, ppt, case, data)
        results.append(data)
        if not args.baseline:
            audit_path.write_text(json.dumps(results, indent=2))
        print(f"{case['kind']}: {data['tyre_energy_kwh']:.12f} kWh; peak {data['peak_tyre_power_kw']:.12f} kW", flush=True)
    (baseline if args.baseline else DEST / "audit.json").write_text(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
