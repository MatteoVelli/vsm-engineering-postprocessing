"""Statistics Value cells must be live formulas with deterministic reader caches."""
from dataclasses import replace
import json

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.formula import Tokenizer
import pytest

from test_profile_math import _channel
from test_sergio_october_15 import dataset_for, profile_for, profile_path, csv_for, remove_channel
from vsm_postprocessing.agro_excel_formulas import write_agro_statistics_formulas
from vsm_postprocessing.agro_reporting import calculate_reporting_kpis, reporting_metric_rows
from vsm_postprocessing.excel_formulas import save_report_workbook
from vsm_postprocessing.excel_report_engine import _bind_agro_excel_sources, generate_profile_excel_report
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.report_metadata import resolve_report_metadata
from vsm_postprocessing.report_selection import ProfileReportSelection


IDS = ('agro_capacity_observed_kg', 'agro_discharge_rate_L_per_m',
       'agro_mass_discharge_rate_kg_per_m', 'agro_discharge_rate_kg_per_km')


def build_book(data, profile, inputs=None):
    # Isolate reporting formulas from unrelated range/energy KPI requirements.
    profile = replace(profile, kpis=())
    metadata = resolve_report_metadata("Unseen machine's run.csv", profile, reporting_inputs=inputs)
    result = calculate_profile_statistics(data, profile)
    result.reporting_agro = calculate_reporting_kpis(result, metadata)
    book = Workbook()
    book.active.title = "Unseen machine's data"
    book.create_sheet('Statistics')
    book.create_sheet('Metadata').sheet_state = 'hidden'
    book.create_sheet('Formula Support').sheet_state = 'hidden'
    _bind_agro_excel_sources(book, result)
    write_agro_statistics_formulas(book, result, metadata)
    start = len(result.statistics) + len(result.kpis) + 8
    cells = {item['metric_id']: book['Statistics'].cell(row, 3)
             for row, item in enumerate(reporting_metric_rows(result, metadata), start)
             if item['metric_id'] in IDS}
    return book, result.reporting_agro, cells


@pytest.mark.parametrize('kind', ('electric', 'hybrid', 'diesel'))
@pytest.mark.parametrize('count', (2, 7, 31))
@pytest.mark.parametrize('source', ('direct_mass', 'force_inference'))
def test_all_three_values_have_formulas_and_matching_caches(tmp_path, kind, count, source):
    data = dataset_for(kind, np.arange(count))
    masses = np.linspace(713, 83, count)
    settings = {'density_kg_per_L': 1.2, 'density_provenance': 'Explicit synthetic fluid fixture'}
    data.values[:, data.channel_index('track_distance')] = np.linspace(17, 317, count)
    if source == 'direct_mass':
        data.values[:, data.channel_index('agrochemical_discharge_force')] = 0
        data.channels.append(_channel('chemical_mass', 'Tank contents', 'kg', len(data.channels)+1))
        data.values = np.column_stack((data.values, masses))
        settings.update(agro_mass_source_policy=source, agro_mass_source='Tank contents',
                        agro_mass_source_provenance='Explicit simulation mass fixture')
    else:
        data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81*masses
    # Exercise kilometer conversion and nonzero counter origins as well.
    profile = profile_for(kind)
    if count == 31:
        channel = data.channels[data.channel_index('track_distance')]
        data.channels[data.channel_index('track_distance')] = replace(channel, unit='km')
        data.values[:, data.channel_index('track_distance')] /= 1000
        profile = replace(profile, raw_channels=tuple(replace(c, unit='km') if c.semantic_name == 'track_distance'
                                                     else c for c in profile.raw_channels))
    book, agro, cells = build_book(data, profile, settings)
    path = tmp_path / 'formulas.xlsx'
    save_report_workbook(book, path)
    formulas, cached = load_workbook(path), load_workbook(path, data_only=True)
    try:
        expected = (agro.capacity.value, agro.rate.value, agro.mass_rate.value, agro.net_mass_rate.value)
        for metric_id, value in zip(IDS, expected):
            cell = formulas['Statistics'][cells[metric_id].coordinate]
            assert cell.data_type == 'f' and cell.value.startswith('=IFERROR(')
            assert Tokenizer(cell.value).items
            assert cached['Statistics'][cell.coordinate].value == pytest.approx(value, rel=1e-13)
            assert 'MAX(' in cell.value if metric_id == IDS[0] else 'SUMPRODUCT(' in cell.value
        assert 'Agro Mass Source' not in formulas
        np.testing.assert_allclose([cached['Metadata'][f'E{r}'].value for r in range(5,count+5)], masses)
        assert formulas['Metadata']['E5'].value == ("='Metadata'!D5" if source == 'direct_mass' else "=-'Metadata'!D5/9.81")
        assert all(formulas['Metadata'].column_dimensions[c].hidden for c in ('D','E','F'))
        assert f'E{count+4}' in cells[IDS[0]].value
        assert 'MEDIAN(' in cells[IDS[2]].value and 'COUNT(' in cells[IDS[2]].value
        density = next(row[1] for row in formulas['Formula Support'] if str(row[0].value).startswith('Agro density'))
        assert density.value == 1.2 and density.coordinate in cells[IDS[1]].value.replace('$', '')
        assert agro.capacity.value == pytest.approx(713)
        assert agro.mass_rate.value == pytest.approx(630/300)
        assert agro.rate.value == pytest.approx(630/(1.2*300))
        assert formulas.calculation.calcMode == 'auto'
        assert formulas.calculation.fullCalcOnLoad and formulas.calculation.forceFullCalc
    finally:
        formulas.close(); cached.close()


@pytest.mark.parametrize('problem', ('no_density', 'zero_distance', 'reversed_distance',
    'negative_distance', 'distance_jump', 'refill', 'mass_jump', 'negative_contents',
    'missing_force', 'zero_force', 'nan_force', 'wrong_unit', 'missing_distance'))
def test_unavailable_values_are_formulas_with_text_na_caches(tmp_path, problem):
    data = dataset_for('electric', np.arange(8))
    mass = np.linspace(100, 30, 8)
    settings = {'density_kg_per_L': 1.2, 'density_provenance': 'Synthetic fixture'}
    distance = data.values[:, data.channel_index('track_distance')]
    if problem == 'no_density':
        settings = {}
    elif problem == 'zero_distance':
        distance[:] = 7
    elif problem == 'reversed_distance':
        distance[3] = -1 + distance[2]
    elif problem == 'negative_distance':
        distance[:] -= 10
    elif problem == 'distance_jump':
        distance[:] = [0, 1, 2, 3, 4, 5, 1000, 1001]
    elif problem == 'refill':
        mass[3] = mass[2] + 3
    elif problem == 'mass_jump':
        mass[:] = [100, 99, 98, 97, 96, 95, 31, 30]
    elif problem == 'negative_contents':
        mass[-1] = -1
    elif problem == 'zero_force':
        mass[:] = 0
    elif problem == 'nan_force':
        mass[3] = np.nan
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81*mass
    if problem == 'missing_force':
        remove_channel(data, 'agrochemical_discharge_force')
    elif problem == 'missing_distance':
        remove_channel(data, 'track_distance')
    elif problem == 'wrong_unit':
        index = data.channel_index('agrochemical_discharge_force')
        data.channels[index] = replace(data.channels[index], unit='kg')
    book, agro, cells = build_book(data, profile_for('electric'), settings)
    path = tmp_path/'unavailable.xlsx'
    save_report_workbook(book, path)
    cached = load_workbook(path, data_only=True)
    try:
        for metric, metric_id in zip((agro.capacity, agro.rate, agro.mass_rate, agro.net_mass_rate), IDS):
            assert cells[metric_id].data_type == 'f'
            value = cached['Statistics'][cells[metric_id].coordinate].value
            assert value == ('N/A' if metric.value is None else pytest.approx(metric.value))
        assert agro.rate.value is None
        assert agro.mass_rate.value is None if problem != 'no_density' else agro.mass_rate.value is not None
    finally:
        cached.close()


@pytest.mark.parametrize('direct_mode', ('volume', 'flow', 'invalid_volume'))
def test_existing_volumetric_modes_keep_their_actual_source(direct_mode, tmp_path):
    data = dataset_for('electric', np.arange(7))
    data.values[:, data.channel_index('agrochemical_discharge_force')] = 0
    settings = {'agro_volume_source' if direct_mode != 'flow' else 'agro_flow_source': 'Spray source',
                'agro_volume_provenance' if direct_mode != 'flow' else 'agro_flow_provenance': 'Explicit synthetic source'}
    if direct_mode != 'invalid_volume':
        data.channels.append(_channel('spray', 'Spray source', 'L' if direct_mode == 'volume' else 'L/min', len(data.channels)+1))
        data.values = np.column_stack((data.values, np.linspace(0, 60, 7) if direct_mode == 'volume' else np.full(7, 600)))
    book, agro, cells = build_book(data, profile_for('electric'), settings)
    path = tmp_path/'volume.xlsx'
    save_report_workbook(book, path)
    cached = load_workbook(path, data_only=True)
    try:
        cell = cells[IDS[1]]
        assert cell.data_type == 'f'
        assert cached['Statistics'][cell.coordinate].value == ('N/A' if direct_mode == 'invalid_volume' else pytest.approx(.2))
        assert 'Agro Mass Source' not in cell.value
        assert agro.capacity.value is None
    finally:
        cached.close()


@pytest.mark.parametrize('kind', ('electric', 'hybrid', 'diesel'))
def test_generated_workbook_preserves_layout_metadata_and_value_formulas(tmp_path, kind):
    data = dataset_for(kind, np.arange(7))
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81*np.linspace(871, 57, 7)
    result = generate_profile_excel_report(csv_for(data, tmp_path/'Future_Carrier.csv'), profile_path(kind),
        tmp_path/'report', reporting_inputs={'density_kg_per_L': 1.2, 'density_provenance': 'Synthetic fluid'},
        selection=ProfileReportSelection(export_channels=('track_time',), plots=()))
    formulas, cached = load_workbook(result.report_path), load_workbook(result.report_path, data_only=True)
    try:
        rows = {row[0].value: row for row in formulas['Statistics'] if row[0].value in IDS}
        expected_formats = {row['metric_id']: row.get('number_format', 'General')
                            for row in reporting_metric_rows(result.statistics_result, result.report_metadata)}
        for metric in (result.statistics_result.reporting_agro.capacity, result.statistics_result.reporting_agro.rate,
                       result.statistics_result.reporting_agro.mass_rate, result.statistics_result.reporting_agro.net_mass_rate):
            cell = rows[metric.metric_id][2]
            assert cell.data_type == 'f' and cell.column == 3
            assert cached['Statistics'][cell.coordinate].value == pytest.approx(metric.value)
            assert cell.number_format == expected_formats[metric.metric_id]
            assert cell.border.left.style == 'thin'
        assert formulas['Formula Support'].sheet_state == 'hidden'
        metadata = {row[0]: row[1] for row in cached['Metadata'].values}
        assert json.loads(metadata['Agro reporting diagnostics']) == result.statistics_result.reporting_agro.to_dict()
        assert len(result.report_channels) == 1  # Selected mass is not in the visible raw export.
        assert 'Agro Mass Source' not in formulas
        assert result.workbook_structure['agro_source']['mass']['sheet'] == 'Metadata'
    finally:
        formulas.close(); cached.close()
