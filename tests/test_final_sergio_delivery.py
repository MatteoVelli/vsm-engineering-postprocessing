"""Final wording and preservation of the approved source/template conventions."""
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from pptx import Presentation
import pytest

from test_sergio_october_15 import csv_for, dataset_for, profile_path
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from vsm_postprocessing.report_profile import load_reporting_profile


LABELS = {
    "suspension_damper_force_fr": ("Suspension_Damper_Force_FR", "Suspension Damper Force FL"),
    "suspension_damper_force_rl": ("Suspension_Damper_Force_RL", "Suspension Damper Force FR"),
    "suspension_damper_force_rr": ("Suspension_Damper_Force_RR", "Suspension Damper Force RL"),
    "suspension_damper_travel_fr": ("Suspension_Damper_Travel_FR", "Suspension Damper Travel FL"),
    "tyre_rollingresistancepower_rr": ("Tyre_RollingResistancePower_RR", "Tyre Rolling Resistance Power FR"),
    "softsoil_sheardisplacementy_in_rr": ("SoftSoil_ShearDisplacementY_IN_RR", "SoftSoil Shear Displacement Y IN FL"),
}


@pytest.mark.parametrize("path", [
    Path("config/report_profiles") / ("robosprayer_" + kind + ".yaml")
    for kind in ("electric", "hybrid")
] + [profile_path(kind) for kind in ("electric", "hybrid")])
def test_original_sources_and_validated_template_display_labels_are_preserved(path):
    profile = load_reporting_profile(path)
    for name, (source, label) in LABELS.items():
        definition = profile.raw_by_semantic_name()[name]
        assert definition.source_name == source and definition.report_name == label


@pytest.fixture(scope="module", params=("electric", "hybrid"))
def labelled_report(request, tmp_path_factory):
    kind = request.param
    folder = tmp_path_factory.mktemp("labels_" + kind)
    data = dataset_for(kind)
    sentinels = {}
    for i, name in enumerate(LABELS, 1):
        sentinels[name] = i * np.array([1., 2., -3., 4.])
        data.values[:, data.channel_index(name)] = sentinels[name]
    report = generate_profile_excel_report(csv_for(data, folder / "source.csv"), profile_path(kind), folder / "excel")
    return kind, report, sentinels, folder


def test_actual_excel_headers_and_mapping_identify_original_source_with_unchanged_values(labelled_report):
    _, report, sentinels, _ = labelled_report
    book = load_workbook(report.report_path, data_only=True)
    try:
        sheet = book.worksheets[0]
        mapping = book["Rename From VSM to Astauto"]
        for name, (source, label) in LABELS.items():
            col = next(i for i, c in enumerate(report.report_channels, 1) if c.channel_id == name)
            assert sheet.cell(3, col).value == label
            row = next(r for r in mapping.iter_rows(min_row=3) if r[5].value == name)
            assert row[0].value == source and row[1].value == label
            np.testing.assert_array_equal([sheet.cell(r, col).value for r in range(5, 9)], sentinels[name])
            np.testing.assert_array_equal(report.plotting_result.values_by_semantic_name[name], sentinels[name])
    finally:
        book.close()


def test_rolling_tyre_detail_page_keeps_six_cards_and_two_plots(labelled_report):
    kind, report, _, folder = labelled_report
    if kind == "hybrid":
        assert "tyre_total" not in report.profile.presentation.slides_by_id()
        return
    ppt = build_profile_powerpoint_report(report, folder / "powerpoint")
    ids = [s.slide_id for s in ppt.powerpoint_result.config.slides]
    index = ids.index("tyre_total")
    assert ids[index:index+3] == ["tyre_total", "soil_bulldozing", "soil_compaction"]
    assert ppt.slide_count == 16
    assert ppt.powerpoint_result.config.slides[-2].slide_id == "torque_maps"
    slide = Presentation(ppt.presentation_path).slides[index]
    text = " ".join(s.text for s in slide.shapes if s.has_text_frame)
    assert "Tyre Rolling Resistance Power & Energy" in text
    assert "includes bulldozing and compaction" in text
    definition = report.profile.plots_by_id()["tyre_total_resistance_energy_accumulated"]
    assert definition.series[0].label == "Accumulated Tyre Resistance Energy"
    assert len(ppt.powerpoint_result.config.slides[index].statistic_ids) == 6
    assert len([s for s in slide.shapes if s.shape_type == 13 and s.top > 3*914400]) == 2
    for name in ("soil_bulldozing", "soil_compaction"):
        assert report.profile.presentation.slides_by_id()[name].subtitle == "Total resistance power [kW] and accumulated energy [kWh] versus time [s]"
