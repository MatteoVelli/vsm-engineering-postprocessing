"""Independent XY pairing, source policy and Office integration acceptance tests."""
from dataclasses import replace
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
import matplotlib.axes
import numpy as np
from openpyxl import load_workbook
from pptx import Presentation
import pytest

from vsm_postprocessing.errors import ConfigurationError
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.importer import ImportOptions
from vsm_postprocessing.profile_math import calculate_profile_math_channels
from vsm_postprocessing.profile_plotting import (_build_semantic_maps, _paired_series_status,
    _prepare_plot_definition, UnavailableProfilePlot)
from vsm_postprocessing.plotting_engine import _render_one_plot, PlotDefaults
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from vsm_postprocessing.report_profile import load_reporting_profile, resolve_profile
from vsm_postprocessing.report_selection import ProfileReportSelection, select_profile
from test_sergio_october_15 import dataset_for, csv_for, profile_for, profile_path

NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
MAPS = ("torque_map_powertrain", "torque_map_wheels")


def maps_for(kind="electric"):
    profile, data = profile_for(kind), dataset_for(kind)
    resolution = resolve_profile(data, profile)
    math = calculate_profile_math_channels(data, profile, resolution)
    channels, values = _build_semantic_maps(data, resolution, math)
    return profile, data, channels, values


def test_independent_signed_pairs_reach_matplotlib_without_lines(tmp_path, monkeypatch):
    profile, data, channels, values = maps_for()
    definition = profile.plots_by_id()[MAPS[0]]
    for i in range(1, 5):
        values[f"electricsystem_em{i}_speed"] = np.array([-i, 0, 10*i, 20*i], dtype=float)
        values[f"electricsystem_em{i}_torque"] = np.array([i, -5*i, 0, 100*i], dtype=float)
    native, summaries = _prepare_plot_definition(definition, channels, values, 4)
    calls = []
    original = matplotlib.axes.Axes.plot
    def spy(self, x, y, **kwargs):
        if kwargs.get("label", "").startswith("EM"):
            calls.append((np.array(x), np.array(y), kwargs))
        return original(self, x, y, **kwargs)
    monkeypatch.setattr(matplotlib.axes.Axes, "plot", spy)
    rendered = _render_one_plot(native, PlotDefaults(), tmp_path, channels, values, 4)
    assert len(calls) == 4 and len(summaries) == 4
    for i, (x, y, style) in enumerate(calls, 1):
        np.testing.assert_array_equal(x, values[f"electricsystem_em{i}_speed"])
        np.testing.assert_array_equal(y, values[f"electricsystem_em{i}_torque"])
        assert style["linestyle"] == "None" and style["marker"] == "o"
    assert rendered.series_x_channel_ids == tuple(f"electricsystem_em{i}_speed" for i in range(1, 5))
    assert rendered.render_mode == "scatter" and Path(rendered.png_file).is_file()


@pytest.mark.parametrize("case,status", [("missing_x", "unavailable"), ("missing_y", "unavailable"),
    ("nan", "unavailable"), ("inf", "unavailable"), ("wrong_rpm_unit", "unavailable"),
    ("wrong_torque_unit", "unavailable"), ("counter_torque", "unavailable"),
    ("misaligned", "unavailable"), ("inactive", "inactive"),
    ("constant", "constant_nonzero"), ("zero_torque_moving", "varying"), ("signed", "varying")])
def test_pair_classification_never_imputes_or_filters(case, status):
    profile, _, channels, values = maps_for()
    definition = profile.plots_by_id()[MAPS[0]]
    item = definition.series[0]
    x, y = item.x, item.semantic_name
    values[x], values[y] = np.array([0., 1, 2, 3]), np.array([1., -2, 3, 4])
    if case.startswith("missing_"):
        del channels[x if case.endswith("x") else y]
    elif case in {"nan", "inf"}:
        values[y][1] = float(case)
    elif case.startswith("wrong_"):
        name = x if case == "wrong_rpm_unit" else y
        channels[name] = replace(channels[name], unit="rad/s" if name == x else "kNm")
    elif case == "counter_torque":
        channels[y] = replace(channels[y], source_name="Wheel_CounterTorque_FL")
    elif case == "misaligned":
        values[x] = values[x][:-1]
    elif case == "inactive":
        values[x][:], values[y][:] = 0, 0
    elif case == "constant":
        values[x][:], values[y][:] = 50, -20
    elif case == "zero_torque_moving":
        values[y][:] = 0
    before = {k: v.copy() for k, v in values.items()}
    record = _paired_series_status(definition, item, channels, values, 4)
    assert record["status"] == status
    assert record["included"] == (status in {"varying", "constant_nonzero"})
    if case == "zero_torque_moving":
        from types import SimpleNamespace
        from vsm_postprocessing.diesel_powerpoint_report import _available_plots
        result = SimpleNamespace(profile=profile, plotting_result=SimpleNamespace(
            rendered_plots=[SimpleNamespace(plot_id=MAPS[0])],
            series_summaries={MAPS[0]: (SimpleNamespace(is_all_zero=True),)}))
        assert MAPS[0] in _available_plots(result)
    for key in values:
        np.testing.assert_array_equal(values[key], before[key])


def test_missing_default_x_does_not_hide_other_valid_components():
    profile, _, channels, values = maps_for()
    definition = profile.plots_by_id()[MAPS[0]]
    del channels[definition.x]
    prepared, _ = _prepare_plot_definition(definition, channels, values, 4)
    assert len(prepared.series) == 3
    assert prepared.x_channel_id == "electricsystem_em2_speed"


def test_no_selected_components_is_explicitly_unavailable():
    profile, _, channels, values = maps_for()
    selected = select_profile(profile, ProfileReportSelection(plot_series={MAPS[0]: ()}))
    assert isinstance(_prepare_plot_definition(selected.plots_by_id()[MAPS[0]], channels, values, 4), UnavailableProfilePlot)


@pytest.mark.parametrize("kind", ["electric", "hybrid", "diesel"])
def test_configuration_preserves_existing_definitions_and_source_locations(kind):
    profile = profile_for(kind)
    for key in MAPS:
        definition = profile.plots_by_id()[key]
        assert definition.render_mode == "scatter" and definition.report_group == "torque_maps"
        assert all(s.x and s.component and s.measurement_location for s in definition.series)
        assert all(not s.required for s in definition.series)
    base = load_reporting_profile(profile_path(kind).parent.parent / profile_path(kind).name)
    assert all(p.render_mode == "line" and p.report_group is None for p in base.plots)
    if kind != "diesel":
        assert profile.raw_by_semantic_name()["driveshaft_torque_rl"].source_name == "KERS_MotGenUnit_GearboxOutputTorque_3"
    if kind == "hybrid":
        source = profile.raw_by_semantic_name()["torque_map_driveshaft_rl"]
        assert source.source_name == "DriveShaft_Torque_RL" and not source.for_report
        assert profile.math_by_semantic_name()["wheel_power_rl"].expression.startswith("wheel_rotationalspeed_rl * driveshaft_torque_rl")


def test_component_selection_validates_names_and_retains_x_dependencies():
    p = profile_for("electric")
    selected = select_profile(p, ProfileReportSelection(plot_series={MAPS[0]: ("electricsystem_em2_torque",)}))
    assert [s.x for s in selected.plots_by_id()[MAPS[0]].series] == ["electricsystem_em2_speed"]
    with pytest.raises(ConfigurationError, match="unknown identifiers"):
        select_profile(p, ProfileReportSelection(plot_series={MAPS[0]: ("invented",)}))


@pytest.fixture(scope="module", params=["electric", "hybrid", "diesel"])
def map_report(request, tmp_path_factory):
    kind = request.param
    folder = tmp_path_factory.mktemp("maps_"+kind)
    data = dataset_for(kind)
    if kind != "diesel":
        for i in range(1, 5):
            data.values[:, data.channel_index(f"electricsystem_em{i}_speed")] = [-i, 10*i, 20*i, 30*i]
            data.values[:, data.channel_index(f"electricsystem_em{i}_torque")] = [0, -10*i, 20*i, 30*i]
    source = csv_for(data, folder / "scenario.csv")
    selection = ProfileReportSelection(export_channels=("track_time",), plot_series={
        MAPS[0]: ("engine_torque",) if kind == "diesel" else ("electricsystem_em2_torque", "electricsystem_em4_torque")})
    report = generate_profile_excel_report(source, profile_path(kind), folder / "excel", selection=selection,
        import_options=ImportOptions(strict=False, data_start_row=4, data_end_row=6))
    ppt = build_profile_powerpoint_report(report, folder / "powerpoint")
    return kind, report, ppt


def test_excel_native_pairs_survive_main_export_deselection_and_interval(map_report):
    kind, report, _ = map_report
    book = load_workbook(report.report_path)
    try:
        sheet = book["Torque Maps"]
        assert len(sheet._charts) == 2 and not sheet._images
        assert sheet.freeze_panes is None and sheet.sheet_view.pane is None
        assert sheet.sheet_view.zoomScale == 80
        assert report.report_channel_count == 1
        assert sheet["A40"].value == 4 and sheet["B40"].value == 1 and sheet["C40"].value == 1
        expected = report.plotting_result.values_by_semantic_name
        for chart, key in zip(sheet._charts, MAPS):
            prepared, _ = _prepare_plot_definition(report.profile.plots_by_id()[key], report.plotting_result.channels_by_semantic_name, expected, 3)
            assert chart.scatterStyle == "marker" and len(chart.series) == len(prepared.series)
            for series, item in zip(chart.series, prepared.series):
                assert series.marker.symbol == "circle" and series.graphicalProperties.line.noFill
                assert series.graphicalProperties.line.solidFill is None
                assert not series.smooth
                def actual(ref):
                    name, extent = ref.split("!")
                    assert name == "'Torque Maps'"
                    cells = sheet[extent.replace("$", "")]
                    return [row[0].value for row in cells]
                np.testing.assert_array_equal(actual(series.xVal.numRef.f), expected[item.x_channel_id])
                np.testing.assert_array_equal(actual(series.yVal.numRef.f), expected[item.channel_id])
        metadata = json.loads((report.plot_assets_dir / "profile_plot_manifest.json").read_text())
        assert metadata["source_data_start_row"] == 4 and metadata["source_data_end_row"] == 6
        assert metadata["scenario"] == "scenario.csv" and metadata["paired_series"][MAPS[0]]
    finally:
        book.close()


def test_powerpoint_maps_precede_summary_and_preserve_numbering(map_report):
    kind, report, ppt = map_report
    prs = Presentation(ppt.presentation_path)
    assert len(prs.slides) == {"electric": 16, "hybrid": 15, "diesel": 11}[kind]
    assert [s.slide_id for s in ppt.powerpoint_result.config.slides[-2:]] == ["torque_maps", "simulation_summary"]
    for number, slide in enumerate(prs.slides, 1):
        text = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
        assert f"{number} / {len(prs.slides)}" in " ".join(text.split())
        assert all(s.left >= 0 and s.top >= 0 and s.left+s.width <= prs.slide_width+10 and s.top+s.height <= prs.slide_height+10 for s in slide.shapes)
    slide = prs.slides[-2]
    subtitle = ppt.powerpoint_result.config.slides[-2].subtitle
    assert subtitle == "VSM simulation operating points showing the relationship between rotational speed and torque."
    caption = next(s for s in slide.shapes if s.has_text_frame and s.text == subtitle)
    reference = Presentation("reference_files/RoboSprayer_Electric_Report_Astauto_Colours.pptx").slides[8].shapes[3]
    assert (caption.left, caption.top, caption.width, caption.height) == (reference.left, reference.top, reference.width, reference.height)
    images = [s for s in slide.shapes if s.shape_type == 13 and s.top > 2*914400]
    assert len(images) == 2 and images[0].left+images[0].width < images[1].left
    assert caption.top+caption.height < min(s.top for s in images)
    rendered = {p.plot_id: p for p in report.plotting_result.rendered_plots}
    for image, key in zip(images, MAPS):
        assert image.image.blob == Path(rendered[key].png_file).read_bytes()


def test_feature_can_be_disabled_without_removing_existing_reports(tmp_path):
    profile, data = profile_for("diesel"), dataset_for("diesel")
    source = csv_for(data, tmp_path/"source.csv")
    selection = ProfileReportSelection(plots=tuple(p.plot_id for p in profile.plots if p.report_group != "torque_maps"),
        powerpoint_plots={"torque_maps": ()})
    report = generate_profile_excel_report(source, profile_path("diesel"), tmp_path/"excel", selection=selection)
    ppt = build_profile_powerpoint_report(report, tmp_path/"powerpoint")
    book = load_workbook(report.report_path)
    try:
        assert "Torque Maps" not in book
    finally:
        book.close()
    assert ppt.slide_count == 10


def test_unavailable_wheel_map_retains_xy_type_without_a_fabricated_point(tmp_path):
    from test_sergio_october_15 import remove_channel
    data = dataset_for("electric")
    for corner in ("fl", "fr", "rl", "rr"):
        remove_channel(data, "driveshaft_torque_"+corner)
    report = generate_profile_excel_report(csv_for(data, tmp_path/"old.csv"), profile_path("electric"), tmp_path/"report")
    book = load_workbook(report.report_path)
    try:
        sheet = book["Torque Maps"]
        assert len(sheet._charts) == 2
        chart = sheet._charts[1]
        assert chart.scatterStyle == "marker" and len(chart.series) == 1
        assert chart.series[0].marker.symbol is None and chart.series[0].graphicalProperties.line.noFill
        assert sheet["A32"].value is None and sheet["B32"].value is None
        assert "source unavailable" in sheet["N28"].value
        assert "torque_map_wheels" not in {p.plot_id for p in report.plotting_result.rendered_plots}
    finally:
        book.close()


def test_current_diesel_without_optional_inputs_keeps_valid_engine_map(tmp_path):
    from test_sergio_october_15 import remove_channel
    data, profile = dataset_for("diesel"), profile_for("diesel")
    for channel in profile.raw_channels:
        if not channel.required:
            remove_channel(data, channel.semantic_name)
    report = generate_profile_excel_report(csv_for(data, tmp_path/"required_only.csv"), profile_path("diesel"), tmp_path/"excel")
    ppt = build_profile_powerpoint_report(report, tmp_path/"powerpoint")
    prs = Presentation(ppt.presentation_path)
    assert ppt.slide_count == 9
    with ZipFile(ppt.presentation_path) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))
    assert [s.slide_id for s in ppt.powerpoint_result.config.slides[-2:]] == ["torque_maps", "simulation_summary"]
    pictures = [s for s in prs.slides[-2].shapes if s.shape_type == 13 and s.top > 2*914400]
    assert len(pictures) == 1
    text = " ".join(s.text for s in prs.slides[-2].shapes if s.has_text_frame)
    assert "PLOT DATA NOT AVAILABLE" in text and "Torque source unavailable" in text
    assert report.statistics_result.is_complete
    assert "Simulation Summary" in " ".join(s.text for s in prs.slides[-1].shapes if s.has_text_frame)
