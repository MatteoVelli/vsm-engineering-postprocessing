from __future__ import annotations

import json
from contextlib import nullcontext
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import numpy as np
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_to_tuple

from conftest import DIESEL_REFERENCE_CSV, DIESEL_REFERENCE_DESCRIPTION, require_private_reference_file
from test_diesel_profile import PROFILE_PATH, _source
from workbook_assertions import assert_universal_bottom_statistics
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from vsm_postprocessing.report_profile import load_reporting_profile


CHARTS = [
    ("Vehicle Speed & Road Gradient", ["chassis_speed", "track_gradient"], ["Vehicle Speed [kph]", "Road Gradient [%]"]),
    ("Engine Speed", ["engine_speed"], ["Engine Speed [rpm]"]),
    ("Engine Torque & Power", ["engine_torque", "engine_power"], ["Engine Torque [Nm]", "Engine Power [kW]"]),
    ("Engine Load Torque", ["engine_load"], ["Engine Load [Nm]"]),
    ("Engine Throttle", ["engine_throttle"], ["Engine Throttle [%]"]),
    ("Fuel Flow", ["fuel_flow"], ["Fuel Flow [L/h]"]),
    ("Cumulative Fuel Mass", ["engine_fuel_consumption"], ["Cumulative Fuel Mass [kg]"]),
    ("Driveshaft Torque", ["driveshaft_torque_" + c for c in ("fl", "fr", "rl", "rr")], ["Driveshaft Torque [Nm]"]),
    ("Tyre Vertical Loads", ["tyre_fz_" + c for c in ("fl", "fr", "rl", "rr")], ["Tyre Vertical Load [N]"]),
    ("Engine Oil Temperature", ["engine_oil_temperature"], ["Oil Temperature [°C]"]),
]


@pytest.fixture(scope="module")
def diesel_report(tmp_path_factory):
    source = require_private_reference_file(DIESEL_REFERENCE_CSV, DIESEL_REFERENCE_DESCRIPTION)
    return generate_profile_excel_report(source, PROFILE_PATH, tmp_path_factory.mktemp("diesel_excel"))


@pytest.fixture(scope="module")
def workbook(diesel_report):
    book = load_workbook(diesel_report.report_path, data_only=True)
    yield book
    book.close()


def _title(title):
    return "".join(run.t for paragraph in title.tx.rich.p for run in paragraph.r)


def _columns(report):
    return {c.channel_id: i for i, c in enumerate(report.report_channels, 1)}


def _summary(sheet, report):
    return {sheet.cell(1, col).value: sheet.cell(2, col).value for col in range(report.report_channel_count + 2, sheet.max_column + 1) if sheet.cell(1, col).value}


def test_reference_workbook_selection_and_units(diesel_report, workbook):
    assert workbook.sheetnames == ["Full Size Sprayer Diesel", "Rename From VSM to Astauto", "Metadata", "Statistics"]
    assert workbook["Metadata"].sheet_state == "hidden"
    assert diesel_report.sample_count == 223 and diesel_report.report_channel_count == 67
    assert_universal_bottom_statistics(diesel_report.report_path, diesel_report.sample_count)
    assert (diesel_report.vsm_count, diesel_report.math_count, diesel_report.statistic_count, diesel_report.kpi_count) == (47, 20, 51, 5)
    ids = _columns(diesel_report)
    assert "track_height" not in ids
    assert not any("electricsystem_em" in c.source_name.lower() or "battery" in c.source_name.lower() for c in diesel_report.report_channels)
    sheet = workbook.worksheets[0]
    for name, unit in [("engine_torque", "Nm"), ("engine_load", "Nm"), ("engine_speed", "rpm"), ("engine_power", "kW"), ("fuel_flow", "l/h"), ("engine_fuel_consumption", "kg"), ("engine_mechanical_energy", "kWh"), ("fuel_volume", "l")]:
        assert sheet.cell(4, ids[name]).value == unit
    mapping = workbook["Rename From VSM to Astauto"]
    assert [mapping.cell(r, 3).value for r in range(3, diesel_report.report_channel_count + 3)].count("MATH") == 20
    assert "cumulative_trapezoid" in mapping.cell(ids["engine_mechanical_energy"] + 2, 8).value
    assert sheet.cell(1, diesel_report.report_channel_count + 2).value == "Maximum Vehicle Speed [kph]"
    manifest = json.loads(diesel_report.manifest_path.read_text(encoding="utf-8"))
    assert manifest["visible_sheet_names"] == [s.title for s in workbook if s.sheet_state == "visible"]


def test_reference_executive_values_match_d1(diesel_report, workbook):
    values = _summary(workbook.worksheets[0], diesel_report)
    expected = {
        "Drive Cycle Time [min]": 3.7, "Distance [km]": 0.999242,
        "Maximum Vehicle Speed [kph]": 22.4686, "Min Road Gradient [%]": -15,
        "Max Road Gradient [%]": 15, "Engine Speed MAX [rpm]": 1996.81,
        "Average Engine Speed [rpm]": 1537.2141644144144, "Engine Torque MAX [Nm]": 826.219,
        "Engine Power MAX [kW]": 165.93862633964403,
        "Net Engine Mechanical Energy [kWh]": 3.7647817811605133, "Total Fuel Consumption [kg]": 0.772836,
        "Average Fuel Flow [l/h]": 15.277633288513513, "Maximum Fuel Flow [l/h]": 41.4741,
    }
    assert values == pytest.approx(expected, rel=1e-12)
    assert not any(word in label.lower() for label in values for word in ["battery", "edu", "economy"])


def test_every_exported_sample_and_statistic_matches_d1(diesel_report, workbook):
    sheet = workbook.worksheets[0]
    for name, col in _columns(diesel_report).items():
        actual = [sheet.cell(row, col).value for row in range(5, 228)]
        np.testing.assert_allclose(actual, diesel_report.plotting_result.values_by_semantic_name[name], rtol=2e-15, atol=1e-14)
    assert min(sheet.cell(r, _columns(diesel_report)["engine_power"]).value for r in range(5, 228)) < -8
    statistics = workbook["Statistics"]
    for row, item in enumerate(diesel_report.statistics_result.statistics, 2):
        assert statistics.cell(row, 1).value == item.definition.statistic_id
        assert statistics.cell(row, 5).value == pytest.approx(item.value, rel=2e-15, abs=1e-14)
    for row, item in enumerate(diesel_report.statistics_result.kpis, diesel_report.statistic_count + 5):
        assert statistics.cell(row, 1).value == item.definition.kpi_id
        assert statistics.cell(row, 3).value == pytest.approx(item.value, rel=2e-15)
    assert [sheet.cell(r, 1).value for r in range(228, 232)] == ["MAX", "MIN", "LAST", "FIRST"]
    assert all(sheet.cell(232, col).value is None for col in range(1, diesel_report.report_channel_count + 2))


@pytest.mark.parametrize("index,spec", list(enumerate(CHARTS)))
def test_native_chart_series_units_and_signed_axes(diesel_report, workbook, index, spec):
    title, channels, units = spec
    sheet = workbook.worksheets[0]
    chart = sheet._charts[index]
    assert _title(chart.title) == title
    assert _title(chart.x_axis.title) == "Time [s]"
    assert [_title(component.y_axis.title) for component in chart._charts] == units
    series = [series for component in chart._charts for series in component.series]
    assert len(series) == len(channels)
    for item, channel in zip(series, channels):
        col = get_column_letter(_columns(diesel_report)[channel])
        assert item.yVal.numRef.f == f"'{sheet.title}'!${col}$5:${col}$227"
        assert item.xVal.numRef.f == f"'{sheet.title}'!$A$5:$A$227"
        assert item.graphicalProperties.line.width == 31750
    for component in chart._charts:
        assert component.y_axis.delete is False and component.y_axis.tickLblPos == "nextTo"
        assert component.y_axis.majorUnit > 0
        assert component.y_axis.numFmt.sourceLinked is False
        assert component.x_axis.scaling.min <= 0 and component.x_axis.scaling.max >= 222
        for series in component.series:
            _, (col, first, _, last) = range_to_tuple(series.yVal.numRef.f)
            data = [sheet.cell(row, col).value for row in range(first, last + 1)]
            assert component.y_axis.scaling.min <= min(data)
            assert component.y_axis.scaling.max >= max(data)
    assert chart.x_axis.delete is False
    if len(chart._charts) == 2:
        assert chart._charts[1].y_axis.axPos == "r"
    assert (chart.legend is not None) == (len(channels) > 1)
    if chart.legend:
        assert chart.legend.position == "b" and chart.legend.overlay is False


def test_native_structure_styles_and_no_overlaps(diesel_report, workbook):
    sheet = workbook.worksheets[0]
    assert len(sheet._charts) == 18 and not sheet._images
    assert sheet.freeze_panes == "B6"
    assert sheet["A3"].fill.fgColor.rgb.endswith("1F4E78")
    assert sheet.cell(3, _columns(diesel_report)["engine_power"]).fill.fgColor.rgb.endswith("C65911")
    assert sheet.cell(1, diesel_report.report_channel_count + 2).font.color.rgb.endswith("FFFFFF")
    assert sheet.cell(2, diesel_report.report_channel_count + 2).number_format == "0.000"
    assert sheet.column_dimensions[get_column_letter(diesel_report.report_channel_count + 2)].width == 13
    assert "AC1:AQ1" not in {str(r) for r in sheet.merged_cells.ranges}
    assert sheet.row_dimensions[3].height == 60
    assert workbook["Statistics"].column_dimensions["D"].width == 42
    for i, chart in enumerate(sheet._charts):
        assert (chart.anchor._from.col, chart.anchor._from.row) == (diesel_report.report_channel_count + 1 + (i % 2) * 8, 2 + (i // 2) * 20)
        assert chart.anchor.ext.cx / 914400 == pytest.approx(637 / 96)
        assert chart.anchor.ext.cy / 914400 == pytest.approx(3.75)
        assert chart.layout.manualLayout.x == pytest.approx(0.13)
        for component in chart._charts:
            expected_x = 0.96 if component.y_axis.axPos == "r" else 0.01
            assert component.y_axis.title.layout.manualLayout.x == pytest.approx(expected_x)
            assert component.y_axis.title.layout.manualLayout.h == pytest.approx(0.60)
    # 7 columns (637 px) inside an 8-column stride; 18 rows inside a 20-row stride.
    assert 637 < 8 * 91 and 3.75 * 72 < 20 * 15
    with ZipFile(diesel_report.report_path) as archive:
        files = archive.namelist()
        assert len([p for p in files if p.startswith("xl/charts/chart") and p.endswith(".xml")]) == 18
        assert not any(p.startswith("xl/media/") or p.startswith("xl/externalLinks/") for p in files)
        for path in files:
            if path.endswith(".xml"):
                assert b"#REF!" not in archive.read(path)
    assert not any(c.data_type == "e" for s in workbook for row in s for c in row)


@pytest.mark.parametrize("missing,expected_charts", [("all", 7), ("fuel_flow", 17), ("engine_load", 17), ("engine_throttle", 17), ("engine_oil_temperature", 17), ("driveshaft_torque_fl", 17)])
def test_missing_optional_channels_omit_only_dependent_outputs(tmp_path, missing, expected_charts):
    profile = load_reporting_profile(PROFILE_PATH)
    omitted = [c.semantic_name for c in profile.raw_channels if not c.required] if missing == "all" else [missing, "track_height"]
    source = _source(tmp_path, profile, omitted=omitted, extra_electric=True)
    result = generate_profile_excel_report(source, PROFILE_PATH, tmp_path / "report")
    book = load_workbook(result.report_path, data_only=True)
    assert result.plot_count == expected_charts and len(book.worksheets[0]._charts) == expected_charts
    assert result.statistics_result.is_complete
    assert_universal_bottom_statistics(result.report_path, result.sample_count)
    assert not set(omitted) & set(_columns(result))
    assert not any("battery" in c.channel_id for c in result.report_channels)
    if missing in {"all", "fuel_flow"}:
        assert "fuel_volume" not in _columns(result)
        assert "Average Fuel Flow [l/h]" not in _summary(book.worksheets[0], result)
    if missing == "all":
        assert result.report_channel_count == 12
    book.close()


def test_optional_height_is_exported_when_available(tmp_path):
    profile = load_reporting_profile(PROFILE_PATH)
    result = generate_profile_excel_report(_source(tmp_path, profile), PROFILE_PATH, tmp_path / "report")
    assert result.statistic_count == 53 and "track_height" in _columns(result)
    assert result.plot_count == 18


def test_powerpoint_reuses_and_preserves_the_excel_workbook(diesel_report, tmp_path):
    before = diesel_report.report_path.read_bytes()
    result = build_profile_powerpoint_report(diesel_report, tmp_path / "ppt")
    assert result.excel_result is diesel_report and result.slide_count == 10
    assert result.presentation_path.exists()
    assert diesel_report.report_path.read_bytes() == before


def test_diesel_ui_uses_combined_reports_and_keys_download_to_current_source(tmp_path, monkeypatch):
    from vsm_postprocessing import ui_app

    output = tmp_path / "diesel.xlsx"
    output.write_bytes(b"workbook")
    presentation = tmp_path / "diesel.pptx"
    presentation.write_bytes(b"presentation")
    result = SimpleNamespace(report_path=output, presentation_path=presentation,
                             excel_result=SimpleNamespace(plotting_result=SimpleNamespace(rendered_plots=[])))
    calls = []
    monkeypatch.setattr(ui_app, "generate_reporting_profile_engineering_report", lambda *args, **kwargs: calls.append((args, kwargs)) or result)
    monkeypatch.setattr(ui_app, "_render_profile_report_completion", lambda *args: None)

    class UI:
        def __init__(self):
            self.session_state = {"profile_validation_key": "source-one"}
            self.clicked = True
            self.downloads = []

        def info(self, *args): pass
        def success(self, *args): pass
        def subheader(self, *args): pass
        def spinner(self, *args): return nullcontext()
        def button(self, label, **kwargs): return self.clicked
        def download_button(self, label, **kwargs): self.downloads.append(label)

    ui = UI()
    definition = SimpleNamespace(profile_path=PROFILE_PATH)
    summary = SimpleNamespace(is_valid=True)
    ui_app._render_diesel_report_workflow(ui, definition, tmp_path / "input.csv", "source-one", summary, "Test Machine")
    assert len(calls) == 1 and calls[0][1]["machine_name_override"] == "Test Machine"
    assert ui.downloads == ["Download Excel Engineering Report", "Download PowerPoint Engineering Report"]
    ui.clicked = False
    ui.downloads.clear()
    ui_app._render_diesel_report_workflow(ui, definition, tmp_path / "other.csv", "source-two", summary, "Test Machine")
    assert ui.downloads == []
