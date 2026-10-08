"""Independent workbook assertions shared by profile and configurable report tests."""

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def assert_universal_bottom_statistics(path, sample_count, *, data_start_row=5):
    workbook = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    try:
        sheet, numeric = workbook.worksheets[0], cached.worksheets[0]
        # The data table ends at the first blank header, before the summary strip.
        channel_count = 0
        while sheet.cell(3, channel_count + 1).value is not None:
            channel_count += 1
        assert channel_count > 0
        data_end_row = data_start_row + sample_count - 1
        rows = {sheet.cell(row, 1).value: row
                for row in range(data_end_row + 1, sheet.max_row + 1)
                if sheet.cell(row, 1).value in ('MAX', 'MIN', 'LAST', 'FIRST')}
        assert set(rows) == {'MAX', 'MIN', 'LAST', 'FIRST'}
        for operation, row in rows.items():
            assert sheet.cell(row, 1).value == operation
            assert sheet.cell(row, 1).data_type == 's'
            assert sheet.cell(row, channel_count + 1).value is None
        count = 0
        for col in range(2, channel_count + 1):
            letter = get_column_letter(col)
            extent = f'{letter}{data_start_row}:{letter}{data_end_row}'
            # Mathematical data cells may now be executable expressions. Their
            # numeric initial values, rather than formula strings, are the data
            # authority for checking the unchanged bottom-statistic caches.
            data = [numeric.cell(row, col).value for row in range(data_start_row, data_end_row + 1)]
            finite = [v for v in data if isinstance(v, (int, float))]
            expected = {'MAX': (f'=MAX({extent})', max(finite) if finite else 0),
                        'MIN': (f'=MIN({extent})', min(finite) if finite else 0),
                        'LAST': (f'={letter}{data_end_row}', data[-1] or 0),
                        'FIRST': (f'={letter}{data_start_row}', data[0] or 0)}
            for operation, (formula, value) in expected.items():
                cell = sheet.cell(rows[operation], col)
                assert cell.data_type == 'f', cell.coordinate
                assert cell.value == formula
                assert numeric[cell.coordinate].value == pytest.approx(value, rel=1e-12, abs=1e-12)
                count += 1
        assert count == 4 * (channel_count - 1)
        assert sum(cell.data_type == 'f' for row in rows.values()
                   for cell in next(sheet.iter_rows(min_row=row, max_row=row, max_col=channel_count))) == 4 * (channel_count - 1)
        assert sheet.freeze_panes == 'B6'
        return channel_count, count
    finally:
        workbook.close()
        cached.close()
