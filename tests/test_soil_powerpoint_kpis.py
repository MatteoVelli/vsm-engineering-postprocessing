"""Numerical and native-slide acceptance for presentation-only soil KPIs."""
from dataclasses import replace
import math
from types import SimpleNamespace

import numpy as np
from pptx import Presentation
from pptx.util import Inches
import pytest

from vsm_postprocessing.profile_soil_kpis import build_soil_powerpoint_metrics, soil_kpi_cards
from vsm_postprocessing.profile_statistics import CanonicalReportMetric
from vsm_postprocessing.profile_powerpoint_report_engine import _format_custom_metric_value
from test_october_soil_plots import soil_report
from test_sergio_october_15 import profile_for
from test_profile_math import _channel


def metric_input(kind="electric", power=(-20, -2, -8, -4), time=(17, 18, 100, 200), offset=37, distance=2):
    p, t = np.array(power, dtype=float), np.array(time, dtype=float)
    energy = offset + np.r_[0, np.cumsum((p[:-1]/2+p[1:]/2)*np.diff(t)/3600)]
    values = {"track_time": t, "distance_km": np.linspace(0, distance, len(t))}
    channels = {"track_time": _channel("track_time", "Track_Time", "s"),
                "distance_km": _channel("distance_km", "Track_Distance / 1000", "Km")}
    for mechanism in ("bulldozing", "compaction"):
        for name, data, unit in ((f"total_{mechanism}_resistance_power", p, "kW"),
                                 (f"tyre_{mechanism}_resistance_energy_accumulated", energy, "kWh")):
            values[name] = data.copy()
            channels[name] = _channel(name, name, unit)
    report = SimpleNamespace(profile=profile_for(kind), sample_count=len(t),
                             plotting_result=SimpleNamespace(values_by_semantic_name=values, channels_by_semantic_name=channels))
    canonical = [CanonicalReportMetric("distance_km_last", "Distance", "DISTANCE", distance, "Km",
                  "distance_km", "Distance", "MATH", "LAST", "LAST(distance_km)", "summary")]
    return report, canonical


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
@pytest.mark.parametrize("offset", (0, 37, -12))
def test_all_twelve_metrics_use_signed_power_time_average_sample_rms_and_net_energy(kind, offset):
    report, canonical = metric_input(kind, offset=offset)
    values = report.plotting_result.values_by_semantic_name
    original = {key: data.copy() for key, data in values.items()}
    metrics = build_soil_powerpoint_metrics(report, canonical)
    assert len(metrics) == (18 if kind == "electric" else 12)
    for mechanism in ("bulldozing", "compaction"):
        p = values[f"total_{mechanism}_resistance_power"]
        e = values[f"tyre_{mechanism}_resistance_energy_accumulated"]
        t = values["track_time"]
        expected = [max(p), sum((p[:-1]+p[1:])/2*np.diff(t))/(t[-1]-t[0]),
                    math.sqrt(sum(p*p)/len(p)), e[-1]-e[0], (e[-1]-e[0])/2, 2]
        actual = [m for m in metrics if m.placement_group == "soil_"+mechanism]
        np.testing.assert_allclose([m.value for m in actual], expected, rtol=3e-14, atol=1e-12)
        assert [m.unit for m in actual] == ["kW", "kW", "kW", "kWh", "kWh/km", "km"]
        assert actual[0].value == -2  # Signed max, never maximum absolute magnitude.
        assert all(m.used_sample_count == len(t) and m.omitted_sample_count == 0 for m in actual)
    for key in values:
        np.testing.assert_array_equal(values[key], original[key])


def test_mean_and_rms_include_low_power_interval_near_2400_seconds():
    report, canonical = metric_input(power=(10, 10, 0, 0, 10, 10), time=(0, 2399, 2400, 2600, 2601, 4000))
    metrics = {m.metric_id: m for m in build_soil_powerpoint_metrics(report, canonical)}
    assert metrics["soil_bulldozing_mean_power"].value == pytest.approx(9.4975)
    assert metrics["soil_bulldozing_rms_power"].value == pytest.approx(math.sqrt(400/6))
    assert metrics["soil_bulldozing_mean_power"].used_sample_count == 6


@pytest.mark.parametrize("invalid", ("missing", "nan", "infinity", "misaligned", "wrong_unit", "empty"))
def test_invalid_power_never_becomes_a_zero_or_contaminates_energy_or_other_mechanism(invalid):
    report, canonical = metric_input()
    name = "total_bulldozing_resistance_power"
    values = report.plotting_result.values_by_semantic_name
    if invalid == "missing": values.pop(name)
    elif invalid == "wrong_unit":
        channels = report.plotting_result.channels_by_semantic_name
        channels[name] = replace(channels[name], unit="W")
    elif invalid in {"empty", "misaligned"}: values[name] = np.array([] if invalid == "empty" else [1])
    else: values[name][1] = np.nan if invalid == "nan" else np.inf
    metrics = {m.metric_id:m for m in build_soil_powerpoint_metrics(report, canonical)}
    assert all(math.isnan(metrics[f"soil_bulldozing_{key}"].value) for key in ("max_power", "mean_power", "rms_power"))
    assert all("Unavailable:" in metrics[f"soil_bulldozing_{key}"].notes for key in ("max_power", "mean_power", "rms_power"))
    assert math.isfinite(metrics["soil_bulldozing_total_energy"].value)
    assert all(math.isfinite(m.value) for m in metrics.values() if m.placement_group == "soil_compaction")


@pytest.mark.parametrize("distance", (0, -1, float("nan")))
def test_invalid_distance_cannot_produce_energy_per_km(distance):
    report, canonical = metric_input(distance=distance)
    metrics = {m.metric_id:m for m in build_soil_powerpoint_metrics(report, canonical)}
    assert math.isnan(metrics["soil_bulldozing_energy_per_km"].value)
    assert math.isfinite(metrics["soil_bulldozing_total_energy"].value)
    if distance == 0:
        assert metrics["soil_bulldozing_distance"].value == 0  # Real zero distance is still a measurement.
    else:
        assert math.isnan(metrics["soil_bulldozing_distance"].value)


def test_distance_reuses_the_canonical_kpi_and_missing_energy_is_explicit():
    report, canonical = metric_input()
    canonical = [replace(canonical[0], value=4)]
    values = report.plotting_result.values_by_semantic_name
    values.pop("tyre_bulldozing_resistance_energy_accumulated")
    metrics = {m.metric_id:m for m in build_soil_powerpoint_metrics(report, canonical)}
    assert metrics["soil_compaction_distance"].value == 4  # Do not recompute a conflicting distance.
    assert metrics["soil_compaction_energy_per_km"].value == pytest.approx(metrics["soil_compaction_total_energy"].value/4)
    assert math.isnan(metrics["soil_bulldozing_total_energy"].value)
    assert math.isnan(metrics["soil_bulldozing_energy_per_km"].value)
    missing = build_soil_powerpoint_metrics(report, [])
    assert all(math.isnan(m.value) for m in missing if m.metric_id.endswith(("_distance", "_energy_per_km")))


@pytest.mark.parametrize("times", ((0, 0, 1, 2), (3, 2, 1, 0), (0, np.nan, 2, 3)))
def test_invalid_time_cannot_produce_a_mean(times):
    report, canonical = metric_input()
    report.plotting_result.values_by_semantic_name["track_time"] = np.asarray(times)
    metrics = {m.metric_id:m for m in build_soil_powerpoint_metrics(report, canonical)}
    assert math.isnan(metrics["soil_bulldozing_mean_power"].value)
    assert math.isfinite(metrics["soil_bulldozing_max_power"].value)


def test_each_slide_has_exactly_six_reference_style_cards_and_unchanged_plots(soil_report):
    _, _, excel, ppt = soil_report
    prs = Presentation(ppt.presentation_path)
    statistics = {s.statistic_id:s for s in ppt.powerpoint_result.statistics_result.statistics}
    steering = prs.slides[10]
    def card_shapes(slide):
        return [s for s in slide.shapes if s.shape_type == 1 and s.top == Inches(1.72) and s.height == Inches(1.28)]
    geometry = lambda s:(s.left,s.top,s.width,s.height,s.fill.fore_color.rgb,s.line.color.rgb,s.line.width)
    ids = [s.slide_id for s in ppt.powerpoint_result.config.slides]
    for slide_id in ("soil_bulldozing", "soil_compaction"):
        number = ids.index(slide_id)
        slide = prs.slides[number]
        cards = card_shapes(slide)
        assert len(cards) == 6
        assert [geometry(s) for s in cards] == [geometry(s) for s in card_shapes(steering)]
        text = [s.text for s in slide.shapes if s.has_text_frame and s.text.strip()
                and Inches(1.72) < s.top < Inches(3)]
        expected = [item for label,key,unit in soil_kpi_cards(slide_id)
                    for item in (label,_format_custom_metric_value(statistics[key],unit))]
        assert text == expected  # Exactly name and actual value per card; no extra explanatory text.
        pictures = [s for s in slide.shapes if s.shape_type == 13 and s.top > Inches(3)]
        template = Presentation("reference_files/RoboSprayer_Electric_Report_Astauto_Colours.pptx")
        reference_top = template.slides[8].shapes[32].top
        assert len(pictures) == 2 and all(s.top == reference_top for s in pictures)
        assert min(s.top for s in pictures) > max(s.top+s.height for s in cards)
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text in text:
                font = shape.text_frame.paragraphs[0].runs[0].font
                assert font.name == "Calibri" and font.size.pt in {8.5,19}
        assert len(ppt.powerpoint_result.config.slides[number].statistic_ids) == 6


def test_new_metrics_respect_profile_and_explicit_card_deselection():
    report, canonical = metric_input()
    report.profile = replace(report.profile, presentation=replace(report.profile.presentation, slides=tuple(
        replace(s, statistics=()) if soil_kpi_cards(s.slide_id) else s for s in report.profile.presentation.slides)))
    assert build_soil_powerpoint_metrics(report, canonical) == []
    report.profile = profile_for("diesel")
    assert build_soil_powerpoint_metrics(report, canonical) == []
