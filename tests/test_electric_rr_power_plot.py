"""The Electric presentation uses actual RR power; shared reports keep RL."""
from __future__ import annotations

import csv
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import pytest
import yaml
from lxml import etree
from matplotlib.axes import Axes
from pptx import Presentation

from test_latest_template_math import dataset_for
from vsm_postprocessing.errors import ConfigurationError
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report, _profile_powerpoint_plots
from vsm_postprocessing.report_profile import load_reporting_profile


PROFILE = Path(__file__).resolve().parents[1] / "config/report_profiles/robosprayer_electric.yaml"
PLOT = "power_at_wheels_and_edu"
RR = ("wheel_power_rr", "edu_mech_power_rr")
RL = ("wheel_power_rl", "edu_mech_power_rl")


@pytest.fixture(scope="module")
def reports(tmp_path_factory):
    root = tmp_path_factory.mktemp("asymmetric_rr_power")
    return build_asymmetric_rr_reports(root, PROFILE)


def build_asymmetric_rr_reports(root, profile_path):
    """Exercise both presentations with independently varying RL/RR sources."""
    profile = load_reporting_profile(profile_path)
    dataset = dataset_for(profile)
    for name, values in {
        "electricsystem_em3_speed": [100, 200, 300, 400],
        "electricsystem_em3_torque": [10, 20, 30, 40],
        "electricsystem_em4_speed": [900, 800, 700, 600],
        "electricsystem_em4_torque": [80, 70, 60, 50],
    }.items():
        dataset.values[:, dataset.channel_index(name)] = values
    source = root / f"Asymmetric_{profile_path.stem}.csv"
    with source.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(c.source_name for c in dataset.channels)
        writer.writerow(c.unit for c in dataset.channels)
        writer.writerows(dataset.values)
    excel = generate_profile_excel_report(source, profile_path, root / "excel")
    before = excel.report_path.read_bytes()
    old_profile = replace(profile, presentation=replace(profile.presentation, plot_series_overrides={}))
    baseline = build_profile_powerpoint_report(replace(excel, profile=old_profile), root / "baseline")
    final = build_profile_powerpoint_report(excel, root / "final")
    assert excel.report_path.read_bytes() == before
    return baseline, final


def test_electric_and_hybrid_presentations_select_rr_with_shared_rl_plots():
    electric = load_reporting_profile(PROFILE)
    assert tuple(s.semantic_name for s in electric.plots_by_id()[PLOT].series) == RL
    assert tuple(s.semantic_name for s in electric.presentation.plot_series_overrides[PLOT]) == RR
    assert tuple(s.label for s in electric.presentation.plot_series_overrides[PLOT]) == ("Wheel Power RR", "EDU Mech Power RR")
    for kind, filename in [("hybrid", "robosprayer_hybrid.yaml"), ("diesel", "full_size_sprayer_diesel.yaml")]:
        profile = load_reporting_profile(PROFILE.with_name(filename))
        if kind == "hybrid":
            assert tuple(s.semantic_name for s in profile.presentation.plot_series_overrides[PLOT]) == RR
            assert tuple(s.semantic_name for s in profile.plots_by_id()[PLOT].series) == RL
        else:
            assert not profile.presentation.plot_series_overrides


def test_rendered_plot_and_metadata_select_actual_rr(reports):
    _, final = reports
    ppt_plot = next(p for p in final.powerpoint_result.plotting_result.rendered_plots if p.plot_id == PLOT)
    excel_plot = next(p for p in final.excel_result.plotting_result.rendered_plots if p.plot_id == PLOT)
    assert ppt_plot.primary_series_ids == RR and not ppt_plot.secondary_series_ids
    assert not set(RL) & set(ppt_plot.primary_series_ids)
    assert ppt_plot.legend_labels == ("Wheel Power RR", "EDU Mech Power RR")
    assert excel_plot.primary_series_ids == RL
    assert ppt_plot.png_file != excel_plot.png_file
    slide = Presentation(final.presentation_path).slides[8]
    assert Path(ppt_plot.png_file).read_bytes() in [s.image.blob for s in slide.shapes if s.shape_type == 13]
    summaries = final.powerpoint_result.plotting_result.series_summaries[PLOT]
    assert tuple(s.semantic_name for s in summaries) == RR


def test_real_renderer_receives_asymmetric_rr_samples(reports, tmp_path, monkeypatch):
    _, final = reports
    values = final.excel_result.plotting_result.values_by_semantic_name
    observed = {}
    original_plot = Axes.plot

    def record(axis, *args, **kwargs):
        observed[kwargs.get("label")] = (np.array(args[0]), np.array(args[1]))
        return original_plot(axis, *args, **kwargs)

    monkeypatch.setattr(Axes, "plot", record)
    _profile_powerpoint_plots(final.excel_result, tmp_path)
    for rr, rl, label in zip(RR, RL, ("Wheel Power RR", "EDU Mech Power RR")):
        assert not np.array_equal(values[rr], values[rl])
        np.testing.assert_array_equal(observed[label][0], values["distance_km"])
        np.testing.assert_array_equal(observed[label][1], values[rr])
    np.testing.assert_allclose(observed["Wheel Power RR"][1],
        values["wheel_rotationalspeed_rr"] * values["driveshaft_torque_rr"] / 9548.8)
    np.testing.assert_allclose(observed["EDU Mech Power RR"][1],
        values["electricsystem_em4_speed"] * values["electricsystem_em4_torque"] / 9548.8)
    assert "Wheel Power RL" not in observed and "EDU Mech Power RL" not in observed


def test_all_six_cards_preserved_and_rr_sources_correct(reports):
    baseline, final = reports
    ids = final.excel_result.profile.presentation.slides_by_id()["traction_auxiliaries"].statistics
    assert ids == ("edu_mech_power_rr_max", "wheel_power_rr_max", "edu_speed_rl_max", "edu_torque_rl_max",
                   "tyre_total_energy_accumulated_last", "auxiliary_energy_accumulated_max")
    before = {s.statistic_id: s for s in baseline.powerpoint_result.statistics_result.statistics}
    after = {s.statistic_id: s for s in final.powerpoint_result.statistics_result.statistics}
    assert all(after[sid] == before[sid] for sid in ids)
    for mid, source in [("edu_mech_power_rr_max", "edu_mech_power_rr"), ("wheel_power_rr_max", "wheel_power_rr"),
                        ("edu_speed_rl_max", "electricsystem_em3_speed"), ("edu_torque_rl_max", "electricsystem_em3_torque"),
                        ("tyre_total_energy_accumulated_last", "tyre_total_resistance_energy_accumulated"),
                        ("auxiliary_energy_accumulated_max", "auxiliary_energy_consumption_accumulated")]:
        assert after[mid].channel_id == source
    values = final.excel_result.plotting_result.values_by_semantic_name
    assert after["edu_mech_power_rr_max"].value == np.max(values["edu_mech_power_rr"])
    assert after["wheel_power_rr_max"].value == np.max(values["wheel_power_rr"])
    texts = [s.text for s in Presentation(final.presentation_path).slides[8].shapes if s.has_text_frame]
    assert {"RR EDU MAX POWER", "RR WHEEL MAX POWER", "EDU MAX SPEED", "EDU MAX TORQUE", "TYRE TOTAL ENERGY", "AUX ENERGY"} <= set(texts)
    assert "RR EDU MAX SPEED" not in texts and "RR EDU MAX TORQUE" not in texts


def test_only_target_image_changes_and_slide_xml_preserves_layout(reports):
    baseline, final = reports
    old, new = Presentation(baseline.presentation_path), Presentation(final.presentation_path)
    changed = []
    for number, (before, after) in enumerate(zip(old.slides, new.slides), 1):
        assert etree.tostring(before.element) == etree.tostring(after.element)
        old_images = [s for s in before.shapes if s.shape_type == 13]
        new_images = [s for s in after.shapes if s.shape_type == 13]
        for a, b in zip(old_images, new_images):
            assert (a.left, a.top, a.width, a.height) == (b.left, b.top, b.width, b.height)
            if a.image.blob != b.image.blob:
                changed.append(number)
    assert changed == [9]
    assert old.slide_width == new.slide_width and old.slide_height == new.slide_height
    with ZipFile(baseline.presentation_path) as a, ZipFile(final.presentation_path) as b:
        assert set(a.namelist()) == set(b.namelist())
        parts = [p for p in a.namelist() if a.read(p) != b.read(p)]
        assert len(parts) == 1 and parts[0].startswith("ppt/media/")


@pytest.mark.parametrize("profile_name", ["hybrid", "diesel"])
def test_unconfigured_reports_reuse_original_plots(reports, tmp_path, profile_name):
    _, final = reports
    name = "robosprayer_hybrid.yaml" if profile_name == "hybrid" else "full_size_sprayer_diesel.yaml"
    profile = load_reporting_profile(PROFILE.with_name(name))
    # Explicitly unconfigured presentations retain the shared rendering path.
    profile = replace(profile, presentation=replace(profile.presentation, plot_series_overrides={}))
    excel = replace(final.excel_result, profile=profile)
    assert _profile_powerpoint_plots(excel, tmp_path) is excel.plotting_result
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("overrides,error", [([], "YAML mapping"), ({PLOT: []}, "non-empty YAML list"),
    ({PLOT: [{"semantic_name": "wheel_power_rr", "axis": "invalid"}]}, "primary or secondary"),
    ({"missing_plot": [{"semantic_name": "wheel_power_rr"}]}, "Unknown presentation.plot_series_overrides")])
def test_invalid_presentation_overrides_fail_validation(tmp_path, overrides, error):
    raw = yaml.safe_load(PROFILE.read_text(encoding="utf-8"))
    raw["presentation"]["plot_series_overrides"] = overrides
    path = tmp_path / "invalid.yaml"
    path.write_text(yaml.safe_dump(raw), encoding="utf-8")
    with pytest.raises(ConfigurationError, match=error):
        load_reporting_profile(path)
