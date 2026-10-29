"""Independent acceptance checks for Sergio's 15 October 2026 revision."""
from dataclasses import replace
from functools import lru_cache
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import pytest
import yaml

from vsm_postprocessing.errors import ConfigurationError, MathChannelError
from vsm_postprocessing.excel_formulas import profile_temporal_formula, statistic_formula
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.math_engine import _interval_energy_kwh
from vsm_postprocessing.profile_math import calculate_profile_math_channels
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.report_profile import load_reporting_profile, resolve_profile
from vsm_postprocessing.statistics_engine import compute_statistic
from vsm_postprocessing.report_selection import ProfileReportSelection, select_profile
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from test_profile_math import _channel, _dataset

ROOT = Path("config/report_profiles/sergio_2026_10")
KINDS = ("electric", "hybrid", "diesel")
WHEELS = ("fl", "fr", "rl", "rr")


def profile_path(kind):
    return ROOT / (("full_size_sprayer_diesel" if kind == "diesel" else "robosprayer_" + kind) + ".yaml")


@lru_cache
def profile_for(kind):
    return load_reporting_profile(profile_path(kind))


def dataset_for(kind, times=(0, 1, 2, 3)):
    profile = profile_for(kind)
    channels = [_channel(c.semantic_name, c.source_name, c.unit, i + 1)
                for i, c in enumerate(profile.raw_channels)]
    data = _dataset(channels, np.ones((len(times), len(channels))))
    values = {c.channel_id: data.values[:, i] for i, c in enumerate(channels)}
    values["track_time"][:] = times
    values["track_distance"][:] = np.linspace(0, 300, len(times))
    if kind == "diesel":
        values["engine_fuel_consumption"][:] = np.linspace(0, 2, len(times))
    else:
        values["electricsystem_battery_energy"][:] = np.linspace(50, 47, len(times))
        values["electricsystem_battery_soc"][:] = np.linspace(100, 94, len(times))
    return data


def remove_channel(data, name):
    index = data.channel_index(name)
    data.channels.pop(index)
    data.values = np.delete(data.values, index, axis=1)


def csv_for(data, path):
    import csv
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow([c.source_name for c in data.channels])
        writer.writerow([c.unit or "" for c in data.channels])
        writer.writerows(data.values)
    return path


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
def test_measured_sources_and_signed_four_corner_power(kind):
    profile, data = profile_for(kind), dataset_for(kind)
    total_torque, total_power = np.zeros(4), np.zeros(4)
    for i, wheel in enumerate(WHEELS, 1):
        definition = profile.raw_by_semantic_name()[f"driveshaft_torque_{wheel}"]
        assert definition.source_name == f"KERS_MotGenUnit_GearboxOutputTorque_{i}"
        assert definition.channel_type == "VSM" and not definition.aliases
        torque = np.array([0, 20, -30, 40]) * i
        speed = np.array([0, 50, 100, -10]) * i
        data.values[:, data.channel_index(definition.semantic_name)] = torque
        data.values[:, data.channel_index(f"wheel_rotationalspeed_{wheel}")] = speed
        total_torque += torque
        total_power += torque * speed * 2 * np.pi / 60000
    result = calculate_profile_math_channels(data, profile)
    np.testing.assert_allclose(result.values_by_semantic_name["wheel_total_torque"], total_torque)
    np.testing.assert_allclose(result.values_by_semantic_name["wheel_power_total"], total_power)
    assert not any(c.semantic_name.startswith("driveshaft_torque") for c in profile.math_channels)


def test_diesel_measured_auxiliary_mass_and_zero_front_wheels():
    data, profile = dataset_for("diesel"), profile_for("diesel")
    data.values[:, data.channel_index("agrochemical_discharge_force")] = -9.80665 * np.array([100, 90, 80, 70])
    data.values[:, data.channel_index("engine_speed")] = 1800
    for i in (1, 2):
        data.values[:, data.channel_index(f"engine_auxiliary_torque_{i}")] = i * 10
    for wheel in WHEELS:
        assert profile.raw_by_semantic_name()[f"driveshaft_torque_{wheel}"].source_name == f"DriveShaft_Torque_{wheel.upper()}"
        data.values[:, data.channel_index(f"driveshaft_torque_{wheel}")] = 0 if wheel in ("fl", "fr") else 100
        data.values[:, data.channel_index(f"wheel_speed_{wheel}")] = 60
    result = calculate_profile_math_channels(data, profile)
    np.testing.assert_allclose(result.values_by_semantic_name["agrochemical_discharge"], [100, 90, 80, 70])
    np.testing.assert_allclose(result.values_by_semantic_name["total_auxiliary_power"], 1800 * 30 * 2 * np.pi / 60000)
    np.testing.assert_allclose(result.values_by_semantic_name["wheel_power_fl"], 0)
    np.testing.assert_allclose(result.values_by_semantic_name["wheel_power_total"], 200 * 60 * 2 * np.pi / 60000)
    assert not any("gearbox-output" in m for m in result.resolution.diagnostics)


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
def test_absent_channels_never_fall_back_to_old_driveshaft(kind):
    data = dataset_for(kind)
    if kind == "hybrid":
        # The maps have explicit, separate driveshaft inputs. Avoid creating
        # duplicate raw headers when substituting the old sources in this test.
        for corner in ("rl", "rr"):
            remove_channel(data, f"torque_map_driveshaft_{corner}")
    for wheel in WHEELS:
        index = data.channel_index(f"driveshaft_torque_{wheel}")
        data.channels[index] = replace(data.channels[index], source_name=f"DriveShaft_Torque_{wheel.upper()}")
    result = calculate_profile_math_channels(data, profile_for(kind))
    assert result.resolution.is_valid
    assert "wheel_power_total" not in result.values_by_semantic_name
    assert "tyre_rolling_resistance_energy_accumulated" in result.values_by_semantic_name
    assert all(any(f"GearboxOutputTorque_{i}" in msg for msg in result.resolution.diagnostics) for i in range(1, 5))


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
@pytest.mark.parametrize("moving", (True, False))
def test_all_zero_warning_is_separate_from_absence(kind, moving):
    data = dataset_for(kind)
    for wheel in WHEELS:
        data.values[:, data.channel_index(f"driveshaft_torque_{wheel}")] = 0
        data.values[:, data.channel_index(f"wheel_rotationalspeed_{wheel}")] = 5 if moving else 0
    data.values[:, data.channel_index("chassis_speed")] = 1 if moving else 0
    result = calculate_profile_math_channels(data, profile_for(kind))
    assert any("All four measured" in msg for msg in result.resolution.diagnostics) == moving
    np.testing.assert_array_equal(result.values_by_semantic_name["wheel_power_total"], 0)
    assert not any(c.definition.semantic_name.startswith("driveshaft") for c in result.unavailable)


@pytest.mark.parametrize("power", (10.0, 0.0, -10.0))
@pytest.mark.parametrize("times", ((0, 3600), (0, 1, 1100, 3600), (17, 18, 1117, 3617)))
def test_analytical_energy_and_sample_count_independence(power, times):
    increments = _interval_energy_kwh(np.full(len(times), power), np.asarray(times))
    assert increments[0] == 0
    assert increments.sum() == pytest.approx(power)


def test_variable_power_uses_trapezoids_and_real_intervals():
    np.testing.assert_allclose(_interval_energy_kwh([0, 10, -10], [0, 1800, 3600]), [0, 2.5, 0])


@pytest.mark.parametrize("power,times", [([1, np.nan], [0, 1]), ([1, 2], [0, 0]),
                                        ([1, 2], [1, 0]), ([1, 2], [0, np.nan]),
                                        ([1], [0]), ([1, 2], [0])])
def test_invalid_energy_inputs_are_explicitly_rejected(power, times):
    with pytest.raises(MathChannelError):
        _interval_energy_kwh(power, times)


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
def test_auxiliary_and_both_soil_energy_aggregations(kind):
    data, profile = dataset_for(kind, (0, 900, 1800, 3600)), profile_for(kind)
    high = profile.raw_by_semantic_name()["auxiliary_high_voltage_power_consumption"]
    assert high.source_name == "User_defined_Data_High_Voltage_Auxiliary_signal1_0"
    data.values[:, data.channel_index(high.semantic_name)] = 8
    data.values[:, data.channel_index("electricsystem_auxiliary_lowvoltagepowerconsumption")] = 2
    for component in ("bulldozing", "compaction"):
        assert profile.math_by_semantic_name()[f"total_{component}_resistance_power"].for_plot
        for wheel in WHEELS:
            data.values[:, data.channel_index(f"tyre_{component}resistancepower_{wheel}")] = 2.5
    result = calculate_profile_math_channels(data, profile)
    values = result.values_by_semantic_name
    assert high.semantic_name not in values  # Measured raw channel, not a fabricated MATH constant.
    for power, energy in [("total_auxiliary_power", "auxiliary_energy_consumption_accumulated"),
                          ("total_bulldozing_resistance_power", "tyre_bulldozing_resistance_energy_accumulated"),
                          ("total_compaction_resistance_power", "tyre_compaction_resistance_energy_accumulated")]:
        np.testing.assert_allclose(values[power], 10)
        assert values[energy][0] == 0 and values[energy][-1] == pytest.approx(10)


@pytest.mark.parametrize("kind", KINDS)
def test_optional_temperature_or_height_absence_preserves_other_results(kind):
    data, profile = dataset_for(kind), profile_for(kind)
    name = "engine_water_temperature" if kind == "diesel" else "track_height"
    remove_channel(data, name)
    result = calculate_profile_math_channels(data, profile)
    assert result.resolution.is_valid and result.is_complete
    assert any(name == item.definition.semantic_name for item in result.resolution.missing_optional)
    assert any(profile.raw_by_semantic_name()[name].source_name in msg for msg in result.resolution.diagnostics)


def test_temperature_source_is_configurable_and_not_ambient(tmp_path):
    data, base = dataset_for("diesel"), profile_for("diesel")
    index = data.channel_index("engine_water_temperature")
    data.channels[index] = replace(data.channels[index], source_name="MeasuredCoolantOutlet")
    data.values[:, index] = [75, 80, 85, 90]
    path = tmp_path / "coolant.yaml"
    path.write_text(yaml.safe_dump({"version": 1, "profile": {"profile_id": base.profile_id, "name": "Diesel", "powertrain": "diesel", "extends": str(profile_path("diesel").resolve())},
                                   "overrides": {"raw": {"engine_water_temperature": {"source_name": "MeasuredCoolantOutlet"}}}}), encoding="utf-8")
    profile = load_reporting_profile(path)
    stats = calculate_profile_statistics(data, profile)
    assert next(s.value for s in stats.statistics if s.definition.statistic_id == "engine_water_temperature_max") == 90


@pytest.mark.parametrize("count", (3, 17, 3853))
def test_rms_counts_window_and_last_numeric(count):
    values = np.resize([3.0, 4.0, np.nan], count)
    finite = values[np.isfinite(values)]
    value, used, omitted = compute_statistic(values, "rms", "omit")
    assert value == pytest.approx(np.sqrt(np.dot(finite, finite) / len(finite)))
    assert used == len(finite) and omitted == count - used
    for operation, expected in [("max", finite.max()), ("min", finite.min()), ("last", finite[-1])]:
        assert compute_statistic(values, operation, "omit")[0] == expected
    assert compute_statistic([0, 3, 4, 100][1:3], "rms")[0] == pytest.approx(np.sqrt(12.5))


@pytest.fixture(scope="module", params=KINDS)
def revised_report(request, tmp_path_factory):
    kind = request.param
    folder = tmp_path_factory.mktemp("oct15_" + kind)
    data = dataset_for(kind, (0, 900, 1800, 3600))
    source = csv_for(data, folder / "measured.csv")
    return kind, generate_profile_excel_report(source, profile_path(kind), folder / "report")


def chart_text(title):
    return "".join(run.t for p in title.tx.rich.p for run in p.r)


def test_report_header_heights_keep_compact_first_row(revised_report):
    kind, report = revised_report
    book = load_workbook(report.report_path)
    try:
        sheet = book.worksheets[0]
        assert [sheet.row_dimensions[row].height for row in (1, 2, 3, 4)] == [
            24, 24, 60 if kind == "diesel" else 42, 22]
    finally:
        book.close()


def test_live_formulas_ranges_caches_and_chart_axes(revised_report):
    kind, report = revised_report
    book, cached = load_workbook(report.report_path), load_workbook(report.report_path, data_only=True)
    try:
        main = book.worksheets[0]
        columns = {c.channel_id: i for i, c in enumerate(report.report_channels, 1)}
        assert book.calculation.fullCalcOnLoad and book.calculation.calcMode == "auto"
        for name in ("wheel_power_total", "total_auxiliary_power", "tyre_rolling_resistance_energy_kwh", "tyre_rolling_resistance_energy_accumulated"):
            col = columns[name]
            assert main.cell(5, col).data_type == "f"
            for offset in range(4):
                assert cached.worksheets[0].cell(5 + offset, col).value == pytest.approx(report.plotting_result.values_by_semantic_name[name][offset])
        for row in book["Statistics"].iter_rows(min_row=2):
            if row[2].value in ("rms", "max", "min", "last", "first", "sum"):
                assert row[4].data_type == "f"
                assert "5:" in row[4].value and "8" in row[4].value
                assert "17417" not in row[4].value
        titles = {chart_text(c.title): c for c in main._charts}
        if kind == "diesel":
            assert "Engine Water Temperature" in titles
            assert chart_text(titles["Engine Water Temperature"].y_axis.title).endswith("[\u00b0C]")
        else:
            for title in ("Tyre-Soil Bulldozing", "Tyre-Soil Compaction"):
                chart = titles[title]
                assert chart_text(chart.y_axis.title) == "Power [kW]"
                assert chart_text(chart._charts[1].y_axis.title) == "Energy [kWh]"
            road = titles["Road Profile"]
            assert chart_text(road.y_axis.title) == "Road Gradient [%]"
            assert chart_text(road._charts[1].y_axis.title) == "Road Height [m]"
            raw = report.profile.raw_by_semantic_name()
            assert raw["auxiliary_high_voltage_power_consumption"].semantic_name in columns
    finally:
        book.close(); cached.close()


def test_temporal_excel_formulas_guard_time_and_do_not_extrapolate():
    columns = {"p": 2, "t": 1}
    first = profile_temporal_formula("interval_energy_kwh(p, t)", columns, 5)
    assert first == "=IF(COUNT($B5,$A5)=2,0,NA())"
    later = profile_temporal_formula("interval_energy_kwh(p, t)", columns, 6)
    assert "$A6>$A5" in later and "($A6-$A5)/3600" in later
    assert "COUNT" in statistic_formula("rms", "C5:C3857")
    assert "LOOKUP" in statistic_formula("last", "C5:C8")


def test_override_typos_cannot_leave_stale_sources_active(tmp_path):
    path = tmp_path / "bad.yaml"
    path.write_text(yaml.safe_dump({"version": 1, "profile": {"profile_id": "bad", "name": "Bad", "extends": str(profile_path("electric").resolve())},
                                   "overrides": {"raw": {"misspelt_name": {"source_name": "Wrong"}}}}), encoding="utf-8")
    with pytest.raises(ConfigurationError, match="Unknown overrides.raw"):
        load_reporting_profile(path)


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
def test_legacy_export_generates_partial_modern_excel_and_powerpoint(kind, tmp_path):
    data = dataset_for(kind)
    for name in [*(f"driveshaft_torque_{w}" for w in WHEELS), "auxiliary_high_voltage_power_consumption", "track_height"]:
        remove_channel(data, name)
    source = csv_for(data, tmp_path / "legacy.csv")
    report = generate_profile_excel_report(source, profile_path(kind), tmp_path / "excel")
    names = {c.channel_id for c in report.report_channels}
    assert "wheel_power_total" not in names and "total_auxiliary_power" not in names
    assert "tyre_bulldozing_resistance_energy_accumulated" in names
    ppt = build_profile_powerpoint_report(report, tmp_path / "ppt")
    from pptx import Presentation
    presentation = Presentation(ppt.presentation_path)
    assert len(presentation.slides) == (16 if kind == "electric" else 15)
    assert ppt.powerpoint_result.config.slides[-2].slide_id == "torque_maps"
    text = " ".join(s.text for slide in presentation.slides for s in slide.shapes if s.has_text_frame)
    assert "PLOT DATA NOT AVAILABLE" in text


def test_nonfinite_optional_source_disables_only_dependent_analysis(tmp_path):
    data = dataset_for("electric")
    data.values[1, data.channel_index("tyre_bulldozingresistancepower_fl")] = np.nan
    data.values[2, data.channel_index("driveshaft_torque_rr")] = np.nan
    report = generate_profile_excel_report(csv_for(data, tmp_path / "gaps.csv"), profile_path("electric"), tmp_path / "report")
    values = report.math_result.values_by_semantic_name
    assert "total_bulldozing_resistance_power" not in values and "wheel_power_total" not in values
    assert "total_compaction_resistance_power" in values
    assert any("non-finite" in msg for msg in report.resolution.diagnostics)


def test_selection_exports_only_requested_channels_and_preserves_plot(tmp_path):
    data, profile = dataset_for("electric"), profile_for("electric")
    selection = ProfileReportSelection(export_channels=("track_time", "track_gradient", "track_height"),
                                       math_channels=(), statistics=(), kpis=(), plots=("road_profile",))
    report = generate_profile_excel_report(csv_for(data, tmp_path / "selection.csv"), profile_path("electric"), tmp_path / "report", selection=selection)
    assert [c.channel_id for c in report.report_channels] == ["track_time", "track_gradient", "track_height"]
    book = load_workbook(report.report_path)
    try:
        # The unexported distance is deliberately not reintroduced. Use the shared PNG.
        assert len(book.worksheets[0]._images) == 1 and not book.worksheets[0]._charts
    finally:
        book.close()
    # The optional reporting KPI retains its mass dependency without adding
    # either mass or distance back to the explicitly selected export columns.
    assert report.math_result.calculation_order == ["distance_km", "agrochemical_discharge"]


def test_new_soil_plots_replace_existing_powerpoint_slots(revised_report, tmp_path):
    kind, report = revised_report
    if kind == "diesel":
        pytest.skip("Soil plots apply to Electric and Hybrid profiles")
    selection = ProfileReportSelection(powerpoint_plots={"traction_auxiliaries": ("tyre_soil_bulldozing", "tyre_soil_compaction")})
    profile = select_profile(report.profile, selection)
    result = build_profile_powerpoint_report(replace(report, profile=profile), tmp_path / "ppt")
    payload = yaml.safe_load(result.config_path.read_text(encoding="utf-8"))
    slide = next(s for s in payload["slides"] if s["slide_id"] == "traction_auxiliaries")
    assert slide["plots"] == ["tyre_soil_bulldozing", "tyre_soil_compaction"]
    from pptx import Presentation
    assert len(Presentation(result.presentation_path).slides) == (16 if kind == "electric" else 15)
    assert result.powerpoint_result.config.slides[-2].slide_id == "torque_maps"


def test_ui_modern_profile_mapping_and_selections(tmp_path, monkeypatch):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    from types import SimpleNamespace
    from vsm_postprocessing import ui_app
    path = csv_for(dataset_for("diesel"), tmp_path / "diesel.csv")
    monkeypatch.setattr(st, "file_uploader", lambda *a, **k: SimpleNamespace(name=path.name, getvalue=path.read_bytes))
    monkeypatch.setattr(ui_app, "UI_WORKSPACE", tmp_path / "workspace")
    at = AppTest.from_string("from vsm_postprocessing.ui_app import main\nmain()", default_timeout=60).run()
    at.selectbox[0].set_value("full_size_sprayer_diesel").run()
    assert not at.exception
    assert any(s.label == "Engine Water Temperature source" for s in at.selectbox)
    assert any(s.label == "Channels exported to Excel" for s in at.multiselect)
    at.button[0].click().run()
    assert not at.exception and not at.button[1].disabled


def test_analysis_window_controls_excel_rms_ranges(tmp_path):
    from vsm_postprocessing.importer import ImportOptions
    data = dataset_for("diesel", (0, 1, 2, 3, 4, 5))
    source = csv_for(data, tmp_path / "window.csv")
    # CSV row 3 starts data. Analyze the central four samples only.
    report = generate_profile_excel_report(source, profile_path("diesel"), tmp_path / "report",
                                           ImportOptions(data_start_row=4, data_end_row=7))
    assert report.sample_count == 4
    assert report.plotting_result.values_by_semantic_name["track_time"].tolist() == [1, 2, 3, 4]
    book = load_workbook(report.report_path)
    try:
        formulas = [c.value for row in book.worksheets[0].iter_rows(min_row=1, max_row=2) for c in row if c.data_type == "f"]
        assert len(formulas) == 2 and all("8" in f and "COUNT" in f for f in formulas)
    finally:
        book.close()


@pytest.mark.parametrize("kind,start_row,first_col,hv_row,hv_col", [
    ("electric", 208, "GY", 227, "HR"), ("hybrid", 216, "HF", 235, "HY")])
def test_real_07_mappings_and_zero_torque_warn_without_validating_traction(kind, start_row, first_col, hv_row, hv_col):
    from conftest import require_private_reference_file
    from openpyxl.utils import column_index_from_string
    path = require_private_reference_file(Path(profile_for(kind).metadata.excel_template), "Sergio _07 authoritative template")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        selection = book.worksheets[1]
        for i in range(4):
            assert selection.cell(start_row + i, 2).value == f"KERS_MotGenUnit_GearboxOutputTorque_{i+1}"
        assert selection.cell(hv_row, 2).value == "User_defined_Data_High_Voltage_Auxiliary_signal1_0"
        rows = list(book.worksheets[0].values)
    finally:
        book.close()
    data = [row for row in rows[4:] if isinstance(row[0], (float, int))]
    assert len(data) == 3853
    headers = {str(name).strip(): i for i, name in enumerate(rows[2])}
    first = column_index_from_string(first_col) - 1
    base = profile_for(kind)
    names = {"track_time", "chassis_speed", *(f"driveshaft_torque_{w}" for w in WHEELS),
             *(f"wheel_rotationalspeed_{w}" for w in WHEELS)}
    profile = replace(base, raw_channels=tuple(c for c in base.raw_channels if c.semantic_name in names),
                      math_channels=tuple(c for c in base.math_channels if c.semantic_name.startswith("wheel_")),
                      statistics=(), kpis=(), plots=(), presentation=None)
    columns, channels = [], []
    for i, c in enumerate(profile.raw_channels, 1):
        col = 0 if c.semantic_name == "track_time" else headers[c.report_name.strip()]
        if c.semantic_name.startswith("driveshaft_torque_"):
            assert col == first + WHEELS.index(c.semantic_name.rsplit("_", 1)[-1])
        columns.append(col)
        channels.append(_channel(c.semantic_name, c.source_name, c.unit, i))
    result = calculate_profile_math_channels(_dataset(channels, [[r[col] for col in columns] for r in data]), profile)
    assert any("All four measured" in msg for msg in result.resolution.diagnostics)
    np.testing.assert_array_equal(result.values_by_semantic_name["wheel_power_total"], 0)
    hv = [r[column_index_from_string(hv_col)-1] for r in data]
    assert min(hv) == 0 and max(hv) == 10


def test_current_diesel_csv_coolant_source_and_corrected_auxiliary_calculations():
    from conftest import DIESEL_REFERENCE_CSV, require_private_reference_file
    from vsm_postprocessing.importer import load_data_file
    source = require_private_reference_file(DIESEL_REFERENCE_CSV, "Diesel measured export")
    data = load_data_file(source)
    result = calculate_profile_math_channels(data, profile_for("diesel"))
    resolution = result.resolution
    assert resolution.resolved["engine_water_temperature"].channel.source_name == "Engine_Water_Temperature"
    speed = data.values[:, data.channel_index(resolution.resolved["engine_speed"].channel.channel_id)]
    for i in (1, 2):
        torque = data.values[:, data.channel_index(resolution.resolved[f"engine_auxiliary_torque_{i}"].channel.channel_id)]
        np.testing.assert_allclose(result.values_by_semantic_name[f"auxiliary_power_{i}"], speed * torque * 2 * np.pi / 60000)


def test_diesel_statistic_deselection_is_respected_in_excel_and_powerpoint(tmp_path):
    profile = profile_for("diesel")
    selection = ProfileReportSelection(math_channels=(), statistics=(), kpis=(), plots=(),
        export_channels=("track_time", "chassis_speed"),
        powerpoint_plots={slide.slide_id: () for slide in profile.presentation.slides},
        powerpoint_statistics=(), bottom_operations=())
    report = generate_profile_excel_report(csv_for(dataset_for("diesel"), tmp_path / "selected.csv"),
                                           profile_path("diesel"), tmp_path / "excel", selection=selection)
    assert not report.statistics_result.statistics and not report.statistics_result.kpis
    book = load_workbook(report.report_path)
    try:
        assert not any(c.value in {"MAX", "MIN", "FIRST", "LAST"} for row in book.worksheets[0] for c in row)
        assert not book.worksheets[0]._charts and not book.worksheets[0]._images
    finally:
        book.close()
    ppt = build_profile_powerpoint_report(report, tmp_path / "ppt")
    assert ppt.presentation_path.is_file()
