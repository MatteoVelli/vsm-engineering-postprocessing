from collections import Counter
from types import SimpleNamespace
from xml.etree import ElementTree
from zipfile import ZipFile

import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill
from openpyxl.utils import get_column_letter

from vsm_postprocessing.excel_report_engine import (
    _engineering_axis_scale, _write_profile_plots_on_report_sheet,
)
from vsm_postprocessing.report_profile import ProfilePlotDefinition, ProfilePlotSeriesDefinition
from test_profile_statistics import _channel


@pytest.mark.parametrize('values,label,expected_format', [
    ([0, 1 / 60, 4 / 60], 'Time [min]', '0.00'),
    ([0, 12], 'Distance [km]', '0.0'),
    ([-30, -10, 4], 'Power [kW]', '0.0'),
    ([0, .012, .045], 'Energy [kWh]', '0.00'),
    ([70, 80, 95], 'SOC [%]', '0.0'),
    ([13000, 17000], 'Wheel Load [N]', '0'),
    ([200, 350], 'Torque [Nm]', '0'),
    ([0, .02, .04], 'Steering Angle [deg]', '0.00'),
    ([0, 5], 'Fuel Consumption [kg]', '0.0'),
    ([0, .05], 'Fuel Flow [l/h]', '0.00'),
    ([20, 40], 'Temperature [C]', '0.0'),
    ([1e-7, 4e-7], 'Steering Angle [deg]', '0.0E+00'),
])
def test_engineering_ticks_and_precision(values, label, expected_format):
    data = np.asarray(values, dtype=float)
    before = data.copy()
    low, high, step, fmt = _engineering_axis_scale(data, label)
    assert low <= min(values) <= max(values) <= high
    assert 5 <= round((high - low) / step) + 1 <= 7
    assert fmt == expected_format
    np.testing.assert_array_equal(data, before)


@pytest.mark.parametrize('values', [[0, 0], [10, 10], [-4, -4], [99.99, 100.01], [np.nan, -2, 3]])
def test_constant_offset_and_signed_axes_do_not_clip_or_force_zero(values):
    data = np.asarray(values)
    low, high, step, _ = _engineering_axis_scale(data, 'Height [m]')
    assert low < high and step > 0
    assert low <= np.nanmin(data) and high >= np.nanmax(data)
    if np.nanmin(data) > 0:
        assert low > 0
    if np.nanmax(data) < 0:
        assert high < 0


@pytest.mark.parametrize('plot_count', [15, 21])
def test_two_column_profile_layout_and_presentation_rules(tmp_path, plot_count):
    workbook = Workbook()
    sheet = workbook.active
    channel_names = ['distance', 'power', 'energy', 'front', 'rear']
    columns = {name: i for i, name in enumerate(channel_names, 1)}
    values = {name: np.array([0., .0123, .066667]) for name in channel_names}
    channels = {name: _channel(name, name.title(), 'km' if name == 'distance' else 'kW')
                for name in channel_names}
    for col, name in enumerate(channel_names, 1):
        for row, value in enumerate(values[name], 5):
            sheet.cell(row, col, value)
    for col in range(10, 25):
        sheet.column_dimensions[get_column_letter(col)].width = 13
    before = {key: cell.value for key, cell in sheet._cells.items()}
    definitions, rendered = [], []
    structures = [('power',), ('power', 'energy'), ('power', 'energy', 'front', 'rear')]
    for i in range(plot_count):
        names = structures[i % 3]
        secondary = ('energy',) if len(names) == 2 else ()
        series = tuple(ProfilePlotSeriesDefinition(name, 'secondary' if name in secondary else 'primary')
                       for name in names)
        definition = ProfilePlotDefinition(f'plot_{i}', f'Engineering plot {i}', 'distance', series,
                                           x_label='Distance [km]', primary_y_label='Power [kW]',
                                           secondary_y_label='Energy [kWh]')
        definitions.append(definition)
        rendered.append(SimpleNamespace(plot_id=definition.plot_id, title=definition.title))
    result = SimpleNamespace(profile=SimpleNamespace(plots=definitions), rendered_plots=rendered,
                             channels_by_semantic_name=channels, values_by_semantic_name=values,
                             sample_count=3)
    _write_profile_plots_on_report_sheet(sheet, result, 6, 10, channel_columns=columns,
        border=Border(), label_fill=PatternFill('solid', fgColor='1F4E78'), white_bold=Font(color='FFFFFF'))
    path = tmp_path / 'layout.xlsx'
    workbook.save(path)
    # openpyxl's chart reader does not restore plot-area spPr; inspect the
    # emitted OOXML rather than mistaking that reader limitation for lost styling.
    ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
          'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with ZipFile(path) as package:
        for i in range(plot_count):
            xml = ElementTree.fromstring(package.read(f'xl/charts/chart{i + 1}.xml'))
            assert xml.find('.//c:plotArea/c:spPr/a:solidFill/a:srgbClr', ns).get('val') == 'FFFFFF'
    workbook = load_workbook(path)
    sheet = workbook.active
    assert not sheet._images
    assert len(sheet._charts) == plot_count
    assert {key: sheet.cell(*key).value for key in before} == before
    assert not sheet.row_dimensions  # Chart layout does not enlarge engineering data rows.
    rows = Counter(chart.anchor._from.row for chart in sheet._charts)
    assert max(rows.values()) == 2
    assert list(rows.values())[-1] == 1
    boxes = []
    for i, chart in enumerate(sheet._charts):
        anchor = chart.anchor._from
        assert chart.anchor.ext.cx / 914400 == pytest.approx(637 / 96)
        assert chart.anchor.ext.cy / 914400 == pytest.approx(3.75)
        assert "".join(r.t for p in chart.title.tx.rich.p for r in p.r) == definitions[i].title
        assert chart.title.txPr.p[0].pPr.defRPr.sz == 1400
        assert chart.title.txPr.p[0].pPr.defRPr.solidFill.srgbClr == "1F4E78"
        left, top = (anchor.col - 9) * 91 * .75, anchor.row * 15
        boxes.append((left, top, left + 637 * .75, top + 3.75 * 72))
        assert sheet.cell(anchor.row + 1, anchor.col + 1).value is None
        count = sum(len(component.series) for component in chart._charts)
        assert (chart.legend is not None) == (count > 1)
        if chart.legend:
            assert chart.legend.position == 'b' and chart.legend.overlay is False
            assert chart.legend.txPr.p[0].pPr.defRPr.sz == 900
        assert chart.graphical_properties.solidFill.srgbClr == 'FFFFFF'
        for axis in (chart.x_axis, chart.y_axis):
            assert axis.majorGridlines.spPr.ln.solidFill.srgbClr == 'DCE3E8'
            assert axis.majorGridlines.spPr.ln.w == 6350
            assert not axis.majorGridlines.spPr.ln.noFill
        for component in chart._charts:
            for axis in (component.x_axis, component.y_axis):
                assert axis.numFmt.formatCode == '0.00'
                assert axis.numFmt.sourceLinked is False
                assert axis.majorUnit > 0
                assert axis.scaling.min <= 0 and axis.scaling.max >= .066667
                assert 5 <= round((axis.scaling.max - axis.scaling.min) / axis.majorUnit) + 1 <= 7
                assert axis.tickLblPos == "nextTo"
                if axis is component.y_axis or component is chart:
                    assert axis.delete is False
                    assert axis.title is not None
                assert axis.txPr.p[0].pPr.defRPr.solidFill.srgbClr == "333333"
                assert axis.txPr.p[0].pPr.defRPr.sz == 1000
            for series in component.series:
                assert series.xVal.numRef.f == "'Sheet'!$A$5:$A$7"
                assert series.graphicalProperties.line.width == 31750
        if len(chart._charts) == 2:
            secondary = chart._charts[1]
            assert secondary.y_axis.axPos == 'r'
            assert secondary.series[0].yVal.numRef.f == "'Sheet'!$C$5:$C$7"
            assert secondary.x_axis.scaling.min == chart.x_axis.scaling.min
            assert secondary.x_axis.scaling.max == chart.x_axis.scaling.max
            assert secondary.y_axis.majorGridlines is None
    for i, (left, top, right, bottom) in enumerate(boxes):
        for other_left, other_top, other_right, other_bottom in boxes[i + 1:]:
            assert right < other_left or other_right < left or bottom < other_top or other_bottom < top



@pytest.mark.parametrize("values,label,zero_expected", [
    ([.02, 14.9], "Speed [kph]", True),
    ([36.4, 40.0], "Battery Energy [kWh]", False),
    ([65, 82], "SOC [%]", False),
    ([-15, 15], "Gradient [%]", True),
    ([-20, 35], "Power [kW]", True),
    ([150, 170], "Height [m]", False),
])
def test_engineering_ranges_preserve_quantity_and_sign(values, label, zero_expected):
    lo, hi, step, fmt = _engineering_axis_scale(np.array(values), label)
    assert lo <= min(values) and hi >= max(values)
    assert (lo <= 0 <= hi) == zero_expected
    assert 5 <= round((hi - lo) / step) + 1 <= 7


def test_visible_axis_flags_are_explicit_in_saved_xml(tmp_path):
    from vsm_postprocessing.excel_report_engine import _build_native_scatter_chart
    from vsm_postprocessing.plotting_engine import PlotDefinition, PlotSeriesDefinition
    workbook = Workbook()
    sheet = workbook.active
    for row in [(0, 0, 40), (.2, 14.9, 38), (1, 5, 36.4)]:
        sheet.append(row)
    definition = PlotDefinition("battery", "Battery Energy - Distance Based", "distance",
        (PlotSeriesDefinition("speed"), PlotSeriesDefinition("energy", "secondary")), "",
        x_label="Distance [km]", primary_y_label="Speed [kph]", secondary_y_label="Energy [kWh]")
    chart = _build_native_scatter_chart(definition, report_sheet=sheet,
        channel_column_by_id={"distance": 1, "speed": 2, "energy": 3},
        data_start_row=1, data_end_row=3, chart_index=0)
    assert chart.y_axis.scaling.min == 0 and chart.y_axis.scaling.max == 15
    assert chart.y_axis.majorUnit == 2.5
    assert chart._charts[1].y_axis.scaling.min > 30
    sheet.add_chart(chart, "E1")
    path = tmp_path / "visible_axes.xlsx"
    workbook.save(path)
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/charts/chart1.xml"))
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        axes = root.findall(".//c:valAx", ns)
        assert len(axes) == 4
        assert sum(a.find("c:delete", ns).get("val") == "1" for a in axes) == 1
        for axis in axes:
            assert axis.find("c:tickLblPos", ns).get("val") == "nextTo"
            assert axis.find("c:majorTickMark", ns).get("val") == "out"
            assert axis.find("c:majorUnit", ns) is not None
            assert axis.find("c:scaling/c:min", ns) is not None
            assert axis.find("c:scaling/c:max", ns) is not None
