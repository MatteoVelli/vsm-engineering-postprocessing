from __future__ import annotations

import csv
import posixpath
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
import numpy as np
from pptx import Presentation
from pptx.util import Inches
import pytest

from conftest import DIESEL_REFERENCE_CSV, DIESEL_REFERENCE_DESCRIPTION, require_private_reference_file
from test_diesel_profile import PROFILE_PATH, PROJECT_ROOT, _source
from vsm_postprocessing.report_profile import load_reporting_profile
from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report


TITLES = [
    "Full Size Sprayer Diesel", "System & Simulation Overview", "Executive Results",
    "Vehicle Operation", "Diesel Engine Performance", "Fuel Consumption & Mechanical Energy",
    "Driveline & Wheel Loads", "Engine Load & Throttle", "Engine Thermal Behaviour", "Simulation Summary",
]
MAJOR = {
    "drive_cycle_minutes": (3.7, "3.70 min"), "distance_km": (0.999242, "1.00 km"),
    "chassis_speed_max": (22.4686, "22.47 kph"), "road_gradient_min": (-15, "-15.00 %"),
    "road_gradient_max": (15, "15.00 %"), "engine_speed_max": (1996.81, "1,997 rpm"),
    "engine_speed_average": (1537.2141644144144, "1,537 rpm"), "engine_torque_max": (826.219, "826.22 Nm"),
    "engine_power_max": (165.93862633964403, "165.94 kW"), "engine_power_min": (-8.230092194143745, "-8.23 kW"),
    "engine_power_rms": (74.57069413573709, "74.57 kW"),
    "engine_mechanical_energy_last": (3.7647817811605133, "3.76 kWh"),
    "fuel_consumed_kg": (0.772836, "0.77 kg"), "fuel_volume_last": (0.9421207194583333, "0.942 L"),
    "fuel_flow_average": (15.277633288513513, "15.28 L/h"), "fuel_flow_max": (41.4741, "41.47 L/h"),
}


@pytest.fixture(scope="module")
def report(tmp_path_factory):
    source = require_private_reference_file(DIESEL_REFERENCE_CSV, DIESEL_REFERENCE_DESCRIPTION)
    return generate_reporting_profile_engineering_report(source, PROFILE_PATH, tmp_path_factory.mktemp("diesel_ppt"))


def _text(slide):
    return "\n".join(s.text for s in slide.shapes if s.has_text_frame)


def _plots(slide):
    return [s for s in slide.shapes if s.shape_type == 13 and s.top > Inches(3)]


def test_combined_report_reopens_with_correct_inventory(report):
    assert report.report_path.exists() and report.presentation_path.exists()
    prs = Presentation(report.presentation_path)
    assert len(prs.slides) == report.slide_count == 10
    assert report.powerpoint_result.plot_count == 10
    for slide, title in zip(prs.slides, TITLES):
        assert title in _text(slide)
    assert [len(_plots(s)) for s in prs.slides] == [0, 0, 0, 1, 2, 2, 2, 2, 1, 0]
    assert report.powerpoint_result.excel_result is report.excel_result
    assert report.report_path.stem == report.presentation_path.stem == DIESEL_REFERENCE_CSV.stem


@pytest.mark.parametrize("metric", list(MAJOR))
def test_major_displayed_values_match_independent_d1_golden_results(report, metric):
    expected, displayed = MAJOR[metric]
    result = report.powerpoint_result.powerpoint_result
    actual = next(s for s in result.statistics_result.statistics if s.statistic_id == metric)
    assert actual.value == pytest.approx(expected, rel=2e-15, abs=1e-14)
    prs = Presentation(report.presentation_path)
    matching_slides = [s for s, definition in zip(prs.slides, result.config.slides) if metric in definition.statistic_ids]
    assert matching_slides
    for slide in matching_slides:
        assert displayed in [shape.text for shape in slide.shapes if shape.has_text_frame]
        assert f"{metric}: {actual.value!r}" in slide.notes_slide.notes_text_frame.text


def test_all_displayed_kpis_retain_canonical_values_and_units(report):
    canonical = {m.metric_id: m for m in report.excel_result.statistics_result.canonical_metrics}
    ppt_result = report.powerpoint_result.powerpoint_result
    for stat in ppt_result.statistics_result.statistics:
        metric = canonical[stat.statistic_id]
        assert stat.value == metric.value and stat.channel_unit == metric.unit
    assert all(s in canonical for slide in ppt_result.config.slides for s in slide.statistic_ids)


def test_diesel_terms_and_existing_plot_bytes_with_no_crop(report):
    prs = Presentation(report.presentation_path)
    text = "\n".join(_text(s) for s in prs.slides)
    for forbidden in ("battery", "soc", "edu", "range extender", "fuel economy", "hybrid", "electric"):
        assert forbidden not in text.lower()
    assert "Powertrain Diesel" in text and "Engine Load" in text
    assert "825.46 Nm" in text and "15.28 L/h" in text
    assert "Road height unavailable" in text and "identical in this run" in text
    plots = {p.plot_id: p for p in report.excel_result.plotting_result.rendered_plots}
    for slide, definition in zip(prs.slides, report.powerpoint_result.powerpoint_result.config.slides):
        pictures = _plots(slide)
        assert len(pictures) == len(definition.plot_ids)
        for picture, plot_id in zip(pictures, definition.plot_ids):
            assert picture.image.blob == Path(plots[plot_id].png_file).read_bytes()
            assert (picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom) == (0, 0, 0, 0)
    values = report.excel_result.plotting_result.values_by_semantic_name
    assert min(values["engine_power"]) == pytest.approx(-8.230092194143745)
    assert np.count_nonzero(values["engine_power"] < 0) == 12


def test_template_style_geometry_branding_and_relationships(report):
    prs = Presentation(report.presentation_path)
    template = Presentation(PROJECT_ROOT / "reference_files/RoboSprayer_Electric_Report_Astauto_Colours.pptx")
    logo = (PROJECT_ROOT / "reference_files/astauto-light-text_web.jpg").read_bytes()
    assert (prs.slide_width, prs.slide_height) == (template.slide_width, template.slide_height)
    for n, (slide, original) in enumerate(zip(prs.slides, template.slides), 1):
        assert etree.tostring(slide.element.cSld.bg) == etree.tostring(original.element.cSld.bg)
        title_index = 3 if n == 1 else 2
        actual = next(s for s in slide.shapes if s.has_text_frame and s.text == TITLES[n - 1])
        expected = original.shapes[title_index]
        assert (actual.left, actual.top, actual.width, actual.height) == (expected.left, expected.top, expected.width, expected.height)
        assert actual.text_frame.paragraphs[0].runs[0].font.name == "Cambria"
        assert actual.text_frame.paragraphs[0].runs[0].font.size == expected.text_frame.paragraphs[0].runs[0].font.size
        assert f"{n}  /  10" in _text(slide)
        assert "POST-PROCESSING TOOL" in _text(slide)
        assert len([s for s in slide.shapes if s.shape_type == 13 and s.image.blob == logo]) == 1
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + Inches(.01)
            assert shape.top + shape.height <= prs.slide_height + Inches(.01)
        for plot in _plots(slide):
            assert plot.top >= Inches(3.15) and plot.top + plot.height < Inches(6.9)
    with ZipFile(report.presentation_path) as package:
        assert package.testzip() is None
        for name in package.namelist():
            if not name.endswith(".rels"):
                continue
            folder = posixpath.dirname(posixpath.dirname(name))
            for rel in etree.fromstring(package.read(name)):
                if rel.get("TargetMode") != "External":
                    target = rel.get("Target")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
                    assert resolved in package.namelist(), (name, target)


@pytest.mark.parametrize("missing,count,plots", [
    ("all", 7, 4), ("fuel_flow", 10, 9), ("engine_oil_temperature", 9, 9),
    ("engine_load", 10, 9), ("shafts", 10, 9), ("wheels", 10, 9),
    ("driveshaft_torque_fl", 10, 10),
])
def test_missing_optional_data_removes_only_dependent_content(tmp_path, missing, count, plots):
    profile = load_reporting_profile(PROFILE_PATH)
    omitted = {
        "all": [c.semantic_name for c in profile.raw_channels if not c.required],
        "shafts": ["driveshaft_torque_" + c for c in ("fl", "fr", "rl", "rr")],
        "wheels": ["tyre_fz_" + c for c in ("fl", "fr", "rl", "rr")],
    }.get(missing, [missing])
    result = generate_reporting_profile_engineering_report(_source(tmp_path, profile, omitted=omitted), PROFILE_PATH, tmp_path / "report")
    prs = Presentation(result.presentation_path)
    assert len(prs.slides) == count and result.powerpoint_result.plot_count == plots
    assert sum(len(_plots(s)) for s in prs.slides) == plots
    text = "\n".join(_text(s) for s in prs.slides)
    assert "n/a" not in text.lower() and "battery" not in text.lower()
    if missing in ("all", "fuel_flow"):
        assert "AVG FUEL FLOW" not in text and "INTEGRATED FUEL VOLUME" not in text
    if missing in ("all", "engine_oil_temperature"):
        assert "Engine Thermal Behaviour" not in text
    for n, slide in enumerate(prs.slides, 1):
        assert f"{n}  /  {count}" in _text(slide)
        pictures = _plots(slide)
        if len(pictures) == 1:
            assert pictures[0].left + pictures[0].width / 2 == pytest.approx(prs.slide_width / 2, abs=1)


def test_all_zero_optional_plots_are_omitted_without_changing_excel(tmp_path):
    profile = load_reporting_profile(PROFILE_PATH)
    source = _source(tmp_path, profile)
    rows = list(csv.reader(source.open(encoding="utf-8")))
    optional = {c.source_name for c in profile.raw_channels if not c.required}
    for col, name in enumerate(rows[0]):
        if name in optional:
            for row in rows[2:]:
                row[col] = "0"
    with source.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows(rows)
    result = generate_reporting_profile_engineering_report(source, PROFILE_PATH, tmp_path / "report")
    assert result.excel_result.plot_count == 10
    assert result.slide_count == 7 and result.powerpoint_result.plot_count == 4
    assert result.excel_result.report_channel_count == 28


def test_synthetic_shifted_counters_and_time_use_d1_cycle_deltas(tmp_path):
    profile = load_reporting_profile(PROFILE_PATH)
    result = generate_reporting_profile_engineering_report(_source(tmp_path, profile), PROFILE_PATH, tmp_path / "report")
    cover = _text(Presentation(result.presentation_path).slides[0])
    assert "0.30 kg" in cover and "10.30 kg" not in cover
    assert "0.05 min" in cover and "0.22 min" not in cover
    values = result.excel_result.plotting_result.values_by_semantic_name
    assert values["engine_power"][-1] < 0
