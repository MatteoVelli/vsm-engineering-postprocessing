"""Live KPI dependency chains and independent Diesel thermal measurements."""
import numpy as np
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
import pytest

from test_sergio_october_15 import dataset_for, profile_for, profile_path, csv_for, remove_channel
from vsm_postprocessing.excel_formulas import save_report_workbook
from vsm_postprocessing.profile_kpi_formulas import write_profile_kpi_formulas
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.report_selection import ProfileReportSelection
from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report


@pytest.mark.parametrize("kind", ("diesel", "electric", "hybrid"))
@pytest.mark.parametrize("count", (3, 17))
@pytest.mark.parametrize("export_raw", (True, False))
def test_all_available_kpis_have_live_dependencies_and_canonical_caches(tmp_path, kind, count, export_raw):
    data = dataset_for(kind, np.linspace(17, 3617, count))
    stats = calculate_profile_statistics(data, profile_for(kind))
    book = Workbook()
    main = book.active
    main.title = "Selected report's data"
    book.create_sheet("Statistics")
    exported = data.channels if export_raw else []
    for col, channel in enumerate(exported, 1):
        for row, value in enumerate(data.values[:, data.channel_index(channel.channel_id)], 5):
            main.cell(row, col, float(value))
    write_profile_kpi_formulas(book, main, stats, exported)
    path = tmp_path / "live.xlsx"
    save_report_workbook(book, path)
    formulas, caches = load_workbook(path), load_workbook(path, data_only=True)
    try:
        start = len(stats.statistics) + 5
        assert len(stats.kpis) == (5 if kind == "diesel" else 12)
        for row, kpi in enumerate(stats.kpis, start):
            cell = formulas["Statistics"].cell(row, 3)
            assert cell.data_type == "f"
            assert "!" in cell.value and "#REF!" not in cell.value
            assert caches["Statistics"].cell(row, 3).value == pytest.approx(kpi.value, rel=1e-14)
        helpers = formulas["Formula Support"]
        assert helpers.sheet_state == "hidden"
        assert any(c.data_type == "f" for row in helpers for c in row)
        text = "\n".join(str(c.value) for sheet in formulas for row in sheet for c in row if c.value is not None)
        assert "$" + str(count+4) in text  # current data extent, including a nonzero time origin
        assert "17417" not in text
        if kind != "diesel":
            assert "MEDIAN(" in text and "minimum valid SOC fraction" in text
            assert "literal 0.85" in text  # documented assumption, referenced from a cell
        else:
            assert "signed trapezoidal integral" in text
    finally:
        formulas.close(); caches.close()


@pytest.mark.parametrize("missing", ("none", "auxiliary_high_voltage_power_consumption", "track_height"))
def test_missing_optional_kpis_are_unavailable_not_zero_or_broken_references(missing):
    data = dataset_for("electric")
    if missing != "none":
        remove_channel(data, missing)
    stats = calculate_profile_statistics(data, profile_for("electric"))
    book = Workbook()
    book.create_sheet("Statistics")
    write_profile_kpi_formulas(book, book.active, stats, [])
    names = {k.definition.kpi_id for k in stats.kpis}
    assert ("total_energy_consumption_tyres_aux" in names) == (missing != "auxiliary_high_voltage_power_consumption")
    assert ("road_height_range" in names) == (missing != "track_height")
    assert all("#REF!" not in str(c.value) for s in book for row in s for c in row)


def test_kpi_formula_chain_survives_a_minimal_export_selection(tmp_path):
    selection = ProfileReportSelection(export_channels=("track_time",), statistics=(),
                  kpis=("range_85_battery_km",), plots=(), powerpoint_statistics=())
    report = generate_profile_excel_report(csv_for(dataset_for("electric"), tmp_path/"selected.csv"),
                  profile_path("electric"), tmp_path/"report", selection=selection)
    assert [c.channel_id for c in report.report_channels] == ["track_time"]
    book = load_workbook(report.report_path)
    try:
        for row, item in enumerate(report.statistics_result.kpis, report.statistic_count+5):
            assert book["Statistics"].cell(row,3).data_type == "f"
        labels = [str(c.value) for c in book["Formula Support"][3] if c.value]
        assert any("raw" in label and "Battery" in label for label in labels)
        assert any("track_distance" in label for label in labels)
    finally:
        book.close()


@pytest.mark.parametrize("mode", ("identical", "different", "missing_oil", "missing_water", "missing_both", "zero"))
def test_current_diesel_thermal_charts_are_independent(tmp_path, monkeypatch, mode):
    from vsm_postprocessing import profile_plotting
    data = dataset_for("diesel", (17, 18, 22, 30))
    oil = np.array([60., 65., 80., 70.])
    water = np.array([30., 35., 40., 45.]) if mode == "different" else oil.copy()
    if mode == "zero":
        oil[:] = 0; water[:] = 0
    data.values[:, data.channel_index("engine_oil_temperature")] = oil
    data.values[:, data.channel_index("engine_water_temperature")] = water
    missing = {"missing_oil": "engine_oil_temperature", "missing_water": "engine_water_temperature"}.get(mode)
    if missing:
        remove_channel(data, missing)
    if mode == "missing_both":
        for name in ("engine_oil_temperature", "engine_water_temperature"):
            remove_channel(data, name)
    captured = {}
    original = profile_plotting._render_one_plot

    def capture(definition, defaults, destination, channels, values, count):
        if definition.plot_id.startswith("diesel_engine_") and definition.plot_id.endswith("temperature"):
            captured[definition.plot_id] = (definition, values[definition.series[0].channel_id].copy())
        return original(definition, defaults, destination, channels, values, count)

    monkeypatch.setattr(profile_plotting, "_render_one_plot", capture)
    report = generate_reporting_profile_engineering_report(csv_for(data, tmp_path / "measured.csv"),
                      profile_path("diesel"), tmp_path / "report")
    prs = Presentation(report.presentation_path)
    assert len(prs.slides) == 11
    assert report.powerpoint_result.powerpoint_result.config.slides[-2].slide_id == "torque_maps"
    thermal = prs.slides[8]
    plots = [s for s in thermal.shapes if s.shape_type == 13 and s.top > Inches(3)]
    assert len(plots) == (0 if mode == "missing_both" else 1 if missing else 2)
    text = "\n".join(s.text for s in thermal.shapes if s.has_text_frame)
    stats = {s.definition.statistic_id:s.value for s in report.excel_result.statistics_result.statistics}
    if mode == "missing_both":
        assert "Engine Oil Temperature; Engine Water Temperature measurement unavailable" in text
        assert "Engine temperature measurements unavailable for this dataset" in text
        assert not captured
        return
    if missing:
        assert "Available measured engine temperature; missing sources are identified below" in text
    for key, expected, source in (("oil", oil, "Engine_Oil_Temperature"), ("water", water, "Engine_Water_Temperature")):
        name = f"engine_{key}_temperature"
        if name == missing:
            assert f"Engine {key.title()} Temperature measurement unavailable" in text
            assert f"diesel_{name}" not in captured
            continue
        definition, values = captured[f"diesel_{name}"]
        np.testing.assert_array_equal(values, expected)
        assert definition.x_channel_id == "track_time" and definition.x_label == "Time [s]"
        assert definition.primary_y_label == "Temperature [°C]"
        assert len(definition.series) == 1 and definition.series[0].channel_id == name
        assert profile_for("diesel").raw_by_semantic_name()[name].source_name == source
        assert stats[name+"_min"] == min(expected) and stats[name+"_max"] == max(expected)
    if not missing:
        assert plots[0].width == plots[1].width and plots[0].height == plots[1].height
        assert plots[0].top == plots[1].top and plots[0].left + plots[0].width < plots[1].left
        definitions = [v[0] for v in captured.values()]
        assert definitions[0].primary_y_limits == definitions[1].primary_y_limits
        assert definitions[0].series[0].colour != definitions[1].series[0].colour
        config = report.powerpoint_result.powerpoint_result.config.slides[8]
        assert len(config.statistic_ids) == 4
