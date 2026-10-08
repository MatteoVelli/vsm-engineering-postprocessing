"""Diesel selects six RMS statistics and exposes only two as live top results."""
from __future__ import annotations

import math

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from test_diesel_profile import PROFILE_PATH, _source
from workbook_assertions import assert_universal_bottom_statistics
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.report_profile import load_reporting_profile


HEADLINES = {"engine_torque_rms": ("engine_torque", "Engine Torque RMS", "Nm"),
             "engine_power_rms": ("engine_power", "Engine Power RMS", "kW")}
SUPPORT = {f"driveshaft_torque_{corner}_rms" for corner in ("fl", "fr", "rl", "rr")}
REMOVED = {"engine_speed_rms", *(f"chassis_acceleration_{axis}_rms" for axis in ("longitudinal", "lateral", "vertical"))}


@pytest.fixture(scope="module", params=[False, True], ids=["all_channels", "required_channels"])
def report(request, tmp_path_factory):
    root = tmp_path_factory.mktemp("diesel_rms")
    profile = load_reporting_profile(PROFILE_PATH)
    omitted = [c.semantic_name for c in profile.raw_channels if not c.required] if request.param else []
    return generate_profile_excel_report(_source(root, profile, omitted=omitted), PROFILE_PATH, root / "report")


def test_profile_selects_only_six_rms_metrics_and_preserves_raw_inputs():
    profile = load_reporting_profile(PROFILE_PATH)
    definitions = {s.statistic_id: s for s in profile.statistics}
    assert {sid for sid, s in definitions.items() if s.operation == "rms"} == HEADLINES.keys() | SUPPORT
    assert not REMOVED & definitions.keys()
    assert {sid for sid, s in definitions.items() if s.placement_group == "top_rms"} == HEADLINES.keys()
    assert all(definitions[sid].placement_group != "top_rms" for sid in SUPPORT)
    assert {"engine_speed", *(f"chassis_acceleration_{axis}" for axis in ("longitudinal", "lateral", "vertical"))} <= profile.raw_by_semantic_name().keys()


@pytest.mark.parametrize("metric_id", HEADLINES)
def test_headline_is_live_sample_rms_in_shared_top_block(report, metric_id):
    formulas = load_workbook(report.report_path, data_only=False)
    cached = load_workbook(report.report_path, data_only=True)
    try:
        sheet, numeric = formulas.worksheets[0], cached.worksheets[0]
        target, label, unit = HEADLINES[metric_id]
        columns = {c.channel_id: i for i, c in enumerate(report.report_channels, 1)}
        col = columns[target]
        heading = next(cell for cell in sheet[1] if cell.value == label)
        block = next(r for r in sheet.merged_cells.ranges if heading.coordinate in r)
        assert block.min_row == block.max_row == 1 and block.min_col <= col <= block.max_col
        # The shared Electric writer moves overlapping blocks as a unit when
        # optional columns are absent. The formula still targets the source.
        value = next(sheet.cell(2, c) for c in range(block.min_col, block.max_col + 1)
                     if sheet.cell(2, c).data_type == "f")
        letter = get_column_letter(col)
        extent = f"{letter}5:{letter}{report.sample_count + 4}"
        assert value.data_type == "f"
        assert value.value == f"=SQRT(SUMSQ({extent})/COUNT({extent}))"
        assert heading.fill.fgColor.rgb.endswith("143642")
        assert value._style.fillId == heading._style.fillId
        assert value._style.fontId == heading._style.fontId
        assert value.alignment.horizontal == "center"
        assert value.number_format == f'0.000" {unit}"'
        # Evaluate the source range independently: sample RMS, including signed
        # samples and the zero-speed power sample; no elapsed-time weighting.
        samples = [numeric.cell(row, col).value for row in range(5, report.sample_count + 5)]
        expected = math.sqrt(math.fsum(x * x for x in samples) / len(samples))
        backend = next(s for s in report.statistics_result.statistics if s.definition.statistic_id == metric_id)
        assert numeric[value.coordinate].value == pytest.approx(expected, rel=2e-15)
        assert backend.value == pytest.approx(expected, rel=2e-15)
        assert backend.channel_unit == unit == sheet.cell(4, col).value
        assert formulas.calculation.fullCalcOnLoad and formulas.calculation.forceFullCalc
    finally:
        formulas.close()
        cached.close()


def test_main_has_no_banners_or_bottom_rms_and_summary_closes_gap(report):
    book = load_workbook(report.report_path, data_only=False)
    try:
        sheet = book.worksheets[0]
        text = [cell.value for row in sheet for cell in row if isinstance(cell.value, str)]
        assert not any("Diesel Executive Results" in value or "Complete results: Statistics." in value for value in text)
        assert "RMS" not in text
        end = report.sample_count + 4
        assert [sheet.cell(end + i, 1).value for i in range(1, 5)] == ["MAX", "MIN", "LAST", "FIRST"]
        assert all(sheet.cell(end + 5, col).value is None for col in range(1, report.report_channel_count + 2))
        assert_universal_bottom_statistics(report.report_path, report.sample_count)
        start = report.report_channel_count + 2
        assert sheet.cell(3, start).value == "Maximum Vehicle Speed [kph]"
        assert sheet.cell(4, start).value.startswith("=")
        assert sheet.cell(1, start).value is None
        assert sheet.cell(2, start).value is None
        assert not any(r.min_col >= start for r in sheet.merged_cells.ranges)
        assert sheet.row_dimensions[1].height == 60
        assert sheet.cell(2, 1).alignment.wrap_text
        assert sheet._charts[0].anchor._from.row == 7
        for chart in sheet._charts:
            assert chart.x_axis.majorGridlines is not None
            assert chart.x_axis.majorGridlines.spPr.ln.w == 6350
        assert not any("Drive" in value and "RMS" in value for value in text)
    finally:
        book.close()


def test_supporting_rms_remains_in_statistics_and_backend(report):
    book = load_workbook(report.report_path, data_only=True)
    try:
        rows = {row[0]: row for row in book["Statistics"].iter_rows(values_only=True)}
        backend = {s.definition.statistic_id: s for s in report.statistics_result.statistics}
        assert not REMOVED & rows.keys()
        expected = SUPPORT if report.vsm_count > 7 else set()
        assert SUPPORT & rows.keys() == SUPPORT & backend.keys() == expected
        for mid in expected:
            item = backend[mid]
            series = report.plotting_result.values_by_semantic_name[item.target_channel]
            independent = math.sqrt(math.fsum(x * x for x in series) / len(series))
            assert rows[mid][2] == "rms" and rows[mid][5] == "Nm"
            assert rows[mid][4] == pytest.approx(independent, rel=2e-15)
            assert item.value == pytest.approx(independent, rel=2e-15)
    finally:
        book.close()


@pytest.mark.parametrize("powertrain", ["electric", "hybrid"])
def test_electric_and_hybrid_keep_established_rms_selection_and_helpers(powertrain):
    profile = load_reporting_profile(PROFILE_PATH.with_name(f"robosprayer_{powertrain}.yaml"))
    rms = {s.statistic_id: s for s in profile.statistics if s.operation == "rms"}
    assert set(rms) == {"battery_power_rms", "battery_heatflow_rms"}
    assert all(s.placement_group == "top_rms" and s.unit == "kW" for s in rms.values())
    helpers = {c.semantic_name: c for c in profile.math_channels if c.semantic_name in {"battery_power_squared", "battery_heatflow_squared"}}
    assert len(helpers) == 2
    assert all(c.unit == "kW^2" for c in helpers.values())
