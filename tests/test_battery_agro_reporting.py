from dataclasses import asdict

import numpy as np
import pytest

from vsm_postprocessing.errors import ConfigurationError
from vsm_postprocessing.reporting_inputs import ReportingInputs
from vsm_postprocessing.agro_reporting import calculate_agro_kpis
from vsm_postprocessing.report_metadata import resolve_report_metadata


def agro(mass=(500, 250, 0), distance=(0, 500, 1000), **kwargs):
    options = dict(density_kg_per_L=1.0, density_provenance='Hypothetical unit test', hitch_force_is_remaining_agro_mass=True,
                   agro_mass_source_policy='force_inference', agro_mass_source='Test force',
                   agro_mass_source_provenance='Explicit synthetic source mapping')
    options.update(kwargs.pop('options', {}))
    return calculate_agro_kpis(
        np.asarray(distance, dtype=float), 'm',
        force_N=-9.81 * np.asarray(mass, dtype=float), force_unit='N',
        inputs=ReportingInputs(**options), **kwargs)


def test_explicit_density_mass_drop():
    result = agro()
    assert result.rate.value == pytest.approx(.5)
    assert result.capacity.value == 500
    assert result.trace['mass_rate_kg_per_m'] == pytest.approx(.5)
    assert result.rate.unit == 'L/m'


def test_missing_density_and_force_interpretation():
    result = agro(options={'density_kg_per_L': None})
    assert result.rate.value is None
    assert 'density required' in result.rate.reason
    assert agro(options={'hitch_force_is_remaining_agro_mass': False}).rate.value == .5


@pytest.mark.parametrize('mass,distance,reason', [
    ([500, 0, -100], [0, 500, 1000], 'negative'),
    ([0, 0, 0], [0, 500, 1000], 'all-zero'),
    ([500, 300, 400], [0, 500, 1000], 'refill'),
    ([500, np.nan, 0], [0, 500, 1000], 'nonfinite'),
    ([500, np.inf, 0], [0, 500, 1000], 'nonfinite'),
    ([500, 250, 0], [0, 0, 0], 'positive'),
    ([500, 250, 0], [0, 500, -1], 'reversed'),
    ([500, 250, 0], [0, 500, 499], 'reversed'),
    ([500, 250, 0], [0, np.nan, 1000], 'nonfinite'),
])
def test_bad_traces_are_unavailable(mass, distance, reason):
    result = agro(mass, distance)
    assert result.rate.value is None
    assert reason in result.rate.reason


def test_units_missing_data_and_direct_volume_precedence():
    result = calculate_agro_kpis([0, 1], 'km', volume_L=[10, 510], volume_unit='L',
                                force_N=[0, 100000], inputs=ReportingInputs(agro_volume_source='Sprayed', agro_volume_provenance='validated cumulative volume'))
    assert result.rate.value == .5
    assert result.trace['method'] == 'direct cumulative volume'
    assert calculate_agro_kpis([0, 1], 'ft', inputs=ReportingInputs()).rate.value is None
    assert calculate_agro_kpis([0, 1], 'm', inputs=ReportingInputs()).rate.value is None
    assert calculate_agro_kpis([0, 1], 'm', volume_L=[0, 5], volume_unit='kg', inputs=ReportingInputs(agro_volume_source='v', agro_volume_provenance='trusted')).rate.value is None
    assert calculate_agro_kpis([0, 1e308], 'km', volume_L=[0, 5], volume_unit='L', inputs=ReportingInputs(agro_volume_source='v', agro_volume_provenance='trusted')).rate.value is None


@pytest.mark.parametrize('key', ['battery_mass_kg', 'density_kg_per_L'])
@pytest.mark.parametrize('value', [0, -1, np.nan, np.inf, True])
def test_positive_finite_inputs(key, value):
    with pytest.raises(ConfigurationError):
        ReportingInputs.from_mapping({key: value})


def test_pack_identity_and_provisional_status():
    from vsm_postprocessing.report_profile import load_reporting_profile
    profile = load_reporting_profile('config/report_profiles/robosprayer_electric.yaml')
    known = resolve_report_metadata('RoboSprayer_3500Kg_Electric.csv', profile)
    assert known.reporting_inputs.battery_mass_kg is None
    assert known.reporting_inputs.battery_mass_status == 'not_provided'
    assert resolve_report_metadata('Unrelated_3500Kg_Electric.csv', profile).reporting_inputs.battery_mass_kg is None
    assert resolve_report_metadata('RoboSprayer_3500Kg_Diesel.csv', profile).reporting_inputs.battery_mass_kg is None
    assert resolve_report_metadata('Sprayer_Caiman_SP_9300Kg_Hybrid.csv').reporting_inputs.battery_mass_kg is None
    provisional = ReportingInputs.from_mapping({'battery_mass_kg': 1700, 'battery_mass_status': 'estimated', 'battery_mass_provenance': 'Sergio preliminary; calculation in progress'})
    assert provisional.battery_display == '~1,700 kg (estimated)'
    assert ReportingInputs().battery_display == 'Not specified'
    assert asdict(provisional)['battery_mass_status'] == 'estimated'
    with pytest.raises(ConfigurationError):
        resolve_report_metadata('RoboSprayer_3500Kg_Electric.csv', profile, machine_name_override='Caiman SP', reporting_inputs={'machine_configuration_id': 'robosprayer_v5_3500kg'})


@pytest.mark.parametrize('fixture_name,profile_id,expected,reason', [
    ('ROBOSPRAYER_LATEST_ELECTRIC_CSV', 'robosprayer_electric', .499832, ''),
    ('ROBOSPRAYER_LATEST_HYBRID_CSV', 'robosprayer_hybrid', None, 'negative'),
    ('ROBOSPRAYER_REFERENCE_CSV', 'robosprayer_electric', None, 'all-zero'),
])
def test_actual_sources(fixture_name, profile_id, expected, reason):
    import conftest
    from vsm_postprocessing.importer import load_data_file
    from vsm_postprocessing.report_profile import load_reporting_profile
    from vsm_postprocessing.profile_statistics import calculate_profile_statistics, build_canonical_report_metrics
    from vsm_postprocessing.agro_reporting import calculate_reporting_kpis
    source = conftest.require_private_reference_file(getattr(conftest, fixture_name), fixture_name)
    profile = load_reporting_profile('config/report_profiles/' + profile_id + '.yaml')
    stats = calculate_profile_statistics(load_data_file(source), profile)
    before = [(m.metric_id, m.value) for m in build_canonical_report_metrics(stats)]
    settings = {'agro_density_kg_per_L': 1.0, 'density_provenance': 'Hypothetical unit test only', 'hitch_force_is_remaining_agro_mass': True}
    metadata = resolve_report_metadata(source, profile, reporting_inputs=settings)
    result = calculate_reporting_kpis(stats, metadata)
    if expected is not None:
        assert result.rate.value == pytest.approx(expected, abs=1e-6)
        settings['agro_density_kg_per_L'] = None
        assert calculate_reporting_kpis(stats, resolve_report_metadata(source, profile, reporting_inputs=settings)).rate.value is None
    else:
        assert result.rate.value is None
        assert reason in result.rate.reason
    assert [(m.metric_id, m.value) for m in build_canonical_report_metrics(stats)] == before


def test_direct_flow_nonzero_start_and_discontinuities():
    options = ReportingInputs(agro_flow_source='SprayFlow', agro_flow_provenance='validated L/min')
    result = calculate_agro_kpis([200, 1200], 'm', inputs=options, flow=[30, 30], flow_unit='L/min', time_s=[0, 1000])
    assert result.rate.value == .5
    assert agro(distance=[200, 700, 1200]).rate.value == .5
    result = agro(mass=[500, 499, 498, 497, 496, 495, 0], distance=list(range(7)))
    assert result.rate.value is None and 'discontinuity' in result.rate.reason
    result = calculate_agro_kpis([0, 1, 2], 'm', inputs=ReportingInputs(agro_volume_source='V', agro_volume_provenance='validated'), volume_L=[0, 10, 1], volume_unit='L')
    assert result.rate.value is None


def test_profile_input_inheritance_and_serializable_config(tmp_path):
    import yaml
    from vsm_postprocessing.report_profile import load_reporting_profile
    root = __import__('pathlib').Path('config/report_profiles/robosprayer_electric.yaml').resolve()
    parent = tmp_path / 'parent.yaml'
    parent.write_text(yaml.safe_dump({'version': 1, 'profile': {
        'profile_id': 'metadata_parent', 'name': 'Electric', 'powertrain': 'electric', 'extends': str(root),
        'reporting_inputs': {'agro_density_kg_per_L': .9, 'density_provenance': 'validated test input'}}}), encoding='utf-8')
    child = tmp_path / 'child.yaml'
    child.write_text(yaml.safe_dump({'version': 1, 'profile': {'profile_id': 'metadata_child', 'name': 'Child',
        'powertrain': 'electric', 'extends': str(parent)}}), encoding='utf-8')
    profile = load_reporting_profile(child)
    assert profile.metadata.reporting_inputs.density_kg_per_L == .9
    metadata = resolve_report_metadata('RoboSprayer_Electric.csv', profile)
    assert metadata.reporting_inputs.density_provenance == 'validated test input'


@pytest.mark.parametrize('kind,density', [('electric', 1.0), ('hybrid', None)])
def test_new_metrics_reach_native_reports_and_manifests(tmp_path, kind, density):
    import json
    from openpyxl import load_workbook
    from pptx import Presentation
    from test_sergio_october_15 import dataset_for, csv_for, profile_path
    from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report
    from vsm_postprocessing.report_selection import ProfileReportSelection
    data = dataset_for(kind, (0, 500, 1000))
    data.values[:, data.channel_index('track_distance')] = [0, 500, 1000]
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81 * np.array([500, 250, 0])
    source = csv_for(data, tmp_path / 'Full_Size_Sprayer_8500Kg.csv')
    settings = {'battery_mass_kg': 1700, 'battery_mass_status': 'estimated',
                'battery_mass_provenance': 'Sergio preliminary; calculation in progress',
                'agro_density_kg_per_L': density, 'density_provenance': 'Explicit hypothetical unit-test density' if density else None,
                'hitch_force_is_remaining_agro_mass': True}
    report = generate_reporting_profile_engineering_report(source, profile_path(kind), tmp_path / 'report',
        reporting_inputs=settings, machine_name_override='Full Size Sprayer',
        selection=ProfileReportSelection(plots=(), math_channels=(), statistics=(), kpis=()))
    prs = Presentation(report.presentation_path)
    text = '\n'.join(s.text for s in prs.slides[7].shapes if s.has_text_frame)
    assert 'AGRO CAPACITY' in text and 'AGRO DISCHARGE RATE' in text
    assert '500.00 kg/km' in text and 'L/m' not in text
    assert any(s.has_text_frame and '1,700 kg' in s.text for s in prs.slides[4].shapes)
    assert prs.core_properties.title == 'Full Size Sprayer ' + kind.title()
    payload = json.loads(report.manifest_path.read_text())
    assert payload['reporting_inputs']['battery_mass_status'] == 'estimated'
    assert payload['agro_reporting']['trace']['density_kg_per_L'] == density
    book = load_workbook(report.report_path, data_only=True)
    try:
        rows = {row[0]: row for row in book['Statistics'].iter_rows(values_only=True) if row[0]}
        assert rows['battery_mass_kg'][2] == '1,700 kg'
        assert rows['agro_discharge_rate_L_per_m'][2] == (.5 if density else 'N/A')
        assert rows['agro_discharge_rate_L_per_m'][3] == 'L/m'
        assert rows['agro_capacity_observed_kg'][2] == 500
    finally:
        book.close()


def test_reporting_widgets_roundtrip_and_density_has_no_default(monkeypatch):
    from streamlit.testing.v1 import AppTest
    from vsm_postprocessing import ui_app
    from vsm_postprocessing.ui_config import get_reporting_profile_definition
    definition = get_reporting_profile_definition(ui_app.PROJECT_ROOT, 'robosprayer_electric')
    monkeypatch.setattr(ui_app, '_metadata_test_definition', definition, raising=False)
    app = AppTest.from_string('''
import streamlit as st
from vsm_postprocessing import ui_app
st.session_state['reporting_inputs_test'] = ui_app._render_reporting_metadata_inputs(st, ui_app._metadata_test_definition, 'RoboSprayer_3500Kg_Electric.csv')
''').run()
    assert not app.exception
    assert app.session_state['reporting_inputs_test'] == {}
    assert not app.selectbox
    next(w for w in app.checkbox if w.label == 'Override automatically detected reporting information').check().run()
    assert not app.exception
    for _ in range(2):
        app.run()
        assert app.session_state['reporting_inputs_test'] == {}
    next(w for w in app.checkbox if w.label == 'Specify agrochemical density [kg/L]').check().run()
    assert app.session_state['reporting_inputs_test'].get('density_kg_per_L') is None
