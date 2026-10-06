"""Release acceptance checks for Diesel input isolation and portable reporting."""
from __future__ import annotations

import os
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
import streamlit as st
from streamlit.testing.v1 import AppTest

from test_diesel_profile import PROFILE_PATH, PROJECT_ROOT, _source
from vsm_postprocessing import ui_app
from vsm_postprocessing.errors import DataValidationError
from vsm_postprocessing.release_builder import build_client_release
from vsm_postprocessing.report_profile import load_reporting_profile


@pytest.fixture
def app(tmp_path, monkeypatch):
    profile = load_reporting_profile(PROFILE_PATH)
    path = _source(tmp_path, profile)
    upload = SimpleNamespace(name="Diesel.csv", getvalue=path.read_bytes)
    state = {"upload": upload}
    # AppTest in the supported Streamlit 1.49 lacks a file-upload driver.
    # Only that widget is supplied here; main(), validation, and report generation run normally.
    monkeypatch.setattr(st, "file_uploader", lambda *args, **kwargs: state["upload"])
    monkeypatch.setattr(ui_app, "UI_WORKSPACE", tmp_path / "uploads")
    monkeypatch.setattr(ui_app, "UI_RUNS", tmp_path / "runs")
    at = AppTest.from_string("from vsm_postprocessing.ui_app import main\nmain()", default_timeout=60).run()
    at.selectbox[0].set_value("full_size_sprayer_diesel").run()
    assert not at.exception
    return at, state


@pytest.mark.parametrize("target", ["robosprayer_electric", "robosprayer_hybrid"])
def test_profile_switch_invalidates_diesel_validation(app, target):
    at, _ = app
    at.button[0].click().run()
    assert not at.button[1].disabled
    at.selectbox[0].set_value(target).run()
    assert at.button[1].disabled
    assert "profile_validation_summary" not in at.session_state
    at.selectbox[0].set_value("full_size_sprayer_diesel").run()
    assert at.button[1].disabled and not at.exception


@pytest.mark.parametrize("replacement", ["changed", "removed", "malformed", "missing_required"])
def test_upload_change_clears_validation_and_both_download_states(app, replacement):
    at, state = app
    at.button[0].click().run()
    at.session_state["profile_report_result"] = "old Electric result"
    at.session_state["diesel_report_result"] = "old Diesel result"
    data = state["upload"].getvalue()
    if replacement == "removed":
        state["upload"] = None
    else:
        data = {"changed": data.replace(b"10.1", b"10.2"),
                "malformed": b"not a valid CSV",
                "missing_required": data.replace(b"Engine_Torque", b"Unknown_Torque")}[replacement]
        state["upload"] = SimpleNamespace(name="Diesel.csv", getvalue=lambda: data)
    at.run()
    assert not at.exception and at.button[1].disabled
    for key in ("profile_validation_summary", "profile_report_result", "diesel_report_result"):
        assert key not in at.session_state
    assert not at.get("download_button")
    if replacement == "missing_required":
        at.button[0].click().run()
        assert not at.session_state["profile_validation_summary"].is_valid
        assert any("Engine_Torque" in message.value for message in at.warning)


def test_failed_revalidation_cannot_reuse_previous_success(app, monkeypatch):
    at, _ = app
    at.button[0].click().run()
    assert not at.button[1].disabled
    def fail(*args, **kwargs):
        raise DataValidationError("Diesel channel contains non-finite values")
    monkeypatch.setattr(ui_app, "validate_reporting_profile_source", fail)
    at.button[0].click().run()
    assert at.button[1].disabled and not at.exception
    assert "profile_validation_summary" not in at.session_state


def test_actual_app_generates_both_reports_and_respects_machine_name(app):
    from pptx import Presentation
    at, _ = app
    at.button[0].click().run()
    at.text_input[0].set_value("Acceptance Sprayer").run()
    at.button[1].click().run()
    assert not at.exception and not at.error
    result = at.session_state["diesel_report_result"]
    assert result.report_path.name == "Diesel.xlsx"
    assert result.presentation_path.name == "Diesel.pptx"
    assert result.report_metadata.machine_name == "Acceptance Sprayer"
    cover = Presentation(result.presentation_path).slides[0]
    assert any(s.has_text_frame and s.text == "Acceptance Sprayer Diesel" for s in cover.shapes)
    assert len(at.get("download_button")) == 2
    assert len(at.get("image") or at.get("imgs")) == 18
    at.text_input[0].set_value("Another Sprayer").run()
    assert not at.get("download_button")
    assert not at.exception


def test_extracted_client_package_generates_diesel_without_private_data(tmp_path):
    profile = load_reporting_profile(PROFILE_PATH)
    source = _source(tmp_path, profile, omitted=[c.semantic_name for c in profile.raw_channels if not c.required])
    release = build_client_release(PROJECT_ROOT, tmp_path / "release")
    with ZipFile(release.archive_path) as archive:
        names = archive.namelist()
        assert not any("Invoice" in name or "/logs/" in name or name.endswith(".csv") for name in names)
        for suffix in ("config/report_profiles/full_size_sprayer_diesel.yaml",
                       "src/vsm_postprocessing/diesel_validation.py",
                       "src/vsm_postprocessing/diesel_powerpoint_report.py"):
            assert any(name.endswith("/" + suffix) for name in names)
        archive.extractall(tmp_path / "extracted")
    root = next((tmp_path / "extracted").iterdir())
    code = """
from pathlib import Path
import sys
import vsm_postprocessing.ui_config as ui
from vsm_postprocessing.doctor import run_doctor
root = Path.cwd()
assert Path(ui.__file__).is_relative_to(root)
assert [p.profile_id for p in ui.discover_reporting_profiles(root)][-1] == 'full_size_sprayer_diesel'
profile = root / 'config/report_profiles/full_size_sprayer_diesel.yaml'
assert ui.validate_reporting_profile_source(sys.argv[1], profile).is_valid
result = ui.generate_reporting_profile_engineering_report(sys.argv[1], profile, root / 'outputs/diesel')
assert result.report_path.exists() and result.presentation_path.exists()
assert (result.sample_count, result.math_count, result.plot_count, result.slide_count) == (3, 5, 7, 7)
print('Extracted Diesel package: PASS')
"""
    env = dict(os.environ, PYTHONPATH=str(root / "src"))
    completed = subprocess.run([sys.executable, "-c", code, str(source)], cwd=root,
                               env=env, capture_output=True, text=True, timeout=90)
    assert completed.returncode == 0, completed.stdout + completed.stderr
