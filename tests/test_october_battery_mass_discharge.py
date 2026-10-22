"""October 08/04 acceptance tests; workbook coordinates belong only to references."""
from dataclasses import replace
from functools import lru_cache
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook
from pptx import Presentation

from vsm_postprocessing.agro_reporting import calculate_net_mass_discharge_rate, calculate_reporting_kpis
from vsm_postprocessing.reporting_inputs import ReportingInputs
from vsm_postprocessing.report_metadata import resolve_report_metadata
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.excel_report_engine import _profile_template_file
from test_sergio_october_15 import profile_for, profile_path, dataset_for, csv_for
from test_profile_math import _channel, _dataset


@lru_cache
def reference_samples(kind):
    name = ('Diesel Template_04.xlsx' if kind == 'diesel' else
            f'Robo_Sprayer_Electrification_Tamplate_{kind.title()}_08.xlsx')
    path = Path('reference_files') / name
    if not path.exists():
        pytest.skip('Private reference workbook unavailable: ' + name)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = iter(book.worksheets[0].values)
        next(rows); next(rows)
        labels, units = next(rows), next(rows)
        specs = [('Time', 's'), ('Distance', 'm'), ('Agrochemical Discharge', 'N'),
                 ('Agrochemical Discharge', 'kg'), ('Distance', 'km')]
        indices = [next(i for i, (label, unit) in enumerate(zip(labels, units))
                        if str(label).strip() == name and str(unit).casefold() == expected.casefold())
                   for name, expected in specs]
        samples = []
        for row in rows:
            if not isinstance(row[indices[0]], (int, float)):
                break
            samples.append([row[i] for i in indices])
        return np.asarray(samples, dtype=float)
    finally:
        book.close()


@pytest.mark.parametrize('kind,expected_mass,expected_rate,count', [
    ('electric', 474.53109072375099, 39.545575746170783, 3853),
    ('hybrid', 474.53109072375099, 39.545575746170783, 3853),
    ('diesel', 3999.8677111959746, 35.062138398308328, 17872),
])
def test_reference_workbook_endpoints_and_profile_conversion(kind, expected_mass, expected_rate, count):
    samples = reference_samples(kind)
    assert len(samples) == count  # Bottom summary rows are excluded.
    profile = profile_for(kind)
    names = ('track_time', 'track_distance', 'agrochemical_discharge_force')
    if kind == 'diesel':
        names += ('engine_fuel_consumption',)
    raw = tuple(profile.raw_by_semantic_name()[name] for name in names)
    profile = replace(profile, raw_channels=raw,
                      math_channels=(profile.math_by_semantic_name()['agrochemical_discharge'],),
                      statistics=(), kpis=(), plots=(), presentation=None)
    data = _dataset([_channel(c.semantic_name, c.source_name, c.unit, i + 1)
                     for i, c in enumerate(raw)],
                    np.column_stack((samples[:, :3], np.zeros(len(samples)))) if kind == 'diesel' else samples[:, :3])
    stats = calculate_profile_statistics(data, profile)
    np.testing.assert_allclose(stats.math_result.values_by_semantic_name['agrochemical_discharge'],
                               samples[:, 3], rtol=1e-13, atol=1e-10)
    report = calculate_reporting_kpis(stats, resolve_report_metadata('Sample.csv', profile))
    assert report.net_mass_rate.value == pytest.approx(expected_rate, rel=2e-14)
    assert report.trace['net_mass_discharge']['discharged_mass_kg'] == pytest.approx(expected_mass, rel=2e-14)
    assert report.net_mass_rate.unit == 'kg/km'
    assert _profile_template_file(profile, None).name.endswith('_04.xlsx' if kind == 'diesel' else '_08.xlsx')
    if kind == 'diesel':
        assert any('does not start at zero' in w for w in report.trace['warnings'])


@pytest.mark.parametrize('length', [2, 7, 43])
def test_variable_lengths_existing_kg_and_endpoint_semantics(length):
    mass = np.linspace(600, 100, length)
    metric, trace = calculate_net_mass_discharge_rate(mass, 'Kg', np.linspace(0, 2000, length), 'm')
    assert metric.value == 250  # Already kg; never divide by gravity again.
    assert trace['discharged_mass_kg'] == 500
    metric, trace = calculate_net_mass_discharge_rate([500, 100, 400, 0], 'kg', [1, 2, 3, 4], 'km')
    assert metric.value == 125  # Net drop / final distance, including refills.
    assert len(trace['warnings']) == 2


@pytest.mark.parametrize('mass,mu,distance,du,reason', [
    (None, 'kg', [0, 1], 'km', 'mass channel'),
    ([2, 1], 'kg', None, 'km', 'distance channel'),
    ([2, 1], 'N', [0, 1], 'km', 'units'),
    ([2, 1], 'kg', [0, 1], 'ft', 'units'),
    ([2], 'kg', [0], 'm', 'two aligned'),
    ([2, float('nan')], 'kg', [0, 1], 'km', 'nonfinite'),
    ([float('inf'), 1], 'kg', [0, 1], 'km', 'nonfinite'),
    ([2, 1], 'kg', [0, float('inf')], 'km', 'nonfinite'),
    ([2, -1], 'kg', [0, 1], 'km', 'negative remaining'),
    ([1, 2], 'kg', [0, 1], 'km', 'exceeds initial'),
    ([2, 1], 'kg', [0, 0], 'km', 'positive finite'),
    ([2, 1], 'kg', [1, 0], 'km', 'nonmonotonic'),
])
def test_optional_invalid_rate(mass, mu, distance, du, reason):
    metric, trace = calculate_net_mass_discharge_rate(mass, mu, distance, du)
    assert metric.value is None and reason in metric.reason and trace['warnings']


@pytest.mark.parametrize('kind,mass', [('electric', 470), ('hybrid', 725), ('electric', None), ('diesel', None)])
def test_native_report_pipeline_and_battery_geometry(tmp_path, kind, mass):
    from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report
    from vsm_postprocessing.report_selection import ProfileReportSelection
    data = dataset_for(kind)
    gravity = 9.80665 if kind == 'diesel' else 9.81
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -gravity * np.array([600, 400, 200, 100])
    source = csv_for(data, tmp_path / ('RoboSprayer_3500Kg_' + kind + '.csv'))
    settings = {'battery_mass_kg': mass} if mass is not None else {}
    report = generate_reporting_profile_engineering_report(source, profile_path(kind), tmp_path / 'report',
        reporting_inputs=settings, selection=ProfileReportSelection(plots=()))
    prs = Presentation(report.presentation_path)
    text = '\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert '1,666.67 kg/km' in text or '1666.67 kg/km' in text
    if kind == 'diesel':
        assert 'BATTERY PACK MASS' not in text
    else:
        page = prs.slides[4]
        labels = [s for s in page.shapes if s.has_text_frame and s.text == 'BATTERY PACK MASS']
        assert len(labels) == 1
        assert not any(s.has_text_frame and 'BATTERY PACK MASS:' in s.text for s in page.shapes)
        assert any(s.has_text_frame and s.text == (f'{mass} kg' if mass else 'N/A') for s in page.shapes)
        cards = [s for s in page.shapes if s.shape_type == 1 and 1.5*914400 < s.top < 3*914400
                 and s.width > 914400 and s.height > .5*914400 and not s.text]
        assert len(cards) == 7
        for i, card in enumerate(cards):
            assert card.left >= 0 and card.top >= 0
            assert card.left + card.width <= prs.slide_width and card.top + card.height <= prs.slide_height
            for other in cards[i+1:]:
                assert card.left + card.width <= other.left or other.left + other.width <= card.left
        assert labels[0].left == max(s.left for s in page.shapes if s.has_text_frame and s.top == labels[0].top)
    book = load_workbook(report.report_path, data_only=True)
    try:
        rows = {r[0]: r for r in book['Statistics'].values if r[0]}
        assert rows['agro_discharge_rate_kg_per_km'][2] == pytest.approx(500 / .3)
        assert rows['agro_discharge_rate_kg_per_km'][3] == 'kg/km'
    finally:
        book.close()


def test_streamlit_manual_input_survives_reruns(monkeypatch, tmp_path):
    from streamlit.testing.v1 import AppTest
    from vsm_postprocessing import ui_app
    from vsm_postprocessing.ui_config import get_reporting_profile_definition
    definition = get_reporting_profile_definition(ui_app.PROJECT_ROOT, 'robosprayer_electric')
    definition = replace(definition, profile_path=profile_path('electric').resolve())
    monkeypatch.setattr(ui_app, '_october_test_definition', definition, raising=False)
    monkeypatch.setattr(ui_app, '_october_test_source', csv_for(dataset_for('electric'), tmp_path / 'Electric.csv'), raising=False)
    monkeypatch.setattr(ui_app, '_october_test_output', tmp_path / 'ui_report', raising=False)
    app = AppTest.from_string("""
import streamlit as st
from vsm_postprocessing import ui_app
st.session_state['settings'] = ui_app._render_reporting_metadata_inputs(st, ui_app._october_test_definition, 'RoboSprayer_3500Kg_Electric.csv')
if st.button('Generate'):
    from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report
    from vsm_postprocessing.report_selection import ProfileReportSelection
    st.session_state['report'] = generate_reporting_profile_engineering_report(
        ui_app._october_test_source, ui_app._october_test_definition.profile_path, ui_app._october_test_output,
        reporting_inputs=st.session_state['settings'], selection=ProfileReportSelection(plots=()))
""", default_timeout=30).run()
    assert not app.exception and app.session_state['settings'] == {}
    next(w for w in app.text_input if w.label == 'Battery pack mass (kg)').set_value('470').run()
    app.run()
    assert app.session_state['settings']['battery_mass_kg'] == 470
    assert app.session_state['settings']['battery_mass_status'] == 'user_provided'
    metadata = resolve_report_metadata('RoboSprayer_3500Kg_Electric.csv', profile_for('electric'),
                                       reporting_inputs=app.session_state['settings'])
    assert metadata.reporting_inputs.battery_card_display == '470 kg'
    app.button[0].click().run()
    assert not app.exception
    report = app.session_state['report']
    assert any(s.has_text_frame and s.text == '470 kg' for s in Presentation(report.presentation_path).slides[4].shapes)


@pytest.mark.parametrize('kind', ['electric', 'hybrid', 'diesel'])
def test_existing_mass_source_missing_force_and_optional_units(kind):
    from test_sergio_october_15 import remove_channel
    from dataclasses import replace
    data = dataset_for(kind)
    remove_channel(data, 'agrochemical_discharge_force')
    profile = profile_for(kind)
    definition = profile.math_by_semantic_name()['agrochemical_discharge']
    data.channels.append(_channel('remaining_mass', definition.source_name, 'kg', len(data.channels) + 1))
    data.values = np.column_stack((data.values, [600, 400, 200, 100]))
    stats = calculate_profile_statistics(data, profile)
    assert stats.resolution.is_valid
    assert 'agrochemical_discharge' not in stats.math_result.values_by_semantic_name
    report = calculate_reporting_kpis(stats, resolve_report_metadata('Test.csv', profile))
    assert report.net_mass_rate.value == pytest.approx(500 / .3)
    # Optional unit mismatches never invalidate otherwise supported CSV data.
    data.channels[-1] = replace(data.channels[-1], unit='lb')
    stats = calculate_profile_statistics(data, profile)
    assert stats.resolution.is_valid
    report = calculate_reporting_kpis(stats, resolve_report_metadata('Test.csv', profile))
    assert report.net_mass_rate.value is None


@pytest.mark.parametrize('count', [2, 7, 17])
def test_xlsx_importer_excludes_labelled_bottom_statistics(tmp_path, count):
    from openpyxl import Workbook
    from vsm_postprocessing.importer import load_data_file
    book = Workbook()
    sheet = book.active
    sheet.append(['Track_Time', 'Track_Distance', 'HitchRear_Force_Z_VehicleCoordinates'])
    sheet.append(['s', 'm', 'N'])
    for t, d, mass in zip(range(count), np.linspace(0, 2000, count), np.linspace(600, 100, count)):
        sheet.append([t, d, -mass * 9.81])
    for label in ['MAX', 'MIN', 'LAST', 'FIRST']:
        sheet.append([label, 999999, 999999])
    path = tmp_path / 'input_with_statistics.xlsx'
    book.save(path)
    data = load_data_file(path)
    assert data.quality.sample_count == count
    assert data.quality.data_end_row == count + 2
    profile = profile_for('electric')
    profile = replace(profile, raw_channels=tuple(profile.raw_by_semantic_name()[name] for name in
        ('track_time', 'track_distance', 'agrochemical_discharge_force')),
        math_channels=(profile.math_by_semantic_name()['agrochemical_discharge'],), statistics=(), kpis=())
    stats = calculate_profile_statistics(data, profile)
    report = calculate_reporting_kpis(stats, resolve_report_metadata(path, profile))
    assert report.net_mass_rate.value == pytest.approx(250)
