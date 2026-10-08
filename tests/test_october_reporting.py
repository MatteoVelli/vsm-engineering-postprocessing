"""October client corrections, independently exercised on all three real datasets."""
import ast
from dataclasses import replace
from pathlib import Path
import re
from types import SimpleNamespace

import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import ScatterChart
from openpyxl.utils.cell import range_to_tuple
from openpyxl.utils import get_column_letter
from openpyxl.styles import Border, Font, PatternFill

from conftest import (DIESEL_REFERENCE_CSV, ROBOSPRAYER_LATEST_ELECTRIC_CSV,
                      ROBOSPRAYER_LATEST_HYBRID_CSV, require_private_reference_file)
from vsm_postprocessing.errors import ExcelReportError
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report, _profile_equation_text, _write_profile_rms_blocks
from vsm_postprocessing.excel_formulas import profile_arithmetic_formula
from vsm_postprocessing.profile_math import DEFAULT_PROFILE_MATH_CONSTANTS
from vsm_postprocessing.report_profile import MathChannelDefinition, StatisticDefinition
from vsm_postprocessing.statistics_engine import compute_statistic


CASES = [("full_size_sprayer_diesel", DIESEL_REFERENCE_CSV),
         ("robosprayer_hybrid", ROBOSPRAYER_LATEST_HYBRID_CSV),
         ("robosprayer_electric", ROBOSPRAYER_LATEST_ELECTRIC_CSV)]


def title(value):
    return "".join(run.t for p in value.tx.rich.p for run in p.r)


@pytest.fixture(scope="module", params=CASES, ids=[c[0] for c in CASES])
def report(request, tmp_path_factory):
    profile, source = request.param
    source = require_private_reference_file(source, profile)
    return generate_profile_excel_report(source, Path("config/report_profiles") / f"{profile}.yaml",
                                         tmp_path_factory.mktemp(profile + "_october"))


def test_complete_authoritative_equation_inventory_and_raw_distinction(report):
    book = load_workbook(report.report_path)
    try:
        mapping, main = book["Rename From VSM to Astauto"], book.worksheets[0]
        assert mapping.sheet_state == "visible"
        configured = report.profile.math_by_semantic_name()
        calculated = set(report.math_result.calculation_order)
        exported = {c.channel_id for c in report.report_channels if c.kind == "math"}
        assert exported == calculated and exported <= configured.keys()
        displayed = set()
        for col, channel in enumerate(report.report_channels, 1):
            row = col + 2
            assert mapping.cell(row, 6).value == channel.channel_id
            assert mapping.cell(row, 3).value == channel.kind.upper()
            text = mapping.cell(row, 8).value
            if channel.kind != "math":
                if channel.channel_id in configured:
                    assert text.startswith("Raw ") and "fallback" in text
                else:
                    assert text is None
                continue
            definition = configured[channel.channel_id]
            assert text.splitlines()[0] == f"Formula: {channel.channel_id} = {definition.expression}"
            assert mapping.cell(row, 8).data_type == "s"
            assert mapping.cell(row, 2).value == definition.report_name
            assert mapping.cell(row, 5).value == (channel.unit or "-")
            assert mapping.cell(row, 8).alignment.wrap_text
            assert mapping.row_dimensions[row].height >= 30
            assert main.cell(3, col).fill.fgColor.rgb.endswith("C65911")
            if "rpm_nm_to_kw_divisor" in definition.expression:
                assert f"rpm_nm_to_kw_divisor = {DEFAULT_PROFILE_MATH_CONSTANTS['rpm_nm_to_kw_divisor']}" in text
            if "sample_energy_kwh(" in definition.expression:
                assert "dt[0] = t[1] - t[0]" in text
            displayed.add(channel.channel_id)
        assert displayed == exported
    finally:
        book.close()


def test_every_exported_value_raw_mapping_and_statistics_remain_aligned(report):
    book = load_workbook(report.report_path, data_only=True)
    try:
        main = book.worksheets[0]
        for col, channel in enumerate(report.report_channels, 1):
            actual = [main.cell(row, col).value for row in range(5, report.sample_count + 5)]
            expected = report.plotting_result.values_by_semantic_name[channel.channel_id]
            np.testing.assert_allclose(actual, expected, rtol=2e-15, atol=1e-12)
            resolved = report.resolution.resolved.get(channel.channel_id)
            if resolved:
                raw = report.dataset.values[:, report.dataset.channel_index(resolved.channel.channel_id)]
                np.testing.assert_allclose(actual, raw, rtol=2e-15, atol=1e-12)
        statistics = book["Statistics"]
        rows = {r[0].value: r for r in statistics if r[0].value}
        for item in report.statistics_result.statistics:
            assert rows[item.definition.statistic_id][4].value == pytest.approx(item.value, rel=2e-15, abs=1e-12)
        for item in report.statistics_result.kpis:
            assert rows[item.definition.kpi_id][2].value == pytest.approx(item.value, rel=2e-15, abs=1e-12)
    finally:
        book.close()


def test_diesel_corrections_and_other_profile_naming_scope(report):
    if report.profile.metadata.powertrain != "diesel":
        raw = report.profile.raw_by_semantic_name()
        for channel in report.report_channels:
            definition = raw.get(channel.channel_id) or report.profile.math_by_semantic_name()[channel.channel_id]
            assert channel.display_name == definition.report_name
        if report.profile.metadata.powertrain == "hybrid":
            assert raw["generator_torque_1"].report_name.strip() == "ICE Generator Torque"
            assert raw["generator_torque_1"].source_name == "Engine_AuxiliaryTorque_1"
        return
    book = load_workbook(report.report_path)
    try:
        main = book.worksheets[0]
        columns = {c.channel_id: i for i, c in enumerate(report.report_channels, 1)}
        for i in (1, 2):
            name = f"engine_auxiliary_torque_{i}"
            assert report.profile.raw_by_semantic_name()[name].source_name == f"Engine_AuxiliaryTorque_{i}"
            assert main.cell(3, columns[name]).value == f"Auxiliary Torque_{i}"
            assert main.cell(4, columns[name]).value == "Nm"
        titles = [title(c.title) for c in main._charts]
        assert titles.count("Auxiliary Power") == 1
        assert "Engine Power and at Wheels" in titles and "Engine Torque & Power" in titles
        assert not any("Generator Torque" in t or "Auxiliary Torque" in t for t in titles)
        chart = main._charts[titles.index("Auxiliary Power")]
        assert isinstance(chart, ScatterChart)
        assert chart.x_axis.__class__.__name__ == "NumericAxis"
        assert title(chart.x_axis.title) == "Distance [km]"
        assert title(chart.y_axis.title) == "Power [kW]"
        names = ["total_auxiliary_power", "auxiliary_power_1", "auxiliary_power_2"]
        assert [s.tx.v for s in chart.series] == ["Total Auxiliary Power", "Auxiliary Power_1", "Auxiliary Power_2"]
        for series, name in zip(chart.series, names):
            sheet, (col, first, _, last) = range_to_tuple(series.yVal.numRef.f)
            assert sheet == main.title and col == columns[name]
            assert (first, last) == (5, report.sample_count + 4)
            _, (xcol, xfirst, _, xlast) = range_to_tuple(series.xVal.numRef.f)
            assert xcol == columns["distance_km"] and (xfirst, xlast) == (first, last)
        assert chart.legend.position == "b" and chart.legend.overlay is False
        index = titles.index("Auxiliary Power")
        assert (chart.anchor._from.col, chart.anchor._from.row) == (report.report_channel_count + 1 + index % 2 * 8,
                                                                  7 + index // 2 * 20)
        values = report.plotting_result.values_by_semantic_name
        for i in (1, 2):
            np.testing.assert_allclose(values[f"auxiliary_power_{i}"],
                                       values["engine_speed"] * values[f"engine_auxiliary_torque_{i}"] * 2 * np.pi / 60000,
                                       rtol=2e-14, atol=1e-12)
        np.testing.assert_allclose(values[names[0]], values[names[1]] + values[names[2]], rtol=2e-14, atol=1e-12)
    finally:
        book.close()


def test_historical_excel_formula_cannot_override_engine_expression():
    definition = MathChannelDefinition("energy", "Energy", "Energy", "kWh", ("power", "time"),
                                       "sample_energy_kwh(power, time)", "=CA5/3600")
    text = _profile_equation_text(definition, "math")
    assert text.startswith("Formula: energy = sample_energy_kwh(power, time)")
    assert "=CA5" not in text
    assert "dt[0] = t[1] - t[0]" in text
    with pytest.raises(ExcelReportError, match="energy"):
        _profile_equation_text(replace(definition, expression=None), "math")


def test_raw_fallback_is_explicit_and_does_not_claim_a_calculation():
    definition = MathChannelDefinition("auxiliary", "Auxiliary", "Auxiliary", "kW", expression="0",
                                       fallback_when_raw_missing=True)
    assert _profile_equation_text(definition, "vsm").startswith("Raw VSM channel used")
    calculated = _profile_equation_text(definition, "math")
    assert calculated.startswith("Formula: auxiliary = 0")
    assert "raw channel unavailable; deterministic fallback used" in calculated


def test_main_sheet_equations_and_live_reference_inventory(report):
    book = load_workbook(report.report_path)
    try:
        sheet = book.worksheets[0]
        columns = {c.channel_id: i for i, c in enumerate(report.report_channels, 1)}
        definitions = report.profile.math_by_semantic_name()
        live, retained = set(), set()
        for channel in report.report_channels:
            col = columns[channel.channel_id]
            header = sheet.cell(3, col)
            if channel.kind != "math":
                assert header.comment is None
                assert all(sheet.cell(r, col).data_type != "f" for r in range(5, report.sample_count + 5))
                continue
            definition = definitions[channel.channel_id]
            assert header.comment and f"Formula: {channel.channel_id} = {definition.expression}" in header.comment.text
            assert channel.display_name in header.comment.text and f"[{channel.unit or '-'}]" in header.comment.text
            tree = ast.parse(definition.expression, mode="eval")
            names = {n.id for n in ast.walk(tree) if isinstance(n, ast.Name)}
            pointwise = not any(isinstance(n, ast.Call) for n in ast.walk(tree))
            references = {get_column_letter(columns[n]) for n in names if n in columns}
            for row in range(5, report.sample_count + 5):
                cell = sheet.cell(row, col)
                if pointwise:
                    assert cell.data_type == "f"
                    assert {c for c, r in re.findall(r"\$([A-Z]+)(\d+)", cell.value) if int(r) == row} == references
                    assert not any(int(r) != row for _, r in re.findall(r"\$([A-Z]+)(\d+)", cell.value))
                    assert "{row}" not in cell.value
                else:
                    assert cell.data_type != "f" and isinstance(cell.value, (int, float))
            (live if pointwise else retained).add(channel.channel_id)
        assert live | retained == set(report.math_result.calculation_order)
        assert live and retained
    finally:
        book.close()


def test_all_upper_rms_cells_reference_the_actual_sample_range(report):
    book = load_workbook(report.report_path)
    cached = load_workbook(report.report_path, data_only=True)
    try:
        sheet = book.worksheets[0]
        columns = {c.channel_id: i for i, c in enumerate(report.report_channels, 1)}
        items = [s for s in report.statistics_result.statistics if s.definition.placement_group == "top_rms"]
        assert len(items) == 2
        for item in items:
            heading = next(c for c in sheet[1] if c.value == (item.definition.display_name or item.channel_display_name))
            block = next(b for b in sheet.merged_cells.ranges if heading.coordinate in b)
            cell = next(sheet.cell(2, col) for col in range(block.min_col, block.max_col + 1) if sheet.cell(2, col).data_type == "f")
            letter = get_column_letter(columns[item.target_channel])
            extent = f"{letter}5:{letter}{report.sample_count + 4}"
            formula = f"SQRT(SUMSQ({extent})/COUNT({extent}))"
            if report.profile.metadata.powertrain != "diesel":
                assert item.definition.nan_policy == "error"
                formula = f"IF(COUNT({extent})=ROWS({extent}),{formula},NA())"
            assert cell.value == "=" + formula
            assert cached.worksheets[0][cell.coordinate].value == pytest.approx(item.value, rel=2e-15)
        first_row = 7 if report.profile.metadata.powertrain == "diesel" else 5
        for i, chart in enumerate(sheet._charts):
            assert chart.anchor._from.row == first_row + i // 2 * 20
            assert chart.anchor._from.col == report.report_channel_count + 1 + i % 2 * 8
    finally:
        book.close()
        cached.close()


@pytest.mark.parametrize("expression", ["sample_energy_kwh(p, t)", "cumulative_sum(p)", "cumulative_trapezoid(p, t)", "clip(p, 0, 1)", "unknown * 2"])
def test_complex_or_unexported_expressions_retain_python_values(expression):
    assert profile_arithmetic_formula(expression, {"p": 2, "t": 1}, {}) is None


def test_excel_arithmetic_preserves_grouping_constants_and_missing_input_guard():
    template = profile_arithmetic_formula("-(p + q) * p / divisor", {"p": 2, "q": 4}, {"divisor": 9548.8})
    assert template.format(row=12) == "=IF(COUNT($B12,$D12)=2,(((-($B12+$D12))*$B12)/9548.8),NA())"
    assert profile_arithmetic_formula("0", {}, {}) == "=0.0"
    assert profile_arithmetic_formula("p ** 2", {"p": 2}, {}).format(row=5) == "=IF(COUNT($B5)=1,($B5^2.0),NA())"


@pytest.mark.parametrize("count", [3, 7])
@pytest.mark.parametrize("policy", ["error", "omit", "propagate"])
def test_rms_formula_tracks_dynamic_rows_and_engine_nan_policy(count, policy):
    samples = np.arange(1, count + 1, dtype=float)
    if policy != "error":
        samples[1] = np.nan
    expected, used, omitted = compute_statistic(samples, "rms", policy)
    definition = StatisticDefinition("signal_rms", "signal", "rms", nan_policy=policy, placement_group="top_rms")
    item = SimpleNamespace(definition=definition, target_channel="signal", channel_display_name="Signal",
                           channel_unit="kW", value=expected, sample_count=count)
    statistics = SimpleNamespace(statistics=[item], profile=SimpleNamespace(metadata=SimpleNamespace(powertrain="electric")))
    book = Workbook()
    _write_profile_rms_blocks(book.active, statistics, {"time": 1, "signal": 2}, title_fill=PatternFill(), border=Border(), white_bold=Font())
    extent = f"B5:B{count + 4}"
    formula = f"SQRT(SUMSQ({extent})/COUNT({extent}))"
    if policy != "omit":
        formula = f"IF(COUNT({extent})=ROWS({extent}),{formula},NA())"
    assert book.active['B2'].value == "=" + formula
    if policy == "omit":
        assert used == count - 1 and omitted == 1
        assert expected == pytest.approx(np.sqrt(np.mean(samples[np.isfinite(samples)] ** 2)))
    elif policy == "propagate":
        assert np.isnan(expected)
