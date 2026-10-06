"""Canonical workbook references and matching solid data traces across renderers."""
from pathlib import Path
import re

import pytest
from matplotlib.axes import Axes
from openpyxl import load_workbook

from conftest import (ROBOSPRAYER_LATEST_ELECTRIC_CSV, ROBOSPRAYER_LATEST_HYBRID_CSV,
                      require_private_reference_file)
from test_diesel_profile import PROFILE_PATH, _source
from test_profile_plotting import _channel, _dataset, _profile
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.plot_colours import series_colour
from vsm_postprocessing.plotting_engine import PlotDefaults, PlotStyle
from vsm_postprocessing.profile_plotting import render_profile_plots
from vsm_postprocessing.report_profile import (ProfilePlotDefinition, ProfilePlotSeriesDefinition,
                                             RawChannelDefinition, load_reporting_profile)


@pytest.mark.parametrize("profile_id,source", [
    ("full_size_sprayer_diesel", None),
    ("robosprayer_electric", ROBOSPRAYER_LATEST_ELECTRIC_CSV),
    ("robosprayer_hybrid", ROBOSPRAYER_LATEST_HYBRID_CSV),
])
def test_every_summary_value_is_a_direct_canonical_reference(tmp_path, profile_id, source):
    profile_path = Path("config/report_profiles") / (profile_id + ".yaml")
    if source is None:
        source = _source(tmp_path, load_reporting_profile(PROFILE_PATH))
    else:
        source = require_private_reference_file(source, profile_id + " source")
    report = generate_profile_excel_report(source, profile_path, tmp_path / "report")
    formulas = load_workbook(report.report_path)
    cached = load_workbook(report.report_path, data_only=True)
    try:
        sheet = formulas.worksheets[0]
        row = 2 if profile_id == "full_size_sprayer_diesel" else 4
        count = 0
        for col in range(report.report_channel_count + 2, sheet.max_column + 1):
            cell = sheet.cell(row, col)
            if cell.value is None:
                continue
            assert cell.data_type == "f"
            assert re.fullmatch(r"=(?:'Statistics'!)?[A-Z]+[0-9]+", cell.value)
            target = cell.value[1:].split("!")
            canonical = cached[target[0].strip("'")] if len(target) == 2 else cached.worksheets[0]
            assert cached.worksheets[0][cell.coordinate].value == pytest.approx(canonical[target[-1]].value)
            count += 1
        assert count > 10
        # Existing RMS and SUM result cells take precedence over introducing
        # a second numeric authority on the Statistics sheet.
        for statistic in report.statistics_result.statistics:
            if statistic.definition.placement_group != "top_rms" and statistic.definition.operation != "sum":
                continue
            label = statistic.definition.display_name or statistic.channel_display_name
            matches = [sheet.cell(row, col) for col in range(report.report_channel_count + 2, sheet.max_column + 1)
                       if str(sheet.cell(row - 1, col).value).split(" [")[0] == label]
            for cell in matches:
                assert "!" not in cell.value
                if statistic.definition.placement_group == "top_rms":
                    assert re.fullmatch(r"=[A-Z]+2", cell.value)
        for chart in sheet._charts:
            series = [s for component in chart._charts for s in component.series]
            colours = [s.graphicalProperties.line.solidFill.srgbClr for s in series]
            assert len(colours) == len(set(colours))
            assert all(s.graphicalProperties.line.dashStyle == "solid" for s in series)
            assert all(s.marker.symbol in {None, "none"} for s in series)
            assert colours == [series_colour(i) for i in range(len(series))]
        assert not sheet._images
        assert formulas["Statistics"].sheet_state == "visible"
        if profile_id != "full_size_sprayer_diesel":
            statistics = formulas["Statistics"]
            range_row = next(row for row in statistics.iter_rows()
                             if row[0].value == "range_85_battery_km")
            assert range_row[1].value == "Range for 85% Battery"
    finally:
        formulas.close()
        cached.close()


def test_matplotlib_primary_and_secondary_series_use_matching_distinct_solid_colours(tmp_path, monkeypatch):
    plotted = []
    original = Axes.plot

    def capture(self, *args, **kwargs):
        lines = original(self, *args, **kwargs)
        plotted.extend(lines)
        return lines

    monkeypatch.setattr(Axes, "plot", capture)
    names = ["first", "second", "third", "fourth"]
    channels = [_channel("t", "Time", "s")] + [_channel(n, n.title(), "kW", i + 2) for i, n in enumerate(names)]
    profile = _profile(
        raw_channels=[RawChannelDefinition("time", "Time", "Time", "VSM", unit="s")]
                     + [RawChannelDefinition(n, n.title(), n.title(), "VSM", unit="kW") for n in names],
        plots=[ProfilePlotDefinition("power", "Power", "time", tuple(
            ProfilePlotSeriesDefinition(n, axis="secondary" if i % 2 else "primary") for i, n in enumerate(names)))],
    )
    render_profile_plots(_dataset(channels, [[0, 1, 2, 3, 4], [1, 2, 3, 4, 5], [2, 3, 4, 5, 6]]),
                         profile, tmp_path,
                         defaults=PlotDefaults(style=PlotStyle(primary_line_style=":", secondary_line_style="--")))
    assert [line.get_linestyle() for line in plotted] == ["-"] * 4
    assert [line.get_color() for line in plotted] == ["#" + series_colour(i) for i in range(4)]
    assert all(line.get_marker() in {None, "None", ""} for line in plotted)


def test_palette_does_not_repeat_for_large_multiseries_figures():
    assert len({series_colour(i) for i in range(40)}) == 40
