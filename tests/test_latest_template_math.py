from dataclasses import replace
from pathlib import Path
import re

import numpy as np
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from vsm_postprocessing.errors import ConfigurationError, ExcelReportError, MathChannelError
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report, _template_channel_name
from vsm_postprocessing.math_engine import _compile_expression, _evaluate_node
from vsm_postprocessing.profile_math import calculate_profile_math_channels
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.report_profile import load_reporting_profile, resolve_profile
from test_profile_math import _channel, _dataset

WHEELS = ("fl", "fr", "rl", "rr")
COMPONENTS = ("rolling", "compaction", "bulldozing")


def profile_for(kind):
    return load_reporting_profile(f"config/report_profiles/robosprayer_{kind}.yaml")


def dataset_for(profile, times=(0, 1, 2, 3)):
    channels = [_channel(c.semantic_name, c.source_name, c.unit, i + 1)
                for i, c in enumerate(profile.raw_channels)]
    data = _dataset(channels, np.ones((len(times), len(channels))))
    data.values[:, data.channel_index("track_time")] = times
    data.values[:, data.channel_index("track_distance")] = [0, 100, 200, 300]
    data.values[:, data.channel_index("electricsystem_battery_energy")] = [50, 49, 48, 47]
    data.values[:, data.channel_index("electricsystem_battery_soc")] = [100, 98, 96, 94]
    for i, wheel in enumerate(WHEELS, 1):
        data.values[:, data.channel_index(f"driveshaft_torque_{wheel}")] = i * np.array([2, -3, 5, 0])
        data.values[:, data.channel_index(f"wheel_rotationalspeed_{wheel}")] = (i + 1) * np.array([10, 20, -40, 50])
        for j, component in enumerate(COMPONENTS, 1):
            data.values[:, data.channel_index(f"tyre_{component}resistancepower_{wheel}")] = i * j * np.array([1, 2, 4, -8])
    return data


@pytest.mark.parametrize("kind,count", [("electric", 37), ("hybrid", 40)])
@pytest.mark.parametrize("times", [(0, 1, 2, 3), (0, .5, 2, 5)])
def test_four_wheel_and_tyre_math_and_final_kpis(kind, count, times):
    profile = profile_for(kind)
    data = dataset_for(profile, times)
    result = calculate_profile_math_channels(data, profile)
    assert result.calculated_math_count == count
    values = result.values_by_semantic_name
    torque = np.zeros(4)
    power = np.zeros(4)
    for i, wheel in enumerate(WHEELS, 1):
        name = f"driveshaft_torque_{wheel}"
        definition = profile.raw_by_semantic_name()[name]
        assert definition.channel_type == "VSM" and definition.required
        assert name not in profile.math_by_semantic_name()
        assert name not in values and name not in result.calculation_order
        t = i * np.array([2, -3, 5, 0])
        p = t * (i + 1) * np.array([10, 20, -40, 50]) / 9548.8
        np.testing.assert_allclose(values[f"wheel_power_{wheel}"], p)
        torque += t
        power += p
    np.testing.assert_allclose(values["wheel_total_torque"], torque)
    np.testing.assert_allclose(values["wheel_power_total"], power)
    dt = np.r_[times[1] - times[0], np.diff(times)]
    total = np.zeros(4)
    for j, component in enumerate(COMPONENTS, 1):
        expected = 10 * j * np.array([1, 2, 4, -8])
        energy = expected * dt / 3600
        np.testing.assert_allclose(values[f"total_{component}_resistance_power"], expected)
        np.testing.assert_allclose(values[f"tyre_{component}_resistance_energy_kwh"], energy)
        np.testing.assert_allclose(values[f"tyre_{component}_resistance_energy_wh"], energy * 1000)
        np.testing.assert_allclose(values[f"tyre_{component}_resistance_energy_accumulated"], np.cumsum(energy))
        total += expected
    energy = total * dt / 3600
    np.testing.assert_allclose(values["tyre_total_resistance_power"], total)
    np.testing.assert_allclose(values["tyre_total_resistance_energy_kwh"], energy)
    np.testing.assert_allclose(values["tyre_total_resistance_energy_wh"], energy * 1000)
    np.testing.assert_allclose(values["tyre_total_resistance_energy_accumulated"], np.cumsum(energy))
    stats = calculate_profile_statistics(data, profile, math_result=result)
    by_id = {s.definition.statistic_id: s.value for s in stats.statistics}
    kpis = {k.definition.kpi_id: k.value for k in stats.kpis}
    # The last sample reduces the accumulated energy: this must use last, not max.
    assert by_id["tyre_total_energy_accumulated_last"] == pytest.approx(energy.sum())
    assert "tyre_rr_energy_accumulated_last" not in by_id
    assert kpis["total_energy_consumption_tyres_aux"] == pytest.approx(
        energy.sum() + by_id["auxiliary_energy_accumulated_max"])


@pytest.mark.parametrize("kind", ["electric", "hybrid"])
@pytest.mark.parametrize("wheel", WHEELS)
def test_every_wheel_contributes_independently(kind, wheel):
    profile = profile_for(kind)
    data = dataset_for(profile)
    for w in WHEELS:
        data.values[:, data.channel_index(f"driveshaft_torque_{w}")] = 7 if w == wheel else 0
        for component in COMPONENTS:
            data.values[:, data.channel_index(f"tyre_{component}resistancepower_{w}")] = 3 if w == wheel else 0
    values = calculate_profile_math_channels(data, profile).values_by_semantic_name
    np.testing.assert_allclose(values["wheel_total_torque"], 7)
    np.testing.assert_allclose(values["wheel_power_total"], values[f"wheel_power_{wheel}"])
    for component in COMPONENTS:
        np.testing.assert_allclose(values[f"total_{component}_resistance_power"], 3)
    np.testing.assert_allclose(values["tyre_total_resistance_power"], 9)


def test_electric_and_hybrid_share_identical_wheel_and_tyre_logic():
    electric, hybrid = (profile_for(k) for k in ("electric", "hybrid"))
    for name, definition in electric.math_by_semantic_name().items():
        if name.startswith(("wheel_", "tyre_", "total_")):
            assert hybrid.math_by_semantic_name()[name].expression == definition.expression


@pytest.mark.parametrize("kind", ["electric", "hybrid"])
def test_road_height_remains_optional(kind):
    profile = profile_for(kind)
    data = dataset_for(profile)
    index = data.channel_index("track_height")
    data.channels.pop(index)
    data.values = np.delete(data.values, index, axis=1)
    result = calculate_profile_math_channels(data, profile)
    assert result.resolution.is_valid and result.is_complete
    assert [m.definition.semantic_name for m in result.resolution.missing_optional] == ["track_height"]


@pytest.mark.parametrize("kind", ["electric", "hybrid"])
@pytest.mark.parametrize("wheel", WHEELS)
def test_missing_wheel_torque_is_required_source_and_never_synthesized(kind, wheel, tmp_path):
    profile = profile_for(kind)
    data = dataset_for(profile)
    name = f"driveshaft_torque_{wheel}"
    index = data.channel_index(name)
    data.channels.pop(index)
    data.values = np.delete(data.values, index, axis=1)
    resolution = resolve_profile(data, profile)
    assert [m.definition.semantic_name for m in resolution.missing_required] == [name]
    result = calculate_profile_math_channels(data, profile, resolution)
    assert name not in result.values_by_semantic_name
    assert "wheel_total_torque" not in result.values_by_semantic_name
    assert "wheel_power_total" not in result.values_by_semantic_name
    source = tmp_path / "missing.csv"
    source.write_text(
        ",".join(c.source_name for c in data.channels) + "\n" +
        ",".join(c.unit or "" for c in data.channels) + "\n" +
        "\n".join(",".join(map(str, row)) for row in data.values), encoding="utf-8")
    with pytest.raises(ExcelReportError) as error:
        generate_profile_excel_report(source, f"config/report_profiles/robosprayer_{kind}.yaml", tmp_path / "out")
    message = str(error.value)
    assert f"DriveShaft_Torque_{wheel.upper()}" in message
    assert "required source/VSM" in message
    assert "optional source" in message
    assert "MATH" in message and "not requested from the source CSV" in message


@pytest.mark.parametrize("kind", ["electric", "hybrid"])
@pytest.mark.parametrize("source", ["driveshaft_torque_fl", "tyre_bulldozingresistancepower_fr",
                                    "tyre_compactionresistancepower_rl", "tyre_rollingresistancepower_rr"])
def test_nan_is_rejected_without_omitting_a_wheel(kind, source):
    profile = profile_for(kind)
    data = dataset_for(profile)
    data.values[2, data.channel_index(source)] = np.nan
    with pytest.raises(MathChannelError, match="non-finite values.*sample index 2"):
        calculate_profile_math_channels(data, profile)


@pytest.mark.parametrize("kind,version", [("electric", "05"), ("hybrid", "06")])
def test_latest_workbook_channel_classifications_and_selection(kind, version):
    profile = profile_for(kind)
    path = Path(f"reference_files/Robo_Sprayer_Electrification_Tamplate_{kind.title()}_{version}.xlsx")
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True)
        entries = list(workbook.worksheets[1].iter_rows(min_row=3, values_only=True))
        workbook.close()
    remaining = list(profile.raw_channels + profile.math_channels)
    for row in entries:
        if not row[1]:
            continue
        candidates = [c for c in remaining if c.channel_type == row[3] and
                      _template_channel_name(c.source_name) == _template_channel_name(row[1])]
        if not candidates:
            candidates = [c for c in remaining if c.channel_type == row[3] and
                          _template_channel_name(c.report_name) == _template_channel_name(row[2])]
        assert candidates, row
        channel = candidates[0]
        assert channel.for_plot == bool(row[4]), channel.semantic_name
        remaining.remove(channel)
    assert [c.semantic_name for c in remaining] == (["track_height"] if kind == "electric" else [])
    assert not profile.raw_by_semantic_name()["track_height"].required


@pytest.mark.parametrize("kind,version,torque_col,total_col", [
    ("electric", "05", "HB", "HG"), ("hybrid", "06", "HJ", "HO")])
def test_workbook_formulas_at_individual_samples_and_final_energy(kind, version, torque_col, total_col):
    profile = profile_for(kind)
    path = Path(f"reference_files/Robo_Sprayer_Electrification_Tamplate_{kind.title()}_{version}.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=False)
    rows = list(workbook.worksheets[0].values)
    workbook.close()
    # Data rows precede the workbook's bottom statistics. Read raw values only.
    samples = []
    for row in rows[4:]:
        if not isinstance(row[0], (int, float)):
            break
        samples.append(row)
    assert len(samples) > 100
    columns = {}
    for i, (label, unit) in enumerate(zip(rows[2], rows[3])):
        columns.setdefault((str(label).strip(), str(unit).strip()), []).append(i)
    raw = [c for c in profile.raw_channels if c.semantic_name == "track_time" or
           c.semantic_name.startswith(("driveshaft_torque_", "wheel_rotationalspeed_")) or
           (c.semantic_name.startswith("tyre_") and "resistancepower_" in c.semantic_name)]
    math = [c for c in profile.math_channels if c.semantic_name.startswith(("wheel_", "tyre_")) or
            c.semantic_name in [f"total_{component}_resistance_power" for component in COMPONENTS]]
    # Consume duplicate FR/RR rolling-power labels in wheel order, as in the workbook.
    indices = [0 if c.semantic_name == "track_time" else columns[(c.report_name.strip(), c.unit)].pop(0) for c in raw]
    data = _dataset([_channel(c.semantic_name, c.source_name, c.unit, i + 1) for i, c in enumerate(raw)],
                    [[r[index] for index in indices] for r in samples])
    result = calculate_profile_math_channels(data, replace(profile, raw_channels=tuple(raw), math_channels=tuple(math)))
    assert result.is_complete
    context = {get_column_letter(i + 1): np.array([r[i] for r in samples], dtype=float)
               for i in range(len(rows[2])) if all(isinstance(r[i], (int, float)) for r in samples)}
    formula_by_label = {}
    # Evaluate workbook arithmetic through the safe expression engine, resolving column dependencies.
    for i, formula in enumerate(samples[0]):
        if not isinstance(formula, str) or not formula.startswith("=") or "SUM(" in formula:
            continue
        if not any(c.report_name.split(" [")[0].replace("Accumulated", "Accomulated") == str(rows[2][i]).strip().split(" [")[0]
                   for c in math):
            continue
        expression = re.sub(r"([A-Z]+)5", r"\1", formula[1:])
        compiled = _compile_expression(expression, context="template regression")
        values = _evaluate_node(compiled.tree.body, context)
        context[get_column_letter(i + 1)] = values
        formula_by_label[(_template_channel_name(rows[2][i]), rows[3][i])] = values
    for c in math:
        actual = result.values_by_semantic_name[c.semantic_name]
        if c.semantic_name.endswith("_accumulated"):
            energy = result.values_by_semantic_name[c.semantic_name.replace("_accumulated", "_kwh")]
            np.testing.assert_allclose(actual, np.cumsum(energy))
            assert actual[-1] == pytest.approx(energy.sum())
            index = next(i for i, label in enumerate(rows[2])
                         if _template_channel_name(str(label)) == _template_channel_name(c.report_name))
            # Document the reference's row-6 anchor bug instead of reproducing it.
            assert re.fullmatch(r"=SUM\([A-Z]+5:[A-Z]+\$6\)", samples[0][index])
            assert actual[-1] - energy[1:].sum() == pytest.approx(energy[0], abs=1e-10)
        else:
            # Templates incorrectly label total tyre power kWh; the application uses kW.
            unit = "kWh" if c.semantic_name == "tyre_total_resistance_power" else c.unit
            expected = formula_by_label[(_template_channel_name(c.report_name), unit)]
            np.testing.assert_allclose(actual[[0, 1, 17, 100, -1]], expected[[0, 1, 17, 100, -1]])
    np.testing.assert_allclose(result.values_by_semantic_name["wheel_total_torque"], context[torque_col])
    np.testing.assert_allclose(result.values_by_semantic_name["wheel_power_total"], context[total_col])


@pytest.mark.parametrize("overrides", ["{unknown: true}", "{track_time: 'yes'}", "[]"])
def test_plot_overrides_are_strictly_validated(tmp_path, overrides):
    parent = Path("config/report_profiles/robosprayer_electric.yaml").resolve().as_posix()
    path = tmp_path / "invalid.yaml"
    path.write_text(f"version: 1\nprofile:\n  profile_id: test\n  name: Test\n  extends: {parent}\n"
                    f"channels:\n  plot_overrides: {overrides}\n")
    with pytest.raises(ConfigurationError, match="plot_overrides"):
        load_reporting_profile(path)
