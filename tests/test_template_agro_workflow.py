"""Automatic template mapping, familiar Excel layout and optional expert controls."""
from dataclasses import asdict
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
import pytest
import yaml

from test_profile_math import _channel
from test_sergio_october_15 import dataset_for, profile_for, profile_path, csv_for
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.report_selection import ProfileReportSelection


@pytest.mark.parametrize('filename,force,mass', [
    ('Robo_Sprayer_Electrification_Tamplate_Electric.xlsx', 'BU', 'BV'),
    ('Robo_Sprayer_Electrification_Tamplate_Hybrid_06.xlsx', 'CC', 'CD'),
    ('Robo_Sprayer_Electrification_Tamplate_Hybrid.xlsx', 'CB', 'CC'),
    ('Sprayer_Caiman_SP_9300Kg_Electrification_03.xlsx', 'AK', 'AL'),
])
def test_original_template_mapping_and_gravity(filename, force, mass):
    book = load_workbook(Path('reference_files') / filename, read_only=True)
    try:
        sheet = book.worksheets[0]
        assert sheet[force+'4'].value == 'N'
        assert sheet[mass+'4'].value.lower() == 'kg'
        assert sheet[mass+'5'].value == f'=-1*{force}5/9.81'
    finally:
        book.close()


@pytest.mark.parametrize('kind', ('electric', 'hybrid', 'diesel'))
@pytest.mark.parametrize('load', (2000, 4000))
def test_unknown_simulation_automatically_uses_template_mapping(tmp_path, kind, load):
    data = dataset_for(kind, np.arange(7))
    mass = np.linspace(load, 51, 7)
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81*mass
    before = calculate_profile_statistics(data, profile_for(kind))
    result = generate_profile_excel_report(csv_for(data, tmp_path / f'Unknown_Carrier_{load}_cycle_Z.csv'),
        profile_path(kind), tmp_path/'report', selection=ProfileReportSelection(plots=()))
    assert result.report_metadata.configuration_detection == 'unmatched'
    assert result.report_metadata.reporting_inputs.hitch_force_is_remaining_agro_mass
    assert result.statistics_result.reporting_agro.capacity.interpretation_status == 'user_confirmed'
    assert result.statistics_result.reporting_agro.capacity.value == pytest.approx(load)
    assert result.statistics_result.reporting_agro.capacity.status == 'inferred'
    assert result.statistics_result.reporting_agro.mass_rate.value == pytest.approx((load-51)/300)
    assert result.statistics_result.reporting_agro.rate.value is None
    for name, values in before.math_result.values_by_semantic_name.items():
        np.testing.assert_array_equal(result.math_result.values_by_semantic_name[name], values)
    assert [s.value for s in before.statistics] == [s.value for s in result.statistics_result.statistics]
    assert [k.value for k in before.kpis] == [k.value for k in result.statistics_result.kpis]
    book = load_workbook(result.report_path)
    try:
        location = result.workbook_structure['agro_source']['mass']
        if kind in {'electric', 'hybrid'}:
            assert 'Agro Mass Source' not in book
            assert location['sheet'] == book.worksheets[0].title
        else:
            assert location['sheet'] == 'Metadata'  # Legacy MATH retains 9.80665.
        assert book.sheetnames == [book.worksheets[0].title,'Rename From VSM to Astauto','Metadata','Statistics','Formula Support']
        columns = {c.channel_id: i for i,c in enumerate(result.report_channels,1)}
        assert book.worksheets[0].cell(5,columns['agrochemical_discharge']).data_type == 'f'
        rows = {r[0].value:r for r in book['Statistics'] if str(r[0].value).startswith('agro_')}
        assert len(rows) == 4 and all(r[2].data_type == 'f' for r in rows.values())
        manifest = json.loads(result.manifest_path.read_text())
        assert manifest['visible_sheet_names'] == [s.title for s in book if s.sheet_state=='visible']
        assert manifest['hidden_sheet_names'] == [s.title for s in book if s.sheet_state!='visible']
        assert 'Agro Mass Source' not in manifest['visible_sheet_names']
    finally:
        book.close()


@pytest.mark.parametrize('export', (('track_time',), ('track_time','track_distance','agrochemical_discharge'),
                                  ('track_time','track_distance','agrochemical_discharge_force')))
def test_independent_export_selection_keeps_hidden_live_source(tmp_path, export):
    data = dataset_for('electric', np.arange(7))
    data.values[:,data.channel_index('agrochemical_discharge_force')] = -9.81*np.linspace(2000,50,7)
    result = generate_profile_excel_report(csv_for(data,tmp_path/'Unseen.csv'),profile_path('electric'),tmp_path/'report',
        selection=ProfileReportSelection(export_channels=export, plots=()))
    assert {c.channel_id for c in result.report_channels} == set(export)
    book = load_workbook(result.report_path)
    try:
        assert 'Agro Mass Source' not in book
        assert book['Metadata']['E5'].data_type == 'f'
        assert all(book['Metadata'].column_dimensions[c].hidden for c in ('D','E','F'))
        assert result.statistics_result.reporting_agro.capacity.value == pytest.approx(2000)
        assert all(c.value != '#REF!' for s in book for row in s for c in row)
    finally:
        book.close()


@pytest.mark.parametrize('kind', ('electric','hybrid','diesel'))
@pytest.mark.parametrize('selected', (False,True))
def test_configured_direct_mass_reuses_exported_source_or_hidden_support(tmp_path,kind,selected):
    data = dataset_for(kind,np.arange(7))
    data.values[:,data.channel_index('agrochemical_discharge_force')] = 0
    data.channels.append(_channel('chemical_mass','Simulation remaining chemical mass','kg',len(data.channels)+1))
    data.values = np.column_stack((data.values,np.linspace(1800,50,7)))
    extension = tmp_path/'direct.yaml'
    metadata = asdict(profile_for(kind).metadata)
    # A future profile can export its own kg channel without inheriting a fixed
    # Sergio template's column selection; the ordinary templates stay intact.
    metadata.update(extends=str(profile_path(kind).resolve()), profile_id='future_direct_'+kind, excel_template=None,
        reporting_inputs={**metadata['reporting_inputs'], 'agro_mass_source_policy':'direct_mass',
                          'agro_mass_source':'chemical_remaining',
                          'agro_mass_source_provenance':'Explicit simulation kg source fixture'})
    extension.write_text(yaml.safe_dump({'version':1,'profile':metadata,
        'channels':{'raw':[{'semantic_name':'chemical_remaining','source_name':'Simulation remaining chemical mass',
            'report_name':'Remaining Chemical Mass','unit':'kg','required':False,'for_report':True,
            'channel_type':'VSM'}]}}),encoding='utf-8')
    result = generate_profile_excel_report(csv_for(data,tmp_path/'Prototype.csv'),extension,tmp_path/'report',
        selection=ProfileReportSelection(plots=(),export_channels=('track_time',) if selected else None))
    assert result.statistics_result.reporting_agro.capacity.value == pytest.approx(1800)
    assert result.statistics_result.reporting_agro.capacity.status == 'simulated'
    book = load_workbook(result.report_path)
    try:
        if selected:
            assert 'Agro Mass Source' not in book
            assert result.workbook_structure['agro_source']['mass']['sheet'] == 'Metadata'
            assert len(result.report_channels) == 1
        else:
            assert 'Agro Mass Source' not in book
            location = result.workbook_structure['agro_source']['mass']
            assert location['sheet'] == book.worksheets[0].title
            assert book.worksheets[0][location['column']+'5'].value == 1800
    finally:
        book.close()


@pytest.mark.parametrize('kind',('electric','hybrid','diesel'))
def test_normal_ui_and_density_do_not_expose_source_configuration(monkeypatch,kind):
    from streamlit.testing.v1 import AppTest
    from vsm_postprocessing import ui_app
    from vsm_postprocessing.ui_config import get_reporting_profile_definition
    profile_id = 'full_size_sprayer_diesel' if kind=='diesel' else 'robosprayer_'+kind
    monkeypatch.setattr(ui_app,'_auto_agro_test_definition',get_reporting_profile_definition(ui_app.PROJECT_ROOT,profile_id),raising=False)
    app = AppTest.from_string('''
import streamlit as st
from vsm_postprocessing import ui_app
st.session_state['settings'] = ui_app._render_reporting_metadata_inputs(st,ui_app._auto_agro_test_definition,'Future_Carrier.csv')
''').run()
    assert not app.exception and app.session_state['settings']=={}
    assert not app.selectbox
    assert not any('interpretation' in w.label for w in app.checkbox)
    assert any(e.label=='Advanced Configuration' for e in app.expander)
    next(w for w in app.checkbox if w.label=='Override automatically detected reporting information').check().run()
    assert not any(w.label=='Agro mass source policy' for w in app.selectbox)
    assert not any('interpretation' in w.label for w in app.checkbox)
    next(w for w in app.checkbox if w.label=='Specify agrochemical density [kg/L]').check().run()
    next(w for w in app.text_input if w.label=='Agrochemical density [kg/L] (strictly positive)').set_value('1.2').run()
    next(w for w in app.text_input if w.label=='Density source / assumption').set_value('Explicit test fluid').run()
    assert app.session_state['settings']=={'density_kg_per_L':1.2,'density_provenance':'Explicit test fluid'}
    labels = [w.label for elements in (app.checkbox,app.selectbox,app.text_input) for w in elements]
    assert not any('mass source' in label.lower() or 'source configuration' in label.lower() or 'interpretation' in label.lower() for label in labels)
    assert app.session_state['settings']=={'density_kg_per_L':1.2,'density_provenance':'Explicit test fluid'}
