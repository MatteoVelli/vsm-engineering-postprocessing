"""Acceptance checks for the four separate, time-based October soil plots."""
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches
import pytest
import yaml

from vsm_postprocessing.errors import ConfigurationError, PlottingError
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.profile_math import calculate_profile_math_channels
from vsm_postprocessing.profile_plotting import _prepare_plot_definition, render_profile_plots
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from vsm_postprocessing.report_profile import load_reporting_profile
from vsm_postprocessing.report_selection import ProfileReportSelection, select_profile
from test_sergio_october_15 import csv_for, dataset_for, profile_for, profile_path, remove_channel

SIGNALS = (
    "total_bulldozing_resistance_power",
    "tyre_bulldozing_resistance_energy_accumulated",
    "total_compaction_resistance_power",
    "tyre_compaction_resistance_energy_accumulated",
)
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def signed_dataset(kind):
    data = dataset_for(kind, (17, 18, 1117, 1817, 3617))
    for mechanism, power in (("bulldozing", [-8, -4, 12, -8, 4]),
                             ("compaction", [4, 0, -4, 8, -4])):
        for wheel in ("fl", "fr", "rl", "rr"):
            data.values[:, data.channel_index(f"tyre_{mechanism}resistancepower_{wheel}")] = np.array(power) / 4
    return data


@pytest.fixture(scope="module", params=("electric", "hybrid"))
def soil_report(request, tmp_path_factory):
    kind = request.param
    folder = tmp_path_factory.mktemp("soil_" + kind)
    data = signed_dataset(kind)
    excel = generate_profile_excel_report(csv_for(data, folder / "source.csv"), profile_path(kind), folder / "excel")
    ppt = build_profile_powerpoint_report(excel, folder / "ppt")
    return kind, data, excel, ppt


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
def test_signals_reuse_measured_wheel_power_and_signed_observed_time_energy(kind):
    profile, data = profile_for(kind), signed_dataset(kind)
    result = calculate_profile_math_channels(data, profile)
    time = data.values[:, data.channel_index("track_time")]
    for mechanism in ("bulldozing", "compaction"):
        raw = sum(data.values[:, data.channel_index(f"tyre_{mechanism}resistancepower_{w}")]
                  for w in ("fl", "fr", "rl", "rr"))
        power, energy = f"total_{mechanism}_resistance_power", f"tyre_{mechanism}_resistance_energy_accumulated"
        np.testing.assert_array_equal(result.values_by_semantic_name[power], raw)
        expected = np.r_[0, np.cumsum((raw[:-1] / 2 + raw[1:] / 2) * np.diff(time) / 3600)]
        np.testing.assert_allclose(result.values_by_semantic_name[energy], expected)
        assert expected[0] == 0 and np.any(np.diff(expected) < 0)  # Signed, never clipped or reintegrated.
        for key, unit in ((power, "kW"), (energy, "kWh")):
            definition = profile.plots_by_id()[key]
            assert definition.x == "track_time" and definition.x_label == "Time [s]"
            assert definition.series[0].semantic_name == key and definition.series[0].required
            assert len(definition.series) == 1 and definition.secondary_y_label is None
            assert profile.math_by_semantic_name()[key].unit == unit


def test_all_four_native_excel_charts_have_exact_ranges_units_and_clear_placements(soil_report):
    _, _, report, _ = soil_report
    book = load_workbook(report.report_path)
    cached = load_workbook(report.report_path, data_only=True)
    try:
        sheet = book.worksheets[0]
        columns = {c.channel_id: get_column_letter(i) for i, c in enumerate(report.report_channels, 1)}
        with ZipFile(report.report_path) as archive:
            charts = [n for n in archive.namelist() if n.startswith("xl/charts/chart") and n.endswith(".xml")]
            assert len(charts) == len(report.plotting_result.rendered_plots)
            chart_ids = [p.plot_id for p in report.plotting_result.rendered_plots]
            for key in SIGNALS:
                part = charts[chart_ids.index(key)]
                node = etree.fromstring(archive.read(part))
                for axis, channel in (("xVal", "track_time"), ("yVal", key)):
                    refs = node.xpath(f"//c:ser/c:{axis}/c:numRef/c:f/text()", namespaces=NS)
                    assert refs == [f"'{sheet.title}'!${columns[channel]}$5:${columns[channel]}${report.sample_count + 4}"]
                text = node.xpath("//a:t/text()", namespaces=NS)
                definition = report.profile.plots_by_id()[key]
                assert definition.title in text and "Time [s]" in text and definition.primary_y_label in text
                actual = [cached.worksheets[0][f"{columns[key]}{r}"].value for r in range(5, report.sample_count + 5)]
                np.testing.assert_allclose(actual, report.math_result.values_by_semantic_name[key])
                assert all(sheet[f"{columns[key]}{r}"].data_type == "f" for r in range(5, report.sample_count + 5))
        new = [sheet._charts[chart_ids.index(key)] for key in SIGNALS]
        old = [chart for chart, key in zip(sheet._charts, chart_ids) if key not in SIGNALS and not key.startswith("tyre_total_resistance_")]
        assert len(new) == 4
        assert all((c.anchor.ext.cx, c.anchor.ext.cy) == (old[0].anchor.ext.cx, old[0].anchor.ext.cy) for c in new)
        anchors = [(c.anchor._from.row, c.anchor._from.col) for c in sheet._charts]
        assert len(anchors) == len(set(anchors))
        # Appending can fill the right-hand slot of the existing last row.
        assert min((c.anchor._from.row, c.anchor._from.col) for c in new) > max(
            (c.anchor._from.row, c.anchor._from.col) for c in old)
        assert {"tyre_soil_bulldozing", "tyre_soil_compaction"} <= {p.plot_id for p in report.plotting_result.rendered_plots}
    finally:
        book.close()
        cached.close()


def test_both_new_powerpoint_pages_reuse_exact_plots_and_approved_style(soil_report):
    _, _, excel, ppt = soil_report
    prs = Presentation(ppt.presentation_path)
    assert len(prs.slides) == (16 if excel.profile.metadata.powertrain == "electric" else 15)
    assert ppt.powerpoint_result.config.slides[-2].slide_id == "torque_maps"
    assert ppt.powerpoint_result.config.slides[-1].slide_id == "simulation_summary"
    assert [s.slide_id for s in ppt.powerpoint_result.config.slides[9:11]] == ["road_profile", "wheel_steering_angles"]
    plots = {p.plot_id: p for p in excel.plotting_result.rendered_plots}
    reference = prs.slides[8]
    ids = [s.slide_id for s in ppt.powerpoint_result.config.slides]
    for number, expected in ((ids.index("soil_bulldozing"), SIGNALS[:2]), (ids.index("soil_compaction"), SIGNALS[2:])):
        slide = prs.slides[number]
        assert ppt.powerpoint_result.config.slides[number].plot_ids == expected
        assert slide.element.cSld.bg.xml == reference.element.cSld.bg.xml
        pictures = [s for s in slide.shapes if s.shape_type == 13 and s.top > Inches(2)]
        assert len(pictures) == 2
        assert pictures[0].left + pictures[0].width < pictures[1].left
        for picture, key in zip(pictures, expected):
            assert picture.image.blob == Path(plots[key].png_file).read_bytes()
            template = Presentation("reference_files/RoboSprayer_Electric_Report_Astauto_Colours.pptx")
            assert (picture.width, picture.height) == (template.slides[8].shapes[32].width, template.slides[8].shapes[32].height)
            assert picture.left > 0 and picture.top + picture.height < Inches(7)
        assert any(s.shape_type == 13 and s.top < Inches(.5) for s in slide.shapes)  # Existing Astauto logo.
        title = next(s for s in slide.shapes if s.has_text_frame and s.top == Inches(.44))
        font = title.text_frame.paragraphs[0].runs[0].font
        assert font.name == "Cambria" and font.size.pt == 25


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
@pytest.mark.parametrize("invalid", ("missing", "nan", "infinite", "wrong_unit"))
def test_optional_invalid_sources_disable_only_affected_soil_pair(kind, invalid, tmp_path):
    data, profile = signed_dataset(kind), profile_for(kind)
    name = "tyre_bulldozingresistancepower_fl"
    if invalid == "missing":
        remove_channel(data, name)
    elif invalid == "wrong_unit":
        index = data.channel_index(name)
        data.channels[index] = replace(data.channels[index], unit="W")
    else:
        data.values[1, data.channel_index(name)] = np.nan if invalid == "nan" else np.inf
    profile = replace(profile, plots=tuple(profile.plots_by_id()[key] for key in SIGNALS))
    result = render_profile_plots(data, profile, tmp_path)
    assert [p.plot_id for p in result.rendered_plots] == list(SIGNALS[2:])
    missing = {p.definition.plot_id: p for p in result.unavailable_plots}
    assert set(missing) == set(SIGNALS[:2])
    assert all(p.missing_semantic_names == (key,) for key, p in missing.items())
    assert any("Tyre_BulldozingResistancePower_FL" in d for d in result.resolution.diagnostics)


def test_missing_pair_is_explicit_in_powerpoint_without_fabricating_charts(tmp_path):
    data = dataset_for("electric")
    remove_channel(data, "tyre_bulldozingresistancepower_fl")
    excel = generate_profile_excel_report(csv_for(data, tmp_path / "missing.csv"), profile_path("electric"), tmp_path / "excel")
    ppt = build_profile_powerpoint_report(excel, tmp_path / "ppt")
    prs = Presentation(ppt.presentation_path)
    ids = [s.slide_id for s in ppt.powerpoint_result.config.slides]
    slide = prs.slides[ids.index("soil_bulldozing")]
    text = " ".join(s.text for s in slide.shapes if s.has_text_frame)
    assert text.count("PLOT DATA NOT AVAILABLE") == 2
    assert all(key in text for key in SIGNALS[:2])
    assert not any(s.shape_type == 13 and s.top > Inches(2) for s in slide.shapes)
    assert len(prs.slides) == 16
    assert ppt.powerpoint_result.config.slides[-2].slide_id == "torque_maps"


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
def test_measured_zeros_remain_valid_plots(kind, tmp_path):
    data, profile = dataset_for(kind), profile_for(kind)
    for m in ("bulldozing", "compaction"):
        for w in ("fl", "fr", "rl", "rr"):
            data.values[:, data.channel_index(f"tyre_{m}resistancepower_{w}")] = 0
    profile = replace(profile, plots=tuple(profile.plots_by_id()[key] for key in SIGNALS))
    result = render_profile_plots(data, profile, tmp_path)
    assert len(result.rendered_plots) == 4 and not result.unavailable_plots
    for key in SIGNALS:
        np.testing.assert_array_equal(result.values_by_semantic_name[key], 0)


@pytest.mark.parametrize("key", SIGNALS)
def test_empty_or_misaligned_series_cannot_be_rendered(key):
    definition = profile_for("electric").plots_by_id()[key]
    from test_profile_math import _channel
    channels = {"track_time": _channel("t", "Track_Time", "s"), key: _channel(key, key, "kW")}
    with pytest.raises(PlottingError, match="no finite values"):
        _prepare_plot_definition(definition, channels, {"track_time": np.array([]), key: np.array([])}, 0)
    with pytest.raises(PlottingError, match="not aligned"):
        _prepare_plot_definition(definition, channels, {"track_time": np.array([0, 1]), key: np.array([1])}, 2)


def test_diesel_and_explicit_legacy_profiles_keep_their_original_plot_inventories():
    for path in (profile_path("diesel"), Path("config/report_profiles/robosprayer_electric.yaml"),
                 Path("config/report_profiles/robosprayer_hybrid.yaml")):
        profile = load_reporting_profile(path)
        assert not set(SIGNALS) & profile.plots_by_id().keys()
        additional = [s.slide_id for s in profile.presentation.slides if s.mode == "additional_plot_pair"]
        assert additional == (["torque_maps"] if profile.metadata.powertrain == "diesel" else [])
    selected = select_profile(profile_for("electric"), ProfileReportSelection(
        plots=(), powerpoint_plots={"soil_bulldozing": SIGNALS[:2], "soil_compaction": ()}))
    assert set(SIGNALS[:2]) <= selected.plots_by_id().keys()
    assert not selected.presentation.slides_by_id()["soil_compaction"].plots


@pytest.mark.parametrize("invalid", ("unknown_plot", "duplicate_id", "missing_title", "too_many_plots", "invalid_mode"))
def test_additional_slide_configuration_fails_clearly_for_invalid_definitions(invalid, tmp_path):
    slide = {"slide_id": "soil_extra", "mode": "additional_plot_pair", "title": "Soil", "plots": [SIGNALS[0]]}
    if invalid == "unknown_plot": slide["plots"] = ["unknown_quantity"]
    if invalid == "duplicate_id": slide["slide_id"] = "vehicle_operation"
    if invalid == "missing_title": del slide["title"]
    if invalid == "too_many_plots": slide["plots"] = list(SIGNALS)
    if invalid == "invalid_mode": slide["mode"] = "unrecognised"
    payload = {"version": 1, "profile": {"profile_id": "invalid", "name": "Invalid",
                "extends": str(profile_path("electric").resolve())},
               "overrides": {"additional_presentation_slides": [slide]}}
    path = tmp_path / "invalid.yaml"
    path.write_text(yaml.safe_dump(payload), encoding="utf-8")
    with pytest.raises(ConfigurationError):
        load_reporting_profile(path)
