"""Prevent specifications leaking between machines, templates and powertrains."""
from dataclasses import replace
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from pptx import Presentation
import pytest

from conftest import require_private_reference_file
from vsm_postprocessing.agro_reporting import calculate_reporting_kpis
from vsm_postprocessing.errors import ConfigurationError
from vsm_postprocessing.importer import ImportOptions, load_data_file
from vsm_postprocessing.machine_specifications import validate_specifications, machine_metadata_pills
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.report_metadata import resolve_report_metadata
from vsm_postprocessing.report_profile import load_reporting_profile
from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report


BASE = Path('config/report_profiles')
CASES = [
    ('robosprayer_electric', 'RoboSprayer_3500Kg_Electric_UoL_Drive_Cycle_12kph_DriverPC_12-5pcGrad_01', 3500, 12, 498.1569826707441, 45.33840184907765),
    ('robosprayer_electric', 'RoboSprayer_8500Kg_Electric_02_UoL_Drive_Cycle_20kph_DriverPC_Flat_02', 8500, 20, 3996.788990825688, 363.1218540582613),
    ('robosprayer_hybrid', 'Sprayer_Caiman_SP_9300Kg_Hybrid_12x1Km_4000Kg_Chem', 9300, None, 3998.6544342507646, None),
    ('full_size_sprayer_diesel', 'Full_Size_Sprayer_8500Kg_Diesel_Test_03_Crop_Field_MatLab_20kph_Gradient', 8500, 20, None, None),
]


@pytest.mark.parametrize('profile_id,stem,mass,speed,capacity,rate', CASES)
def test_exact_raw_reference_calculations_and_confidence(profile_id, stem, mass, speed, capacity, rate):
    source = require_private_reference_file(Path('reference_files') / (stem + '.csv'), stem)
    profile = load_reporting_profile(BASE / (profile_id + '.yaml'))
    dataset = load_data_file(source, ImportOptions(strict=False))
    result = calculate_profile_statistics(dataset, profile)
    metadata = resolve_report_metadata(source, profile)
    agro = calculate_reporting_kpis(result, metadata)
    assert agro.capacity.value == pytest.approx(capacity) if capacity is not None else agro.capacity.value is None
    assert agro.net_mass_rate.value == pytest.approx(rate) if rate is not None else agro.net_mass_rate.value is None
    assert metadata.machine_specifications['gvw_kg']['value'] == mass
    assert metadata.machine_specifications['gvw_kg']['status'] == 'filename_claim'
    assert metadata.machine_specifications['target_speed_kph']['value'] == speed
    assert metadata.machine_specifications['battery']['value'] is None
    if rate is not None:
        force = result.resolution.resolved['agrochemical_discharge_force']
        raw = dataset.values[:, dataset.channel_index(force.channel.channel_id)]
        distance = result.resolution.resolved['track_distance']
        d = dataset.values[:, dataset.channel_index(distance.channel.channel_id)]
        gravity = 9.81
        assert agro.net_mass_rate.value == pytest.approx(((-raw[0] / gravity) - (-raw[-1] / gravity)) / (d[-1] / 1000), rel=1e-14)
    if mass == 8500 and profile_id == 'robosprayer_electric':
        assert metadata.configuration_detection == 'unmatched'
        assert 'not physically validated' in agro.net_mass_rate.definition
        assert agro.trace['net_mass_discharge']['interpretation_status'] == 'assumed'


def configured_profile():
    profile = load_reporting_profile(BASE / 'robosprayer_electric.yaml')
    configurations = [dict(c) for c in profile.metadata.machine_configurations]
    configurations[0]['specifications'] = {
        'gvw_kg': {'value': 3500, 'status': 'confirmed', 'provenance': 'Hypothetical regression fixture; not a vehicle approval'},
        'battery': {'value': '48 V; 20 kWh', 'status': 'user_provided', 'provenance': 'Hypothetical regression fixture'},
        'motor': {'value': '24 kW; 80 rpm', 'status': 'estimated', 'provenance': 'Hypothetical regression fixture'},
    }
    return replace(profile, metadata=replace(profile.metadata, machine_configurations=tuple(configurations)))


def test_configuration_specifications_cannot_leak_after_machine_or_profile_change():
    profile = configured_profile()
    selected = resolve_report_metadata('RoboSprayer_3500Kg_Electric.csv', profile)
    assert selected.machine_specifications['battery']['value'] == '48 V; 20 kWh'
    assert selected.machine_specifications['motor']['status'] == 'estimated'
    for name in ('RoboSprayer_8500Kg_Electric.csv', 'RoboSprayer_8500Kg_Electric_3500Kg_Chem.csv',
                 'Other_3500Kg_Electric.csv', 'RoboSprayer_3500Kg_Hybrid.csv', 'Unknown.csv'):
        metadata = resolve_report_metadata(name, profile)
        assert metadata.machine_specifications['battery']['value'] is None
        assert '48 V' not in str(machine_metadata_pills(metadata))
    # A title rename preserves actual source identity; a new source never borrows it.
    renamed = resolve_report_metadata('RoboSprayer_3500Kg_Electric.csv', profile, machine_name_override='Display Alias')
    assert renamed.machine_name == 'Display Alias'
    assert renamed.machine_specifications == selected.machine_specifications
    assert any('confirm identity' in warning for warning in renamed.warnings)
    with pytest.raises(ConfigurationError, match='mass'):
        resolve_report_metadata('RoboSprayer_8500Kg_Electric.csv', profile,
                                reporting_inputs={'machine_configuration_id': 'robosprayer_v5_3500kg'})
    with pytest.raises(ConfigurationError, match='mass'):
        resolve_report_metadata('RoboSprayer_8500Kg_Electric_3500Kg_Chem.csv', profile,
                                reporting_inputs={'machine_configuration_id': 'robosprayer_v5_3500kg'})


@pytest.mark.parametrize('specifications', [
    {'battery': {'value': '50 kWh', 'status': 'confirmed', 'provenance': ''}},
    {'gvw_kg': {'value': float('nan'), 'status': 'confirmed', 'provenance': 'test'}},
    {'gvw_kg': {'value': True, 'status': 'confirmed', 'provenance': 'test'}},
    {'generator': {'value': '80 kW'}},
    {'unknown': {'value': 'x', 'status': 'confirmed', 'provenance': 'test'}},
])
def test_untraceable_or_invalid_specifications_are_rejected(specifications):
    with pytest.raises(ConfigurationError):
        validate_specifications(specifications)


@pytest.fixture(scope='module', params=CASES, ids=['electric3500', 'electric8500', 'hybrid', 'diesel'])
def generated_machine_report(request, tmp_path_factory):
    profile_id, stem, *_ = request.param
    source = require_private_reference_file(Path('reference_files') / (stem + '.csv'), stem)
    return generate_reporting_profile_engineering_report(source, BASE / (profile_id + '.yaml'),
               tmp_path_factory.mktemp('machine_' + profile_id), ImportOptions(strict=False))


def test_every_slide_and_workbook_use_same_source_and_machine_metadata(generated_machine_report):
    result = generated_machine_report
    metadata = result.report_metadata
    prs = Presentation(result.presentation_path)
    assert prs.core_properties.title == metadata.report_title
    assert metadata.machine_name in prs.slides[0].shapes[3].text
    expected = list(machine_metadata_pills(metadata))
    assert [prs.slides[0].shapes[i].text for i in (6, 8, 10, 12, 14)] == expected
    assert [prs.slides[1].shapes[i].text for i in (5, 7, 9, 11, 13)] == expected
    text = '\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    for stale in ('GVW 3,500 kg', '12 kph field', 'Battery 50 kW', 'Motor 63 rpm'):
        assert stale not in text
    assert 'Not specified' in text
    definitions = result.powerpoint_result.powerpoint_result.config.slides
    for slide, definition in zip(prs.slides, definitions):
        notes = json.loads(slide.notes_slide.notes_text_frame.text)
        assert notes['machine_specifications'] == metadata.machine_specifications
        assert notes['machine_name'] == metadata.machine_name
        assert notes['powertrain'] == metadata.powertrain_name
        assert notes['source_filename'] == metadata.source_filename
        assert notes['source_sha256'] == result.excel_result.dataset.quality.source_sha256
        assert notes['plots'] == list(definition.plot_ids)
        for key, item in notes['statistics'].items():
            stat = next(s for s in result.powerpoint_result.powerpoint_result.statistics_result.statistics if s.statistic_id == key)
            assert item['value'] == stat.value if np.isfinite(stat.value) else item['value'] is None
    wb = load_workbook(result.report_path, data_only=True)
    try:
        fields = {r[0]: r[1] for r in wb['Metadata'].iter_rows(values_only=True) if r[0]}
        assert json.loads(fields['Machine specifications']) == metadata.machine_specifications
        agro = result.excel_result.statistics_result.reporting_agro
        rows = {r[0]: r for r in wb['Statistics'].iter_rows(values_only=True) if r[0]}
        assert rows['agro_discharge_rate_kg_per_km'][2] == pytest.approx(agro.net_mass_rate.value) if agro.net_mass_rate.value is not None else rows['agro_discharge_rate_kg_per_km'][2] == 'N/A'
    finally:
        wb.close()
    for file in (result.excel_result.manifest_path, result.powerpoint_result.manifest_path):
        manifest = json.loads(file.read_text())
        assert manifest['machine_specifications'] == metadata.machine_specifications
        assert manifest['configuration_detection'] == metadata.configuration_detection


def test_static_profile_pills_are_never_a_specification_authority():
    from types import SimpleNamespace
    from vsm_postprocessing.profile_powerpoint_report_engine import _cover_metadata_items
    metadata = resolve_report_metadata('NewMachine_8500Kg_Electric_20kph.csv')
    fake = SimpleNamespace(report_metadata=metadata, profile=SimpleNamespace(
        presentation=SimpleNamespace(metadata_pills=('WRONG MACHINE',) * 5)))
    assert _cover_metadata_items(fake) == machine_metadata_pills(metadata)


def test_configured_specifications_reach_slides_and_notes_across_profiles(generated_machine_report, tmp_path):
    from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
    result = generated_machine_report.excel_result
    metadata = result.report_metadata
    # Explicitly hypothetical, configuration-scoped values distinct from source
    # maxima. Propagation must never replace a measured KPI with a component rating.
    tokens = [str(int(metadata.machine_specifications['gvw_kg']['value'])) + 'kg']
    configuration = {'configuration_id': 'regression_only', 'machine_names': [metadata.machine_name],
        'identity_tokens': tokens, 'powertrains': [metadata.powertrain_name.lower()], 'reporting_inputs': {},
        'specifications': {key: {'value': value, 'status': 'confirmed', 'provenance': 'Hypothetical regression fixture'}
                           for key, value in {'battery': '42 kWh', 'motor': '61 kW', 'generator': '67 kW', 'engine': '72 kW'}.items()}}
    profile = replace(result.profile, metadata=replace(result.profile.metadata, machine_configurations=(configuration,)))
    selected = resolve_report_metadata(metadata.source_filename, profile)
    copy = replace(result, profile=profile, report_metadata=selected)
    deck = build_profile_powerpoint_report(copy, tmp_path / 'configured')
    prs = Presentation(deck.presentation_path)
    assert [prs.slides[0].shapes[i].text for i in (6, 8, 10, 12, 14)] == list(machine_metadata_pills(selected))
    text = '\n'.join(s.text for s in prs.slides[1].shapes if s.has_text_frame)
    if metadata.powertrain_name == 'Diesel':
        assert '72 kW' in text and '67 kW' in text
    else:
        assert '42 kWh' in text and '61 kW' in text
        if metadata.powertrain_name == 'Hybrid':
            assert '67 kW' in text and '72 kW' in text
    for slide in prs.slides:
        assert json.loads(slide.notes_slide.notes_text_frame.text)['machine_specifications'] == selected.machine_specifications
    original = {s.definition.statistic_id: s.value for s in result.statistics_result.statistics}
    assert original == {s.definition.statistic_id: s.value for s in copy.statistics_result.statistics}


def test_ui_manual_specifications_reset_when_source_or_profile_changes(monkeypatch):
    from streamlit.testing.v1 import AppTest
    from vsm_postprocessing import ui_app
    from vsm_postprocessing.ui_config import get_reporting_profile_definition
    monkeypatch.setattr(ui_app, '_machine_test_profiles', {
        kind: get_reporting_profile_definition(ui_app.PROJECT_ROOT, 'robosprayer_' + kind)
        for kind in ('electric', 'hybrid')}, raising=False)
    app = AppTest.from_string('''
import streamlit as st
from vsm_postprocessing import ui_app
source = st.selectbox('Test source', ['RoboSprayer_3500Kg_Electric.csv', 'RoboSprayer_8500Kg_Electric.csv'])
kind = st.selectbox('Test profile', ['electric', 'hybrid'])
st.session_state['settings'] = ui_app._render_reporting_metadata_inputs(st, ui_app._machine_test_profiles[kind], source)
''').run()
    assert not app.exception
    next(w for w in app.text_input if w.label == 'Battery pack mass (kg)').set_value('450').run()
    next(w for w in app.checkbox if w.label == 'Override automatically detected reporting information').check().run()
    assert not any('source configuration' in w.label or 'interpretation' in w.label for w in app.checkbox)
    assert app.session_state['settings']['battery_mass_kg'] == 450
    next(w for w in app.selectbox if w.label == 'Test source').select('RoboSprayer_8500Kg_Electric.csv').run()
    assert not app.exception
    assert app.session_state['settings'] == {}
    assert next(w for w in app.text_input if w.label == 'Battery pack mass (kg)').value == ''
    assert not next(w for w in app.checkbox if w.label == 'Override automatically detected reporting information').value
    next(w for w in app.text_input if w.label == 'Battery pack mass (kg)').set_value('1000').run()
    next(w for w in app.selectbox if w.label == 'Test profile').select('hybrid').run()
    assert not app.exception
    assert app.session_state['settings'] == {}
