"""Rolling resistance is authoritative; soil components remain diagnostics."""
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_to_tuple
from pptx import Presentation
import pytest

from test_sergio_october_15 import csv_for, dataset_for, profile_for, profile_path, remove_channel
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report, _ordered_profile_plots
from vsm_postprocessing.profile_math import calculate_profile_math_channels
from vsm_postprocessing.profile_powerpoint_report_engine import build_profile_powerpoint_report
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.report_profile import load_reporting_profile

POWER = "total_rolling_resistance_power"
ENERGY = "tyre_rolling_resistance_energy_accumulated"
RETIRED = {"tyre_total_resistance_" + suffix for suffix in
           ("power", "energy_kwh", "energy_wh", "energy_accumulated")}


@pytest.mark.parametrize("kind", ("electric", "hybrid", "diesel"))
@pytest.mark.parametrize("current", (False, True))
def test_all_profiles_use_rolling_without_duplicate_channels(kind, current):
    path = profile_path(kind)
    profile = load_reporting_profile(path if current else path.parent.parent / path.name)
    assert not RETIRED.intersection(profile.math_by_semantic_name())
    if kind != "diesel":
        assert profile.statistics_by_id()["tyre_total_energy_accumulated_last"].target == ENERGY
        assert [s.semantic_name for s in profile.plots_by_id()["tyres_energy_consumption"].series] == [POWER, ENERGY]


@pytest.mark.parametrize("kind", ("electric", "hybrid"))
@pytest.mark.parametrize("diagnostics", ("zero", "large", "missing"))
def test_soil_diagnostics_cannot_change_tyre_consumption(kind, diagnostics):
    data = dataset_for(kind, (17, 19, 30, 100))
    expected = np.array([10., -20., 50., -80.])
    for wheel in ("fl", "fr", "rl", "rr"):
        data.values[:, data.channel_index(f"tyre_rollingresistancepower_{wheel}")] = expected / 4
        for mechanism in ("bulldozing", "compaction"):
            name = f"tyre_{mechanism}resistancepower_{wheel}"
            if diagnostics == "missing":
                remove_channel(data, name)
            else:
                data.values[:, data.channel_index(name)] = 0 if diagnostics == "zero" else 1000
    profile = profile_for(kind)
    result = calculate_profile_math_channels(data, profile)
    values = result.values_by_semantic_name
    np.testing.assert_array_equal(values[POWER], expected)
    energy = np.r_[0, np.cumsum((expected[:-1] + expected[1:]) * np.diff([17, 19, 30, 100]) / 7200)]
    np.testing.assert_allclose(values[ENERGY], energy, rtol=3e-14, atol=1e-12)
    stats = calculate_profile_statistics(data, profile, math_result=result)
    assert next(s.value for s in stats.statistics if s.definition.statistic_id == "tyre_total_energy_accumulated_last") == pytest.approx(energy[-1])
    for mechanism in ("bulldozing", "compaction"):
        name = f"total_{mechanism}_resistance_power"
        assert (name in values) == (diagnostics != "missing")
        if diagnostics != "missing":
            np.testing.assert_array_equal(values[name], 0 if diagnostics == "zero" else 4000)


@pytest.fixture(scope="module", params=("electric", "hybrid", "diesel", "hybrid_no_aux"))
def reports(request, tmp_path_factory):
    kind = request.param.removesuffix("_no_aux")
    folder = tmp_path_factory.mktemp("rolling_" + request.param)
    data = dataset_for(kind, (0, .5, 2, 5))
    if request.param.endswith("_no_aux"):
        remove_channel(data, "auxiliary_high_voltage_power_consumption")
    report = generate_profile_excel_report(csv_for(data, folder / "source.csv"), profile_path(kind), folder / "excel")
    ppt = build_profile_powerpoint_report(report, folder / "powerpoint")
    return kind, report, ppt


def test_generated_reports_use_correct_channels_formulas_and_metrics(reports):
    kind, report, ppt = reports
    values = report.plotting_result.values_by_semantic_name
    columns = {c.channel_id: i for i, c in enumerate(report.report_channels, 1)}
    assert len(columns) == report.report_channel_count
    assert not RETIRED.intersection(columns)
    book = load_workbook(report.report_path)
    cached = load_workbook(report.report_path, data_only=True)
    try:
        for sheet in book:
            for row in sheet:
                for cell in row:
                    assert cell.data_type != "e" and "#REF!" not in str(cell.value)
        for item in report.statistics_result.kpis:
            row = next(r for r in book["Statistics"].iter_rows() if r[0].value == item.definition.kpi_id)
            assert row[2].data_type == "f"
            assert cached["Statistics"][row[2].coordinate].value == pytest.approx(item.value)
        if kind == "diesel":
            return
        sheet = book.worksheets[0]
        plots = report.plotting_result.rendered_plots
        index = next(i for i, p in enumerate(_ordered_profile_plots(plots)) if p.plot_id == "tyres_energy_consumption")
        chart = sheet._charts[index]
        refs = [s.yVal.numRef.f for c in chart._charts for s in c.ser]
        assert [range_to_tuple(ref)[1][0] for ref in refs] == [columns[POWER], columns[ENERGY]]
        statistic = next(r for r in book["Statistics"].iter_rows() if r[0].value == "tyre_total_energy_accumulated_last")
        assert statistic[4].data_type == "f"
        assert f"{get_column_letter(columns[ENERGY])}5:" in statistic[4].value.replace("$", "")
        assert cached["Statistics"][statistic[4].coordinate].value == pytest.approx(values[ENERGY][-1])
        if "auxiliary_high_voltage_power_consumption" not in report.resolution.resolved:
            assert "total_energy_consumption_tyres_aux" not in {k.definition.kpi_id for k in report.statistics_result.kpis}
            assert "total_auxiliary_power" not in values
        for mechanism in ("bulldozing", "compaction"):
            assert f"tyre_soil_{mechanism}" in {p.plot_id for p in plots}
            assert f"soil_{mechanism}" in {s.slide_id for s in ppt.powerpoint_result.config.slides}
        metrics = {m.statistic_id: m for m in ppt.powerpoint_result.statistics_result.statistics}
        assert metrics["tyre_total_energy_accumulated_last"].value == pytest.approx(values[ENERGY][-1])
        if kind == "electric":
            assert metrics["tyre_total_max_power"].value == max(values[POWER])
            assert metrics["tyre_total_total_energy"].value == pytest.approx(values[ENERGY][-1])
            assert metrics["tyre_total_energy_per_km"].value == pytest.approx(values[ENERGY][-1] / .3)
            slide_index = next(i for i, s in enumerate(ppt.powerpoint_result.config.slides) if s.slide_id == "tyre_total")
            slide = Presentation(ppt.presentation_path).slides[slide_index]
            assert "Tyre Rolling Resistance Power & Energy" in " ".join(s.text for s in slide.shapes if s.has_text_frame)
    finally:
        book.close(); cached.close()
