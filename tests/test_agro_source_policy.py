"""Source-policy eligibility must not depend on machine name or load magnitude."""
from dataclasses import replace
import json
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook
from pptx import Presentation

from vsm_postprocessing.agro_reporting import calculate_agro_kpis, calculate_reporting_kpis
from vsm_postprocessing.errors import ConfigurationError
from vsm_postprocessing.reporting_inputs import ReportingInputs
from vsm_postprocessing.report_metadata import resolve_report_metadata
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.report_selection import ProfileReportSelection
from vsm_postprocessing.ui_config import generate_reporting_profile_engineering_report
from test_sergio_october_15 import dataset_for, profile_for, profile_path, csv_for, remove_channel
from test_profile_math import _channel


def policy(kind='force_inference', **overrides):
    return ReportingInputs(agro_mass_source_policy=kind,
        agro_mass_source='Chemical contents' if kind == 'direct_mass' else 'Rear vertical load',
        agro_mass_source_provenance='Explicit synthetic source-mapping assumption; not a physical approval', **overrides)


@pytest.mark.parametrize('length,load', [(2, 37), (7, 813), (41, 7921)])
@pytest.mark.parametrize('kind', ['electric', 'hybrid', 'diesel'])
def test_unknown_machine_uses_profile_source_policy(kind, length, load):
    data = dataset_for(kind, tuple(range(length)))
    expected = np.linspace(load, load / 5, length)
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81 * expected
    profile = profile_for(kind)
    # A generic unconfirmed configuration must remain distinct from the
    # explicitly approved Astauto tank-load configuration inherited above.
    profile = replace(profile, metadata=replace(profile.metadata, reporting_inputs=replace(
        profile.metadata.reporting_inputs, hitch_force_is_remaining_agro_mass=False,
        agro_mass_provenance=None, agro_mass_confirmation_source=None)))
    for filename in ('Future_OrchardCarrier_X.csv', 'CompletelyUnknown_17234Kg_' + kind + '.csv'):
        metadata = resolve_report_metadata(filename, profile)
        assert metadata.configuration_detection == 'unmatched'
        assert not metadata.reporting_inputs.hitch_force_is_remaining_agro_mass
        result = calculate_reporting_kpis(calculate_profile_statistics(data, profile), metadata)
        assert result.capacity.value == pytest.approx(load)
        assert result.capacity.status == result.net_mass_rate.status == 'inferred'
        assert result.capacity.interpretation_status == 'assumed'
        assert result.net_mass_rate.value == pytest.approx((load - load / 5) / .3)
        np.testing.assert_allclose(result.mass_values, expected)
        assert result.capacity.provenance == metadata.reporting_inputs.agro_mass_source_provenance
        assert 'not nominal' in result.capacity.definition and 'physical proof' in result.capacity.definition


@pytest.mark.parametrize('values,unit,reason', [
    (None, 'N', 'missing'), ([0, 0], 'N', 'all-zero'),
    ([-981, -490.5], 'kg', 'unit mismatch'),
    ([-981, np.nan], 'N', 'nonfinite'), ([-981, np.inf], 'N', 'nonfinite'),
    ([981, 490.5], 'N', 'negative'), ([981, 0], 'N', 'negative'),
    ([-981], 'N', 'insufficient'),
])
def test_unsupported_force_is_unavailable_with_reason(values, unit, reason):
    result = calculate_agro_kpis([0, 1000], 'm', inputs=policy(), force_N=values, force_unit=unit)
    assert result.capacity.value is None and reason in result.capacity.reason
    assert result.rate.value is None and result.mass_rate.value is None


def test_source_mapping_is_required_and_confidence_is_independent():
    unconfigured = calculate_agro_kpis([0, 1000], 'm', inputs=ReportingInputs(), force_N=[-981, -490.5], force_unit='N')
    assert unconfigured.capacity.value is None and 'policy not configured' in unconfigured.capacity.reason
    # Physical confirmation cannot enable an absent/disabled numerical mapping.
    no_policy = calculate_agro_kpis([0, 1000], 'm', inputs=ReportingInputs(
        hitch_force_is_remaining_agro_mass=True, agro_mass_provenance='Explicit confidence only'),
        force_N=[-981, -490.5], force_unit='N')
    assert no_policy.capacity.value is None and 'policy not configured' in no_policy.capacity.reason
    assert no_policy.capacity.interpretation_status == 'unconfigured'
    plain = calculate_agro_kpis([0, 1000], 'm', inputs=policy(), force_N=[-981, -490.5], force_unit='N')
    confirmed = calculate_agro_kpis([0, 1000], 'm', inputs=policy(
        hitch_force_is_remaining_agro_mass=True, agro_mass_provenance='Synthetic confirmation fixture'),
        force_N=[-981, -490.5], force_unit='N')
    assert plain.capacity.value == confirmed.capacity.value == 100
    assert plain.capacity.status == confirmed.capacity.status == 'inferred'
    assert plain.capacity.interpretation_status == 'assumed'
    assert confirmed.capacity.interpretation_status == 'user_confirmed'


@pytest.mark.parametrize('mass', [[0, 0], [641, 200], [100, 100]])
def test_direct_simulation_mass_is_not_divided_by_gravity(mass):
    result = calculate_agro_kpis([0, 1000], 'm', inputs=policy('direct_mass'), mass_kg=mass, mass_unit='kg')
    assert result.capacity.value == max(mass)
    assert result.capacity.status == 'simulated'
    assert 'measured' in result.capacity.definition and 'not asserted' in result.capacity.definition
    assert result.mass_rate.value == pytest.approx((mass[0] - mass[-1]) / 1000)
    assert result.rate.value is None and 'density required' in result.rate.reason


def test_questionable_force_maximum_is_explicit_and_never_clamped():
    result = calculate_agro_kpis([0, 1000], 'm', inputs=policy(), force_N=-9.81*np.array([100, -5]), force_unit='N')
    assert result.capacity.value == 100 and result.capacity.quality_status == 'questionable'
    assert result.mass_values[-1] == -5
    assert result.rate.value is result.mass_rate.value is None
    assert any('negative' in warning for warning in result.trace['warnings'])


@pytest.mark.parametrize('bad', [
    {'agro_mass_source_policy': 'force_inference'},
    {'agro_mass_source_policy': 'direct_mass', 'agro_mass_source': 'Mass'},
    {'agro_mass_source_policy': 'guess_from_filename'},
])
def test_policy_configuration_requires_mapping_and_provenance(bad):
    with pytest.raises(ConfigurationError):
        ReportingInputs.from_mapping(bad)


@pytest.mark.parametrize('kind', ['electric', 'hybrid', 'diesel'])
@pytest.mark.parametrize('source_kind', ['force_inference', 'direct_mass'])
def test_future_machine_cross_format_source_values_and_provenance(tmp_path, kind, source_kind):
    data = dataset_for(kind, (0, 1, 2, 3, 4, 5, 6))
    masses = np.linspace(871, 57, 7)
    settings = {}
    if source_kind == 'direct_mass':
        # The configured kg source deliberately disagrees with the old force MATH
        # channel; no renderer/formula may quietly use that channel's mass.
        data.values[:, data.channel_index('agrochemical_discharge_force')] = 0
        data.channels.append(_channel('chemical_mass', 'Chemical contents', 'kg', len(data.channels)+1))
        data.values = np.column_stack((data.values, masses))
        settings = {'agro_mass_source_policy': 'direct_mass', 'agro_mass_source': 'Chemical contents',
                    'agro_mass_source_provenance': 'Explicit test simulation mass mapping'}
    else:
        data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81*masses
    source = csv_for(data, tmp_path / 'Unseen_Prototype_R67.csv')
    agro_plot = next(p.plot_id for p in profile_for(kind).plots
                     if any(s.semantic_name == 'agrochemical_discharge' for s in p.series))
    result = generate_reporting_profile_engineering_report(source, profile_path(kind), tmp_path/'report',
        reporting_inputs=settings, selection=ProfileReportSelection(plots=(agro_plot,)))
    agro = result.excel_result.statistics_result.reporting_agro
    assert agro.capacity.value == pytest.approx(871)
    assert agro.net_mass_rate.value == pytest.approx((871-57)/.3)
    np.testing.assert_allclose(result.excel_result.plotting_result.values_by_semantic_name['reporting_agro_mass'], masses)
    series = result.excel_result.plotting_result.series_summaries[agro_plot]
    canonical = next(s for s in series if s.semantic_name == 'reporting_agro_mass')
    assert canonical.max_value == pytest.approx(871) and canonical.min_value == pytest.approx(57)
    rendered = next(p for p in result.excel_result.plotting_result.rendered_plots if p.plot_id == agro_plot)
    assert any('Remaining Mass' in label for label in rendered.legend_labels)
    assert all('Capacity' not in label for label in rendered.legend_labels)
    wb = load_workbook(result.report_path, data_only=True)
    formulas = load_workbook(result.report_path, data_only=False)
    try:
        location = result.excel_result.workbook_structure['agro_source']['mass']
        mass_sheet, mass_column = location['sheet'], location['column']
        np.testing.assert_allclose([wb[mass_sheet][f'{mass_column}{row}'].value for row in range(5, 12)], masses)
        assert 'Agro Mass Source' not in wb
        if source_kind == 'force_inference' and kind in {'electric', 'hybrid'}:
            assert 'Agro Mass Source' not in wb
        rows = {r[0]: r for r in wb['Statistics'].values if r[0]}
        assert rows['agro_capacity_observed_kg'][2] == pytest.approx(871)
        assert rows['agro_capacity_observed_kg'][4] == agro.capacity.status
        assert rows['agro_capacity_observed_kg'][5] == agro.capacity.source
        assert rows['agro_discharge_rate_kg_per_km'][2] == pytest.approx(agro.net_mass_rate.value)
        frows = {r[0].value: r for r in formulas['Statistics'].rows if r[0].value}
        assert mass_sheet.replace("'", "''") in frows['agro_capacity_observed_kg'][2].value
        assert mass_sheet.replace("'", "''") in frows['agro_discharge_rate_kg_per_km'][2].value
        assert formulas[mass_sheet][f'{mass_column}5'].data_type == 'f'
        charts = formulas.worksheets[0]._charts
        refs = [s.yVal.numRef.f for chart in charts for part in chart._charts for s in part.series]
        assert any(mass_sheet.replace("'", "''") in ref and mass_column in ref for ref in refs)
    finally:
        wb.close(); formulas.close()
    manifest = json.loads(result.manifest_path.read_text())
    assert manifest['agro_reporting'] == agro.to_dict()
    prs = Presentation(result.presentation_path)
    text = '\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert '871.00 kg' in text and agro.capacity.label.upper() in text
    for slide in prs.slides:
        notes = json.loads(slide.notes_slide.notes_text_frame.text)
        assert notes['agro_reporting'] == agro.to_dict()
        assert notes['agro_reporting']['capacity']['provenance'] == agro.capacity.provenance


def test_missing_direct_mass_does_not_fall_back_to_plausible_force():
    data = dataset_for('electric')
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81*np.array([99, 90, 50, 0])
    metadata = resolve_report_metadata('Future.csv', profile_for('electric'), reporting_inputs={
        'agro_mass_source_policy': 'direct_mass', 'agro_mass_source': 'Uninstalled chemical sensor',
        'agro_mass_source_provenance': 'Explicit unavailable mapping'})
    result = calculate_reporting_kpis(calculate_profile_statistics(data, profile_for('electric')), metadata)
    assert result.capacity.value is result.net_mass_rate.value is None
    assert 'missing' in result.capacity.reason
