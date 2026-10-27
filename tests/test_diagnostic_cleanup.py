"""Resolved provenance is quiet; unresolved engineering conditions stay visible."""
from copy import deepcopy
from dataclasses import replace
import json
import pickle

import numpy as np
from openpyxl import load_workbook
from pptx import Presentation
import pytest

from vsm_postprocessing.diagnostics import Diagnostic, diagnostic_records, visible_diagnostics
from vsm_postprocessing.agro_reporting import calculate_reporting_kpis
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.report_metadata import resolve_report_metadata
from vsm_postprocessing.ui_app import (_render_diagnostics, _render_profile_validation_summary,
                                      _render_profile_report_completion)
from vsm_postprocessing.ui_config import (generate_reporting_profile_engineering_report,
                                        validate_reporting_profile_source)
from test_sergio_october_15 import dataset_for, profile_for, profile_path, csv_for, remove_channel
from test_ui_config import _FakeStreamlit


def test_diagnostic_policy_survives_existing_text_json_copy_and_session_consumers():
    message = Diagnostic('Applied interpretation', code='approved_unit_interpretation',
                         severity='info', user_visible=False)
    assert isinstance(message, str) and 'interpretation' in message
    assert json.loads(json.dumps([message])) == ['Applied interpretation']
    for restored in (deepcopy(message), pickle.loads(pickle.dumps(message))):
        assert restored.to_dict() == message.to_dict()
    st = _FakeStreamlit()
    _render_diagnostics(st, [message, 'Unclassified engineering issue',
                            Diagnostic('Invalid input', code='invalid_input', severity='error')])
    assert st.messages == [('warning', 'Unclassified engineering issue'), ('error', 'Invalid input')]


@pytest.mark.parametrize('kind', ['electric', 'hybrid', 'diesel'])
def test_explicit_astauto_tank_mapping_is_confirmed_independent_of_machine_name(kind):
    profile, data = profile_for(kind), dataset_for(kind, (0, 1, 2, 3))
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81 * np.array([100, 80, 60, 40])
    stats = calculate_profile_statistics(data, profile)
    for filename in ('Unknown.csv', 'Unlisted_Manufacturer_' + kind + '.csv'):
        metadata = resolve_report_metadata(filename, profile)
        agro = calculate_reporting_kpis(stats, metadata)
        assert metadata.configuration_detection == 'unmatched'
        assert not visible_diagnostics(metadata.warnings)
        assert agro.capacity.interpretation_status == 'user_confirmed'
        assert agro.capacity.status == 'inferred'
        assert agro.trace['mass_source']['confirmation_applied']
        assert '27 October 2026' in agro.trace['mass_source']['physical_validation_provenance']
        assert agro.capacity.value == pytest.approx(100)
        assert not any(diagnostic_records([m])[0]['code'] == 'force_interpretation_unconfirmed'
                       for m in visible_diagnostics(agro.trace['warnings']))
        np.testing.assert_array_equal(agro.mass_values, np.array([100., 80., 60., 40.]))


@pytest.mark.parametrize('mode', ['different_source', 'unconfirmed_profile', 'missing_authority'])
def test_physical_confirmation_does_not_leak_to_unapproved_sources_or_configurations(mode):
    profile, data = profile_for('electric'), dataset_for('electric', (0, 1, 2))
    if mode == 'different_source':
        index = data.channel_index('agrochemical_discharge_force')
        data.channels[index] = replace(data.channels[index], source_name='Different vertical force')
        profile = replace(profile, raw_channels=tuple(
            replace(c, source_name='Different vertical force', aliases=())
            if c.semantic_name == 'agrochemical_discharge_force' else c for c in profile.raw_channels))
    else:
        settings = profile.metadata.reporting_inputs
        settings = replace(settings, **({'hitch_force_is_remaining_agro_mass': False}
                           if mode == 'unconfirmed_profile' else {'agro_mass_provenance': None}))
        profile = replace(profile, metadata=replace(profile.metadata, reporting_inputs=settings))
    agro = calculate_reporting_kpis(calculate_profile_statistics(data, profile),
                                   resolve_report_metadata('Unknown.csv', profile))
    assert agro.capacity.interpretation_status == 'assumed'
    assert any(getattr(d, 'code', '') == 'force_interpretation_unconfirmed'
               for d in visible_diagnostics(agro.trace['warnings']))


@pytest.mark.parametrize('values,reason', [([100, -5, -10], 'negative'),
                                         ([100, 10, 90], 'refill'),
                                         ([0, 0, 0], 'all-zero')])
def test_confirmation_never_overrides_data_quality(values, reason):
    profile, data = profile_for('hybrid'), dataset_for('hybrid', (0, 1, 2))
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81 * np.array(values)
    agro = calculate_reporting_kpis(calculate_profile_statistics(data, profile),
                                   resolve_report_metadata('Unknown.csv', profile))
    assert any(reason in d for d in visible_diagnostics(agro.trace['warnings']))
    if reason == 'negative':
        assert agro.net_mass_rate.value is None
        assert any(getattr(d, 'code', '') == 'force_sign_change' for d in agro.trace['warnings'])
        np.testing.assert_array_equal(agro.mass_values, values)
    assert not any(getattr(d, 'code', '') == 'force_interpretation_unconfirmed'
                   for d in visible_diagnostics(agro.trace['warnings']))


@pytest.mark.parametrize('unit', ['-', 'N'])
def test_validation_keeps_unsupported_units_visible_and_approved_units_quiet(tmp_path, unit):
    data = dataset_for('electric', (0, 1, 2))
    name = 'auxiliary_high_voltage_power_consumption'
    index = data.channel_index(name)
    data.channels[index] = replace(data.channels[index], unit=unit)
    source = csv_for(data, tmp_path / 'Unknown.csv')
    summary = validate_reporting_profile_source(source, profile_path('electric'))
    st = _FakeStreamlit()
    _render_profile_validation_summary(st, summary)
    if unit == '-':
        assert any(getattr(d, 'code', '') == 'approved_unit_interpretation' for d in summary.diagnostics)
        assert not any('Engineering unit interpretation' in message for _, message in st.messages)
    else:
        assert any(level == 'warning' and 'unit mismatch' in message for level, message in st.messages)


@pytest.mark.parametrize('kind', ['electric', 'hybrid', 'diesel'])
def test_quiet_resolved_notifications_retain_cross_format_provenance_and_real_warnings(tmp_path, kind):
    data = dataset_for(kind, (0, 1, 2, 3))
    data.values[:, data.channel_index('agrochemical_discharge_force')] = -9.81 * np.array([100, 80, 60, 40])
    if kind != 'diesel':
        index = data.channel_index('auxiliary_high_voltage_power_consumption')
        data.channels[index] = replace(data.channels[index], unit='-')
    result = generate_reporting_profile_engineering_report(csv_for(data, tmp_path / 'Unknown.csv'),
                                                          profile_path(kind), tmp_path / 'reports')
    st = _FakeStreamlit()
    _render_profile_report_completion(st, result)
    assert not any('Engineering unit interpretation' in m or 'identity unmatched' in m
                   or 'Configured force is treated' in m or 'Battery pack mass not specified' in m
                   for _, m in st.messages)
    assert any(level == 'warning' and 'density required' in m for level, m in st.messages)
    manifest = json.loads(result.excel_result.manifest_path.read_text())
    codes = {d['code']: d for d in manifest['diagnostic_records']}
    assert not codes['machine_identity_unmatched']['user_visible']
    assert not codes['simulation_mass_interpretation']['user_visible']
    assert manifest['configuration_detection'] == 'unmatched'
    assert manifest['agro_reporting']['capacity']['interpretation_status'] == 'user_confirmed'
    powerpoint_manifest = json.loads(result.powerpoint_result.manifest_path.read_text())
    assert powerpoint_manifest['diagnostic_records'] == manifest['diagnostic_records']
    book = load_workbook(result.report_path, read_only=True)
    try:
        fields = {r[0]: r[1] for r in book['Metadata'].iter_rows(values_only=True) if r[0]}
        assert any(d['code'] == 'machine_identity_unmatched' for d in json.loads(fields['Diagnostics']))
        assert '27 October 2026' in fields['Agro reporting diagnostics']
    finally:
        book.close()
    deck = Presentation(result.presentation_path)
    for slide in deck.slides:
        notes = json.loads(slide.notes_slide.notes_text_frame.text)
        assert notes['agro_reporting']['capacity']['interpretation_status'] == 'user_confirmed'
        assert '27 October 2026' in notes['reporting_inputs']['agro_mass_provenance']
        assert not any('identity unmatched' in s.text for s in slide.shapes if s.has_text_frame)


def test_missing_essential_channel_still_blocks_reporting(tmp_path):
    data = dataset_for('electric', (0, 1, 2))
    remove_channel(data, 'track_distance')
    source = csv_for(data, tmp_path / 'Unknown.csv')
    summary = validate_reporting_profile_source(source, profile_path('electric'))
    assert not summary.is_valid and summary.missing_required_count
    from vsm_postprocessing.errors import ExcelReportError
    with pytest.raises(ExcelReportError, match='Track_Distance'):
        generate_reporting_profile_engineering_report(source, profile_path('electric'), tmp_path / 'reports')
