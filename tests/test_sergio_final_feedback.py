"""Acceptance for exact cover metadata, tyre totals and Diesel report selection."""
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from pptx import Presentation
import pytest

from test_sergio_october_15 import csv_for, dataset_for, profile_for, profile_path, remove_channel
from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
from vsm_postprocessing.profile_math import calculate_profile_math_channels
from vsm_postprocessing.profile_powerpoint_report_engine import _add_cover_source_filename, build_profile_powerpoint_report
from vsm_postprocessing.report_profile import load_reporting_profile
from vsm_postprocessing.report_selection import ProfileReportSelection


@pytest.mark.parametrize("filename", (
    "Full_Size_Sprayer_8500Kg_Diesel_1000m_20kph_4WS_DriverPC_up_to_25pcGrad_SL_LSD_07.csv",
    "Long source (with spaces)_" + "configuration_" * 15 + "(final).xlsx",
    "W" * 249 + ".csv",
))
def test_cover_preserves_exact_filename_and_has_wrapping(filename):
    from types import SimpleNamespace
    prs = Presentation("reference_files/RoboSprayer_Electric_Report_Astauto_Colours.pptx")
    result = SimpleNamespace(report_metadata=SimpleNamespace(source_filename=filename))
    _add_cover_source_filename(prs.slides[0], result)
    shape = next(s for s in prs.slides[0].shapes if s.name == "VSM source filename")
    assert shape.text == "Source: " + filename
    assert shape.text_frame.word_wrap
    assert shape.text_frame.paragraphs[0].font.size.pt >= 12
    assert shape.text_frame.paragraphs[0]._p.pPr.get("latinLnBrk") == "1"
    assert shape.top + shape.height < prs.slides[0].shapes[5].top
    assert prs.slides[0].shapes[5].top + prs.slides[0].shapes[5].height < prs.slides[0].shapes[15].top


def test_tyre_consumption_uses_rolling_only_with_observed_time_integration():
    data, profile = dataset_for("electric", (17, 18, 123, 1817, 3617)), profile_for("electric")
    expected = np.zeros(5)
    components = []
    for j, mechanism in enumerate(("rolling", "bulldozing", "compaction"), 1):
        component = np.zeros(5)
        for i, wheel in enumerate(("fl", "fr", "rl", "rr"), 1):
            p = j * i * np.array([-2, 0, 8, -4, 1.0])
            data.values[:, data.channel_index(f"tyre_{mechanism}resistancepower_{wheel}")] = p
            component += p
        components.append(component)
        if mechanism == "rolling":
            expected += component
    math = calculate_profile_math_channels(data, profile)
    values = math.values_by_semantic_name
    t = data.values[:, data.channel_index("track_time")]
    energy = np.r_[0, np.cumsum((expected[:-1]+expected[1:])*np.diff(t)/7200)]
    np.testing.assert_array_equal(values["total_rolling_resistance_power"], expected)
    np.testing.assert_allclose(values["tyre_rolling_resistance_energy_accumulated"], energy)
    assert not any(name.startswith("tyre_total_resistance_") for name in values)
    np.testing.assert_array_equal(values["total_bulldozing_resistance_power"], components[1])
    np.testing.assert_array_equal(values["total_compaction_resistance_power"], components[2])
    assert np.any(np.diff(energy) < 0)


@pytest.mark.parametrize("kind", ("electric", "hybrid", "diesel"))
def test_xlsx_ingestion_passes_original_basename_to_generated_cover(kind, tmp_path):
    data = dataset_for(kind)
    source = tmp_path / "Source simulation (review)_with spaces.xlsx"
    book = Workbook()
    book.active.append([c.source_name for c in data.channels])
    book.active.append([c.unit or "" for c in data.channels])
    for row in data.values:
        book.active.append(row.tolist())
    book.save(source)
    original = source.read_bytes()
    selection = ProfileReportSelection(plots=(), math_channels=(), statistics=(), kpis=())
    report = generate_profile_excel_report(source, profile_path(kind), tmp_path / "excel", selection=selection)
    ppt = build_profile_powerpoint_report(report, tmp_path / "ppt", output_filename="different_report_name.pptx")
    prs = Presentation(ppt.presentation_path)
    assert any(s.has_text_frame and s.text == "Source: " + source.name for s in prs.slides[0].shapes)
    assert source.read_bytes() == original


@pytest.mark.parametrize("missing", ("tyre_rollingresistancepower_rr", "tyre_bulldozingresistancepower_fl", "tyre_compactionresistancepower_fr"))
def test_missing_diagnostic_does_not_disable_rolling_consumption(missing):
    data = dataset_for("electric")
    remove_channel(data, missing)
    result = calculate_profile_math_channels(data, profile_for("electric"))
    available = missing != "tyre_rollingresistancepower_rr"
    assert ("total_rolling_resistance_power" in result.values_by_semantic_name) == available
    assert ("tyre_rolling_resistance_energy_accumulated" in result.values_by_semantic_name) == available


@pytest.mark.parametrize("legacy", (False, True))
def test_diesel_raw_mapping_retained_but_report_entries_excluded(legacy):
    path = Path("config/report_profiles/full_size_sprayer_diesel.yaml") if legacy else profile_path("diesel")
    profile = load_reporting_profile(path)
    assert profile.raw_by_semantic_name()["engine_load"].source_name == "Engine_Load"
    assert not profile.raw_by_semantic_name()["engine_load"].for_report
    assert profile.raw_by_semantic_name()["engine_torque"].for_report
    assert all(s.target != "engine_load" for s in profile.statistics)
    assert all(s.semantic_name != "engine_load" for p in profile.plots for s in p.series)
    assert profile.presentation.slides_by_id()["diesel_controls"].plots == ("diesel_engine_throttle",)


def test_diesel_actual_exports_keep_engine_torque_and_shift_formulas_safely(tmp_path):
    data = dataset_for("diesel")
    data.values[:, data.channel_index("engine_load")] = 18000
    data.values[:, data.channel_index("engine_torque")] = [0, 300, 936.59, -20]
    report = generate_profile_excel_report(csv_for(data, tmp_path / "source.csv"), profile_path("diesel"), tmp_path / "excel")
    assert "engine_load" not in {c.channel_id for c in report.report_channels}
    assert "engine_load" in report.resolution.resolved
    assert next(s.value for s in report.statistics_result.statistics if s.definition.statistic_id == "engine_torque_max") == 936.59
    book = load_workbook(report.report_path)
    try:
        assert all("Engine Load" not in str(c.value) and "#REF!" not in str(c.value)
                   for sheet in book for row in sheet for c in row)
        assert len(book.worksheets[0]._charts) == len(report.plotting_result.rendered_plots)
    finally:
        book.close()
    ppt = build_profile_powerpoint_report(report, tmp_path / "ppt")
    prs = Presentation(ppt.presentation_path)
    text = "\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert "Engine Load" not in text and "936.59 Nm" in text
    assert "Engine Throttle" in text
