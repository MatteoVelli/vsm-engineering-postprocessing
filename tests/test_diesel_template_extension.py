"""Final template selection/figure coverage and independent Diesel extension maths."""
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from conftest import require_private_reference_file
from test_diesel_profile import PROFILE_PATH, _source, _values
from vsm_postprocessing.importer import load_data_file
from vsm_postprocessing.profile_math import calculate_profile_math_channels
from vsm_postprocessing.report_profile import load_reporting_profile, normalize_name


def test_every_template_selection_row_has_a_distinct_profile_channel():
    template = require_private_reference_file(Path("reference_files/Diesel Template.xlsx"), "final Diesel template")
    profile = load_reporting_profile(PROFILE_PATH)
    definitions = (*profile.raw_channels, *profile.math_channels)
    overrides = {"Engine Power Required": "engine_power", "Generator Torque_2": "engine_auxiliary_torque_2"}
    seen = set()
    workbook = load_workbook(template, read_only=True)
    try:
        for source, report, kind, _ in workbook.worksheets[1].iter_rows(min_row=3, min_col=2, max_col=5, values_only=True):
            if source is None:
                continue
            if report.strip() in overrides:
                semantic = overrides[report.strip()]
            else:
                candidates = [d for d in definitions if d.channel_type == kind and d.semantic_name not in seen
                              and any(normalize_name(source.replace("Accomulated", "Accumulated")) == normalize_name(name)
                                      for name in (d.source_name, *getattr(d, "aliases", ())))]
                assert len(candidates) == 1, (source, report, kind, candidates)
                semantic = candidates[0].semantic_name
            assert semantic not in seen
            seen.add(semantic)
        assert len(seen) == 56
        template_plots = [p for p in profile.plots if p.reference_chart_number is not None]
        assert {p.reference_chart_number for p in template_plots} == set(range(1, 9))
        assert len(template_plots) == 8
        assert len(profile.plots) == 18
    finally:
        workbook.close()


def test_corner_power_auxiliary_sources_and_energy_use_independent_signed_inputs(tmp_path):
    profile = load_reporting_profile(PROFILE_PATH)
    dataset = load_data_file(_source(tmp_path, profile))
    for i, corner in enumerate(("FL", "FR", "RL", "RR"), 1):
        _values(dataset, "Wheel_RotationalSpeed_" + corner)[:] = [60 * i, 120 * i, 180 * i]
        _values(dataset, "DriveShaft_Torque_" + corner)[:] = [10 * i, -20 * i, 30 * i]
        _values(dataset, "Tyre_RollingResistancePower_" + corner)[:] = [i, 2 * i, 3 * i]
    _values(dataset, "Engine_AuxiliaryTorque_1")[:] = [10, -20, 30]
    _values(dataset, "Engine_AuxiliaryTorque_2")[:] = [40, 50, -60]
    _values(dataset, "HitchRear_Force_Z_VehicleCoordinates")[:] = [-9.80665, -19.6133, 9.80665]
    result = calculate_profile_math_channels(dataset, profile)
    values = result.values_by_semantic_name
    wheel_total = np.zeros(3)
    for i, corner in enumerate(("fl", "fr", "rl", "rr"), 1):
        expected = np.array([60 * i, 120 * i, 180 * i]) * [10 * i, -20 * i, 30 * i] * math.tau / 60000
        np.testing.assert_allclose(values["wheel_power_" + corner], expected, rtol=2e-15)
        wheel_total += expected
    np.testing.assert_allclose(values["wheel_power_total"], wheel_total)
    np.testing.assert_allclose(values["wheel_total_torque"], [100, -200, 300])
    np.testing.assert_allclose(values["auxiliary_power_1"], np.array([0, 600, 1200]) * [10, -20, 30] * math.tau / 60000)
    np.testing.assert_allclose(values["auxiliary_power_2"], np.array([0, 600, 1200]) * [40, 50, -60] * math.tau / 60000)
    np.testing.assert_allclose(values["total_auxiliary_power"], values["auxiliary_power_1"] + values["auxiliary_power_2"])
    np.testing.assert_allclose(values["agrochemical_discharge"], [1, 2, -1])
    np.testing.assert_allclose(values["total_rolling_resistance_power"], [10, 20, 30])
    np.testing.assert_allclose(values["tyre_rolling_resistance_energy_kwh"], np.array([10, 20, 60]) / 3600)
    np.testing.assert_allclose(values["tyre_rolling_resistance_energy_wh"], np.array([10, 20, 60]) / 3.6)
    np.testing.assert_allclose(values["tyre_rolling_resistance_energy_accumulated"], np.array([10, 30, 90]) / 3600)
    np.testing.assert_allclose(values["engine_energy_delivered"], values["engine_power"] * [1, 1, 2] / 3600)
    np.testing.assert_allclose(values["engine_mechanical_energy"], [0, math.pi / 3600, -math.pi / 3600])
