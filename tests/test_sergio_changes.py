from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation

from vsm_postprocessing.excel_report_engine import _build_native_scatter_chart, _profile_right_summary_items
from vsm_postprocessing.plotting_engine import PlotDefinition, PlotSeriesDefinition
from vsm_postprocessing.profile_statistics import calculate_profile_statistics
from vsm_postprocessing.profile_powerpoint_report_engine import (
    _add_reusable_content_layout, _add_astauto_logo, _canonical_metric_as_powerpoint_item,
)
from vsm_postprocessing.report_profile import load_reporting_profile, resolve_profile
from vsm_postprocessing.utils import source_report_filename, protect_source_path
from test_profile_statistics import _dataset, _channel


@pytest.mark.parametrize('profile_name', ['electric', 'hybrid'])
def test_auxiliary_max_uses_shared_profile_statistics(profile_name):
    profile = load_reporting_profile(f'config/report_profiles/robosprayer_{profile_name}.yaml')
    ids = {'auxiliary_energy_accumulated_max', 'total_auxiliary_power_max'}
    definitions = tuple(item for item in profile.statistics if item.statistic_id in ids)
    assert len(definitions) == 2
    assert all(item.operation == 'max' and item.nan_policy == 'omit' for item in definitions)
    power_definition = profile.math_by_semantic_name()['total_auxiliary_power']
    assert set(power_definition.dependencies) == {
        'electricsystem_auxiliary_lowvoltagepowerconsumption', 'auxiliary_high_voltage_power_consumption'}
    profile = replace(profile, statistics=definitions, raw_channels=(), kpis=())
    dataset = _dataset([], [[], [], [], [], []])
    signals = {'auxiliary_energy_consumption_accumulated': np.array([0, .4, 1.2, np.nan, .9]),
               'total_auxiliary_power': np.array([1, 4.5, 2, np.nan, 3])}
    math_result = SimpleNamespace(values_by_semantic_name=signals, calculated_channels=[
        _channel('test__math__' + d.target, d.target, d.unit) for d in definitions])
    result = calculate_profile_statistics(dataset, profile, resolve_profile(dataset, profile), math_result)
    expected = {'auxiliary_energy_accumulated_max': 1.2, 'total_auxiliary_power_max': 4.5}
    for metric in result.canonical_metrics:
        assert metric.value == pytest.approx(expected[metric.metric_id])
        assert metric.statistic == 'MAX'
        assert metric.omitted_sample_count == 1
        assert _canonical_metric_as_powerpoint_item(metric).value == metric.value
    assert {item['value'] for item in _profile_right_summary_items(result)} == {1.2, 4.5}


@pytest.mark.parametrize('suffix', ['.csv', '.xlsx'])
@pytest.mark.parametrize('profile_name', ['Electric', 'Hybrid'])
def test_source_names_preserved_and_input_protected(tmp_path, suffix, profile_name):
    stem = f'  abcdef123456_Machine  A--_{profile_name}_Test_05 '
    source = tmp_path / (stem + suffix)
    source.write_bytes(b'source bytes')
    for extension in ('.xlsx', '.pptx'):
        name = source_report_filename(source, extension)
        assert name == stem + extension
        destination = protect_source_path(tmp_path / name, source)
        assert destination != source
        destination.write_bytes(b'report bytes')
    assert source.read_bytes() == b'source bytes'
    assert source_report_filename('Test:01?.csv', '.xlsx') == 'Test_01_.xlsx'
    assert source_report_filename('CON.csv', '.xlsx') == '_CON.xlsx'


def test_filename_validators_preserve_valid_spaces():
    from vsm_postprocessing.excel_report_engine import _plain_xlsx_filename
    from vsm_postprocessing.powerpoint_report_engine import _pptx_filename
    assert _plain_xlsx_filename('  Test  Run .xlsx', 'output') == '  Test  Run .xlsx'
    assert _pptx_filename('  Test  Run .pptx') == '  Test  Run .pptx'


def test_ui_upload_and_download_preserve_source_name(tmp_path, monkeypatch):
    from vsm_postprocessing import ui_app
    monkeypatch.setattr(ui_app, 'UI_WORKSPACE', tmp_path)
    name = 'abcdef123456_Machine  A--_Test.csv'
    source = ui_app._persist_upload(SimpleNamespace(name=name, getvalue=lambda: b'input'))
    assert source.name == name
    assert source.parent != tmp_path
    paths = [tmp_path / source_report_filename(source, suffix) for suffix in ('.xlsx', '.pptx')]
    for path in paths:
        path.write_bytes(b'report')
    calls = []
    st = SimpleNamespace(download_button=lambda *args, **kwargs: calls.append(kwargs))
    result = SimpleNamespace(report_path=paths[0], presentation_path=paths[1],
                             excel_result=SimpleNamespace(plotting_result=SimpleNamespace(rendered_plots=[])))
    ui_app._render_profile_report_download(st, result)
    assert [call['file_name'] for call in calls] == [path.name for path in paths]


def test_native_chart_ranges_axes_and_readability(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Engineering data'
    for row in [(0, 2, 10), (1, 4, 20), (2, 3, 15)]:
        sheet.append(row)
    before = list(sheet.values)
    definition = PlotDefinition('road', 'Road Profile', 'distance', (
        PlotSeriesDefinition('gradient', 'primary', 'Road Gradient'),
        PlotSeriesDefinition('height', 'secondary', 'Road Height')),
        'road.png', x_label='Distance [km]', primary_y_label='Gradient [%]', secondary_y_label='Height [m]')
    chart = _build_native_scatter_chart(definition, report_sheet=sheet,
        channel_column_by_id={'distance': 1, 'gradient': 2, 'height': 3},
        data_start_row=1, data_end_row=3, chart_index=1)
    sheet.add_chart(chart, 'E2')
    path = tmp_path / 'native.xlsx'
    workbook.save(path)
    restored = load_workbook(path)
    assert list(restored.active.values) == before
    assert not restored.active._images
    chart = restored.active._charts[0]
    assert len(chart._charts) == 2
    for index, component in enumerate(chart._charts):
        series = component.series[0]
        assert series.xVal.numRef.f == "'Engineering data'!$A$1:$A$3"
        col = get_column_letter(index + 2)
        assert series.yVal.numRef.f == f"'Engineering data'!${col}$1:${col}$3"
        assert series.graphicalProperties.line.width == 31750
        assert component.y_axis.txPr.p[0].pPr.defRPr.sz == 1000
        assert component.y_axis.title.txPr.p[0].pPr.defRPr.sz == 1100
    assert chart._charts[1].y_axis.axPos == 'r'
    assert chart._charts[1].y_axis.crosses == 'max'
    assert chart.legend.txPr.p[0].pPr.defRPr.sz == 900


def test_legacy_plot_only_channel_is_editable_and_visible(tmp_path):
    from vsm_postprocessing.excel_report_engine import _native_chart_source_columns
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(5, 1, 0)
    sheet.cell(6, 1, 1)
    definition = PlotDefinition('power', 'Power', 'time',
                                (PlotSeriesDefinition('power', 'primary', 'Power'),), 'power.png')
    columns = _native_chart_source_columns(sheet, [_channel('time', 'Time', 's')],
        {'power': definition}, {'time': [0, 1], 'power': [2, 3]}, 5)
    chart = _build_native_scatter_chart(definition, report_sheet=sheet,
        channel_column_by_id=columns, data_start_row=5, data_end_row=6, chart_index=0)
    sheet.add_chart(chart, 'C1')
    path = tmp_path / 'legacy.xlsx'
    workbook.save(path)
    workbook = load_workbook(path)
    assert workbook['Chart Data'].sheet_state == 'hidden'
    assert workbook['Chart Data']['A6'].value == 3
    chart = workbook.active._charts[0]
    assert not chart.visible_cells_only
    assert chart.series[0].yVal.numRef.f == "'Chart Data'!$A$5:$A$6"


@pytest.mark.parametrize('layout', ['engineering', 'sergio_reference'])
def test_configurable_report_converts_image_selection_to_native_chart(tmp_path, layout):
    import yaml
    from vsm_postprocessing.excel_report_engine import generate_excel_report
    def config(name, values):
        path = tmp_path / name
        path.write_text(yaml.safe_dump({'version': 1, **values}))
        return path
    statistics = config('stats.yaml', {'statistics': [
        {'statistic_id': 'maximum', 'channel_id': 'signal__col_002', 'operation': 'max', 'placement_group': 'bottom_channel'}]})
    plots = config('plots.yaml', {'defaults': {'legend': False}, 'plots': [
        {'plot_id': 'signal', 'title': 'Signal', 'x_channel_id': 'time__col_001',
         'series': [{'channel_id': 'signal__col_002'}]}]})
    report = config('report.yaml', {'channels': ['time__col_001'],
        'layout': {'profile': layout}, 'statistics': {'bottom_operations': []},
        'plots': {'include': ['signal']}, 'output': {'filename': 'old_generic.xlsx'}})
    result = generate_excel_report('tests/fixtures/tiny_vsm_input.csv', report, statistics, plots, tmp_path)
    assert result.report_path.name == 'tiny_vsm_input.xlsx'
    workbook = load_workbook(result.report_path)
    assert len(workbook.active._charts) == result.native_excel_chart_count == 1
    assert not workbook.active._images
    assert result.embedded_plot_image_count == 0
    assert workbook.active._charts[0].legend is None
    assert workbook.active._charts[0].series[0].yVal.numRef.f == "'Chart Data'!$A$5:$A$6"


@pytest.mark.parametrize('include_height', [False, True])
@pytest.mark.parametrize('extension', ['.csv', '.xlsx'])
def test_profile_chart_optional_channel_and_source_workbook(tmp_path, include_height, extension):
    import yaml
    from vsm_postprocessing.excel_report_engine import generate_profile_excel_report
    source = tmp_path / ('Test  Run_05' + extension)
    rows = [['Time', 'Gradient'], ['s', '%'], [0, 1], [1, 3], [2, 2]]
    if include_height:
        for row, height in zip(rows, ['Height', 'm', 10, 20, 15]):
            row.append(height)
    if extension == '.csv':
        source.write_text('\n'.join(','.join(map(str, row)) for row in rows))
    else:
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        workbook.save(source)
    original = source.read_bytes()
    raw = [{'semantic_name': name.lower(), 'source_name': name, 'report_name': name,
            'channel_type': 'VSM', 'unit': unit, 'required': name != 'Height'}
           for name, unit in [('Time', 's'), ('Gradient', '%'), ('Height', 'm')]]
    profile = {'version': 1, 'profile': {'profile_id': 'test', 'name': 'Test'},
               'channels': {'raw': raw, 'math': []}, 'statistics': [], 'kpis': [],
               'plots': [{'plot_id': 'road_profile', 'title': 'Road Profile', 'x': 'time',
                          'x_label': 'Time [s]', 'primary_y_label': 'Gradient [%]',
                          'secondary_y_label': 'Height [m]',
                          'series': [{'semantic_name': 'gradient'},
                                     {'semantic_name': 'height', 'axis': 'secondary', 'required': False}]}]}
    config = tmp_path / 'profile.yaml'
    config.write_text(yaml.safe_dump(profile))
    result = generate_profile_excel_report(source, config, tmp_path, machine_name_override='Other Machine')
    assert result.report_path.name == 'Test  Run_05.xlsx'
    assert result.report_path != source
    assert source.read_bytes() == original
    workbook = load_workbook(result.report_path)
    assert len(workbook.active._charts) == 1
    assert len(workbook.active._charts[0]._charts) == (2 if include_height else 1)
    assert list(workbook.active.values)[4][:2] == (0, 1)
    assert len(result.plotting_result.rendered_plots) == 1


@pytest.mark.parametrize('filename', ['RoboSprayer_Electric_Report_Astauto_Colours.pptx',
                                      'Caiman_SP_Hybrid_Report_Astauto_Colours.pptx'])
def test_reusable_layout_retains_original_slides_and_creates_branded_slide(tmp_path, filename):
    prs = Presentation(Path('reference_files') / filename)
    _add_astauto_logo(prs.slides[8], prs)
    original = [str(slide.element.xml) for slide in prs.slides]
    _add_reusable_content_layout(prs)
    assert [str(slide.element.xml) for slide in prs.slides] == original
    path = tmp_path / filename
    prs.save(path)
    prs = Presentation(path)
    layout = prs.slide_layouts[-1]
    assert layout.name == 'Astauto VSM - Additional content'
    assert len(layout.placeholders) == 2
    assert any(shape.shape_type == 13 for shape in layout.shapes)
    assert layout.element.cSld.bg.xml == prs.slides[8].element.cSld.bg.xml
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = 'Motor specifications'
    prs.save(path)
    assert Presentation(path).slides[-1].shapes.title.text == 'Motor specifications'


@pytest.mark.parametrize('sample_count', [3, 7])
@pytest.mark.parametrize('layout', ['profile', 'engineering', 'sergio_reference'])
def test_october_statistics_formulas_and_summary_dependencies(tmp_path, sample_count, layout):
    import yaml
    from vsm_postprocessing.excel_report_engine import generate_excel_report, generate_profile_excel_report
    from vsm_postprocessing.statistics_engine import compute_statistic
    source = tmp_path / 'dynamic.csv'
    values = [4, -2, 9, 1, 3, 8, 5][:sample_count]
    source.write_text('Time,Signal\ns,kW\n' + '\n'.join(f'{i},{v}' for i, v in enumerate(values)))
    operations = ['max', 'min', 'last', 'first']
    def write(name, payload):
        path = tmp_path / name
        path.write_text(yaml.safe_dump({'version': 1, **payload}))
        return path
    if layout == 'profile':
        # Put MATH beyond Z, so a fixed channel letter cannot pass this test.
        math_channels = [{'semantic_name': f'copy_{i}', 'source_name': 'Signal',
                          'report_name': f'Copy {i}', 'unit': 'kW',
                          'dependencies': ['signal'], 'expression': 'signal * 2'} for i in range(26)]
        definitions = [{'statistic_id': f'{target}_{op}', 'target': target, 'operation': op,
                        'placement_group': 'summary'} for target in ['signal', 'copy_25'] for op in operations]
        profile = write('profile.yaml', {'profile': {'profile_id': 'dynamic', 'name': 'Dynamic'},
            'channels': {'raw': [{'semantic_name': name.lower(), 'source_name': name,
                'report_name': name, 'unit': unit, 'channel_type': 'VSM'}
                for name, unit in [('Time', 's'), ('Signal', 'kW')]], 'math': math_channels},
            'statistics': definitions, 'kpis': [{'kpi_id': 'scaled', 'expression': 'signal_max * 3',
                'dependencies': ['signal_max'], 'unit': 'kW'}], 'plots': []})
        result = generate_profile_excel_report(source, profile, tmp_path / 'out')
        columns = {c.channel_id: i for i, c in enumerate(result.report_channels, 1)}
        cases = [(columns['signal'], values), (columns['copy_25'], [v * 2 for v in values])]
        summary_start = result.report_channel_count + 2
        first_bottom = sample_count + 5
    else:
        stats = write('statistics.yaml', {'statistics': [
            {'statistic_id': op, 'channel_id': 'signal__col_002', 'operation': op, 'placement_group': 'bottom_channel'} for op in operations]})
        plots = write('plots.yaml', {'plots': [{'plot_id': 'signal', 'title': 'Signal', 'x_channel_id': 'time__col_001', 'series': [{'channel_id': 'signal__col_002'}]}]})
        config = write('report.yaml', {'channels': ['time__col_001', 'signal__col_002'],
            'layout': {'profile': layout}, 'statistics': {'kpis': operations,
                'bottom_operations': operations, 'bottom_summary': operations}, 'plots': {'include': []}})
        result = generate_excel_report(source, config, stats, plots, tmp_path / 'out')
        cases, summary_start = [(2, values)], 4
        first_bottom = sample_count + (6 if layout == 'engineering' else 5)
    workbook = load_workbook(result.report_path, data_only=False)
    cached = load_workbook(result.report_path, data_only=True)
    sheet, numeric = workbook.active, cached.active
    assert sheet.freeze_panes == 'B6'
    assert workbook.calculation.calcMode == 'auto'
    assert workbook.calculation.fullCalcOnLoad
    for index, (column, data) in enumerate(cases):
        letter = get_column_letter(column)
        extent = f'{letter}5:{letter}{sample_count + 4}'
        for offset, operation in enumerate(operations):
            cell = sheet.cell(first_bottom + offset, column)
            expected = {'max': f'=MAX({extent})', 'min': f'=MIN({extent})',
                        'first': f'={letter}5', 'last': f'={letter}{sample_count + 4}'}[operation]
            assert cell.data_type == 'f' and cell.value == expected
            # Evaluate the definition against the actual exported range, independently of its cache.
            exported = [sheet.cell(row, column).value for row in range(5, sample_count + 5)]
            evaluated = {'max': max(exported), 'min': min(exported), 'first': exported[0], 'last': exported[-1]}[operation]
            python_value = compute_statistic(data, operation)[0]
            assert evaluated == pytest.approx(python_value)
            upper = sheet.cell(4, summary_start + index * 4 + offset)
            assert upper.value == '=' + cell.coordinate
            assert numeric[cell.coordinate].value == pytest.approx(python_value)
            assert numeric[upper.coordinate].value == pytest.approx(python_value)
    if layout == 'profile':
        assert sheet.cell(4, summary_start + 8).value == max(values) * 3
        assert all(sheet.cell(first_bottom + i, 3).value is None for i in range(4))


@pytest.mark.parametrize('profile_name', ['electric', 'hybrid'])
def test_slide_nine_rr_statistics_are_distinct_and_missing_channels_fail(profile_name):
    from vsm_postprocessing.profile_powerpoint_report_engine import _profile_statistics_as_powerpoint_statistics
    from vsm_postprocessing.errors import PowerPointReportError
    profile = load_reporting_profile(f'config/report_profiles/robosprayer_{profile_name}.yaml')
    slide = profile.presentation.slides_by_id()['traction_auxiliaries']
    assert slide.statistics[:2] == ('edu_mech_power_rr_max', 'wheel_power_rr_max')
    signals = {'edu_mech_power_rr': np.array([1., 17., 3.]), 'wheel_power_rr': np.array([11., 2., 5.]),
               'edu_mech_power_rl': np.array([90., 91., 92.]), 'total_edu_mech_power': np.array([200., 201., 202.]),
               'wheel_power_total': np.array([300., 301., 302.])}
    result = SimpleNamespace(profile=profile, dataset=None, report_path=Path('report.xlsx'),
        statistics_result=SimpleNamespace(profile=profile, statistics=[], kpis=[]),
        plotting_result=SimpleNamespace(values_by_semantic_name=signals,
            channels_by_semantic_name={name: _channel(name, name, 'kW') for name in signals}))
    stats = _profile_statistics_as_powerpoint_statistics(result).statistics
    assert [(item.display_name, item.channel_id, item.operation, item.value) for item in stats] == [
        ('RR EDU MAX POWER', 'edu_mech_power_rr', 'max', 17.),
        ('RR WHEEL MAX POWER', 'wheel_power_rr', 'max', 11.)]
    for missing in ('edu_mech_power_rr', 'wheel_power_rr'):
        value = signals.pop(missing)
        with pytest.raises(PowerPointReportError, match=missing):
            _profile_statistics_as_powerpoint_statistics(result)
        signals[missing] = value


@pytest.mark.parametrize('operation', ['max', 'min', 'first', 'last'])
def test_formula_nonfinite_policy_and_caches(tmp_path, operation):
    from vsm_postprocessing.excel_formulas import write_statistic, write_formula, save_report_workbook
    from vsm_postprocessing.statistics_engine import compute_statistic
    workbook = Workbook()
    values = np.array([np.nan, 2., -3., 7., np.nan])
    for row, value in enumerate(values, 5):
        workbook.active.cell(row, 2, value)
    value = compute_statistic(values, operation, 'omit')[0]
    write_statistic(workbook.active['B10'], operation, 5, 9, values, value, 'omit')
    write_formula(workbook.active['C4'], '=B10', value)
    write_statistic(workbook.active['D10'], operation, 5, 9, values, np.nan, 'propagate')
    path = tmp_path / 'nonfinite.xlsx'
    save_report_workbook(workbook, path)
    formulas = load_workbook(path).active
    expected = {'max': '=MAX(B5:B9)', 'min': '=MIN(B5:B9)',
                'first': '=INDEX(B5:B9,MATCH(TRUE,INDEX(ISNUMBER(B5:B9),0),0))',
                'last': '=LOOKUP(2,1/ISNUMBER(B5:B9),B5:B9)'}
    assert formulas['B10'].value == expected[operation]
    assert formulas['D10'].value.startswith('=IF(COUNT(D5:D9)=ROWS(D5:D9),')
    cached = load_workbook(path, data_only=True).active
    assert cached['B10'].value == cached['C4'].value == pytest.approx(value)
